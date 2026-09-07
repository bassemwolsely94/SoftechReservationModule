import datetime as _dt
import logging
import threading

from rest_framework import viewsets, mixins, status, generics
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import IsAuthenticated, IsAdminUser
from rest_framework.response import Response

from .models import StockBatch, NearExpiryAlert, PurchaseExpiryAuditRun
from .serializers import (
    StockBatchSerializer, StockBatchListSerializer,
    NearExpiryAlertSerializer, FEFORecommendationSerializer,
    PurchaseExpiryAuditRunSerializer,
)
from .service import BatchService

logger = logging.getLogger('elrezeiky.batches')


# ── Purchase-Expiry Physical Audit ────────────────────────────────────────────

def _parse_date(val):
    """Parse 'YYYY-MM-DD' (or 'YYYY-MM') to a date, or None."""
    if not val:
        return None
    s = str(val).strip()
    for fmt in ('%Y-%m-%d', '%Y-%m'):
        try:
            return _dt.datetime.strptime(s[:10] if fmt == '%Y-%m-%d' else s[:7], fmt).date()
        except ValueError:
            continue
    return None


def _branch_list(data):
    """Accept branches as a list, a single value, or a comma string."""
    raw = data.get('branches') or data.get('branch') or []
    if isinstance(raw, str):
        raw = [b.strip() for b in raw.split(',') if b.strip()]
    elif not isinstance(raw, (list, tuple)):
        raw = [str(raw)]
    return [str(b).strip() for b in raw if str(b).strip()]


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def purchase_expiry_candidates(request):
    """
    Physical near-expiry audit report.

    Body:
      from        'YYYY-MM-DD' | 'YYYY-MM'   (required — ENTERED-EXPIRY window start)
      to          'YYYY-MM-DD' | 'YYYY-MM'   (required — expiry window end, inclusive)
      branches    []|str|csv                 (optional — default: all in mirror)
      categories  []                         (optional — default main categories)
      only_in_stock bool                     (default true)
      min_qty     number                     (default 0)

    Returns the candidate item worklist (purchase-entry trigger ∩ in-stock now).
    Hits SOFTECH live for the current-stock check, so it may take a few seconds.
    """
    from .expiry_audit import audit_candidates

    date_from = _parse_date(request.data.get('from'))
    date_to   = _parse_date(request.data.get('to'))
    if not date_from or not date_to:
        return Response({'detail': 'يجب تحديد الفترة (from / to).'},
                        status=status.HTTP_400_BAD_REQUEST)
    if date_from > date_to:
        return Response({'detail': 'تاريخ البداية بعد تاريخ النهاية.'},
                        status=status.HTTP_400_BAD_REQUEST)

    branches   = _branch_list(request.data) or None
    categories = request.data.get('categories') or None
    only_in_stock = request.data.get('only_in_stock', True)
    if isinstance(only_in_stock, str):
        only_in_stock = only_in_stock.lower() != 'false'
    try:
        from decimal import Decimal
        min_qty = Decimal(str(request.data.get('min_qty', 0)))
    except Exception:
        min_qty = 0

    sort = request.data.get('sort') or None
    imported_only = request.data.get('imported_only', False)
    if isinstance(imported_only, str):
        imported_only = imported_only.lower() == 'true'
    min_var = request.data.get('min_value_at_risk')

    try:
        rows = audit_candidates(
            date_from, date_to, branch_codes=branches,
            categories=categories, only_in_stock=only_in_stock, min_qty=min_qty,
            sort=sort, imported_only=imported_only, min_value_at_risk=min_var,
        )
    except Exception as e:
        logger.exception('purchase_expiry_candidates failed')
        return Response({'detail': f'خطأ أثناء توليد التقرير: {e}'},
                        status=status.HTTP_503_SERVICE_UNAVAILABLE)

    from apps.config.models import SystemSetting
    return Response({
        'from': date_from.isoformat(), 'to': date_to.isoformat(),
        'branches': branches or 'ALL',
        'only_in_stock': only_in_stock,
        'count': len(rows),
        'items': rows,
        # A5.2: whether the "request markdown" action is enabled (default OFF)
        'markdown_enabled': bool(SystemSetting.get('near_expiry_markdown_enabled', False)),
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def purchase_expiry_request_markdown(request):
    """
    A5.2 — create a PENDING near-expiry markdown request in the EXISTING
    discount-approvals channel (special_discp). A human approves it there, which
    triggers the established SOFTECH writeback; reversal uses that module's
    rollback. Nothing is written to SOFTECH here — this only files the request.

    Gated by SystemSetting 'near_expiry_markdown_enabled' (default OFF).
    Body: { item_code (required), discount_pct (required), days_to_expiry?, branch?, reason? }
    """
    from apps.config.models import SystemSetting
    if not SystemSetting.get('near_expiry_markdown_enabled', False):
        return Response({'detail': 'خصم قرب انتهاء الصلاحية غير مُفعّل من الإعدادات.'},
                        status=status.HTTP_403_FORBIDDEN)

    from apps.catalog.models import Item
    from apps.discount_approvals.models import ItemPriceChangeRequest, USER_EDITABLE_FIELDS

    item_code = str(request.data.get('item_code') or '').strip()
    try:
        discount_pct = round(float(request.data.get('discount_pct')), 2)
    except (TypeError, ValueError):
        return Response({'detail': 'discount_pct مطلوب.'}, status=status.HTTP_400_BAD_REQUEST)
    if not (0 < discount_pct <= 100):
        return Response({'detail': 'نسبة الخصم يجب أن تكون بين 0 و100.'},
                        status=status.HTTP_400_BAD_REQUEST)

    item = Item.objects.filter(softech_id=item_code).first()
    if not item:
        return Response({'detail': 'الصنف غير موجود.'}, status=status.HTTP_404_NOT_FOUND)

    # Idempotent: one OPEN near-expiry markdown per item (no stacking / double-post).
    existing = ItemPriceChangeRequest.objects.filter(
        item=item, source='near_expiry',
        status__in=[ItemPriceChangeRequest.STATUS_PENDING,
                    ItemPriceChangeRequest.STATUS_APPROVED],
    ).first()
    if existing:
        return Response({'detail': f'يوجد طلب خصم قائم لهذا الصنف (#{existing.pk}).',
                         'request_id': existing.pk}, status=status.HTTP_409_CONFLICT)

    # Snapshot current priceable values exactly like the discount-approvals create view.
    old_values = {}
    for dj, _, _ in USER_EDITABLE_FIELDS:
        v = getattr(item, dj, None)
        old_values[dj] = str(v) if v is not None else '0'
    for af in ('pack_price_tax', 'unit_price'):
        v = getattr(item, af, None)
        old_values[af] = str(v) if v is not None else '0'

    days = request.data.get('days_to_expiry')
    branch = str(request.data.get('branch') or '').strip()
    extra = str(request.data.get('reason') or '').strip()
    reason = (f"خصم قرب انتهاء الصلاحية {discount_pct}% (خصم خاص) — {item.name} — "
              f"أيام حتى الصلاحية: {days} — فرع: {branch or 'الكل'}. "
              f"يُسمح بالبيع تحت التكلفة للتصريف."
              + (f" | {extra}" if extra else ''))

    obj = ItemPriceChangeRequest.objects.create(
        item=item, requested_by=request.user,
        old_values=old_values,
        new_values={'special_discp': str(discount_pct)},
        reason=reason, source='near_expiry',
        status=ItemPriceChangeRequest.STATUS_PENDING,
    )
    try:
        from apps.discount_approvals.notify import notify_admins_new_request
        notify_admins_new_request(obj)
    except Exception:
        logger.exception('notify_admins_new_request failed for markdown #%s', obj.pk)

    return Response({'detail': 'تم إنشاء طلب الخصم — بانتظار الاعتماد في «اعتماد الأسعار».',
                     'request_id': obj.pk}, status=status.HTTP_201_CREATED)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def export_purchase_expiry(request):
    """
    Export the (already filtered/sorted) audit rows the client is showing to xlsx.
    Body: { items: [row, ...], from, to, branch_label }. Posting the displayed
    rows keeps the file identical to what's on screen (client-side filters/sort)
    without re-hitting SOFTECH.
    """
    from django.http import HttpResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    items = request.data.get('items') or []
    date_from = str(request.data.get('from') or '')
    date_to = str(request.data.get('to') or '')
    branch_label = str(request.data.get('branch_label') or 'كل الفروع')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'تدقيق الصلاحيات'
    ws.sheet_view.rightToLeft = True

    cols = [
        ('كود', 'item_code', 12),
        ('الصنف', 'item_name', 42),
        ('الكمية الحالية', 'current_qty', 12),
        ('تكلفة الوحدة', 'unit_cost', 12),
        ('قيمة معرّضة للخطر', 'value_at_risk', 16),
        ('خصم مقترح %', 'markdown_discount_pct', 12),
        ('سعر بعد الخصم', 'markdown_net_price', 13),
        ('خسارة متوقعة', 'expected_loss', 14),
        ('عمر بالفرع (يوم)', 'stock_age_days', 14),
        ('أقدم وصول', 'oldest_arrival_date', 12),
        ('التصنيف العام', 'medicine_type', 18),
        ('مستورد', 'is_imported', 8),
        ('صنف ثلاجة', 'is_fridge', 9),
        ('المنشأ', 'origin', 14),
        ('الشكل', 'shape', 14),
        ('المنتج', 'producer', 20),
        ('العائلة', 'family', 20),
        ('الفروع', 'branches_in_stock', 14),
        ('صلاحية مُدخَلة (من)', 'earliest_entered_expiry', 14),
        ('صلاحية مُدخَلة (إلى)', 'latest_entered_expiry', 14),
        ('منتهية', 'has_entered_expiry_passed', 8),
        ('عدد الإدخالات', 'entry_count', 10),
        ('الموردون', 'suppliers', 40),
    ]

    # Title + meta
    ws.append([f'تقرير تدقيق صلاحيات الشراء — {branch_label} — صلاحية {date_from} : {date_to}'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])

    header_row = 3
    ws.append([c[0] for c in cols])
    hfill = PatternFill('solid', fgColor='022871')
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = hfill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = c[2]

    def _fmt(row, key):
        v = row.get(key)
        if key in ('is_imported', 'is_fridge', 'has_entered_expiry_passed'):
            return 'نعم' if v else ''
        if key == 'branches_in_stock':
            return '، '.join(v or [])
        if key == 'suppliers':
            return '، '.join(v or [])
        return v if v is not None else ''

    red = Font(color='C00000')
    for row in items:
        ws.append([_fmt(row, c[1]) for c in cols])
        if row.get('has_entered_expiry_passed'):
            for i in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=i).font = red

    ws.freeze_panes = f'A{header_row + 1}'

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f'expiry_audit_{branch_label}_{date_from}_{date_to}.xlsx'.replace(' ', '_')
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = f'attachment; filename="expiry_audit.xlsx"'
    return resp


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def purchase_expiry_rebalance_suggest(request):
    """
    A4 — inter-branch rebalancing suggestion for one near-expiry item.
    Body: { item_code (required), from_branch (softech code), days_to_expiry }.
    Read-only; returns a transfer plan (target branches + qty). The frontend
    turns a plan line into a transfer request via the existing /transfers/ API.
    """
    from .expiry_audit import rebalance_suggest

    item_code = str(request.data.get('item_code') or '').strip()
    if not item_code:
        return Response({'detail': 'كود الصنف مطلوب.'}, status=status.HTTP_400_BAD_REQUEST)
    from_branch = str(request.data.get('from_branch') or '').strip() or None
    try:
        days_to_expiry = int(request.data.get('days_to_expiry'))
    except (TypeError, ValueError):
        return Response({'detail': 'days_to_expiry مطلوب (عدد الأيام حتى انتهاء الصلاحية).'},
                        status=status.HTTP_400_BAD_REQUEST)

    try:
        result = rebalance_suggest(item_code, from_branch, days_to_expiry)
    except Exception as e:
        logger.exception('rebalance_suggest failed for %s', item_code)
        return Response({'detail': f'خطأ أثناء حساب الاقتراح: {e}'},
                        status=status.HTTP_503_SERVICE_UNAVAILABLE)
    return Response(result)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def purchase_expiry_prone_items(request):
    """
    C2 — procurement feedback: items chronically bought short-dated (high share of
    main-supplier purchases arriving with < N months shelf life), with velocity,
    flagged for reorder review. Read-only.
    Query: ?months_back=&short_dated_months=&min_short_pct=
    """
    from .expiry_audit import expiry_prone_items

    def _num(name, default):
        v = request.query_params.get(name)
        try:
            return type(default)(v) if v not in (None, '') else default
        except (TypeError, ValueError):
            return default

    rows = expiry_prone_items(
        months_back=_num('months_back', 12),
        short_dated_months=_num('short_dated_months', 6),
        min_short_pct=_num('min_short_pct', 30.0),
    )
    return Response({'count': len(rows), 'items': rows})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def purchase_expiry_supplier_scorecard(request):
    """
    B2 — supplier dating scorecard: how well-dated is each main supplier's stock?
    Query: ?months_back=&short_dated_months=&categories=a,b  (all optional).
    Read-only, pure PG aggregation over the mirror.
    """
    from .expiry_audit import supplier_scorecard

    def _int(name):
        v = request.query_params.get(name)
        try:
            return int(v) if v not in (None, '') else None
        except ValueError:
            return None

    cats = request.query_params.get('categories')
    categories = [c.strip() for c in cats.split(',') if c.strip()] if cats else None
    rows = supplier_scorecard(
        categories=categories,
        months_back=_int('months_back'),
        short_dated_months=_int('short_dated_months') or 6,
    )
    return Response({'count': len(rows), 'suppliers': rows})


class PurchaseExpiryRunListView(generics.ListAPIView):
    """GET — recent purchase-expiry backfill runs (status/counters)."""
    permission_classes = [IsAuthenticated]
    serializer_class   = PurchaseExpiryAuditRunSerializer
    queryset           = PurchaseExpiryAuditRun.objects.all()[:30]


@api_view(['POST'])
@permission_classes([IsAdminUser])
def trigger_purchase_expiry_sync(request):
    """
    Kick off a purchase-expiry backfill in a background thread (admin only).

    Body: years|from|to, categories (all optional). Always mirrors ALL branches
    (any 'branch' is ignored — the mirror must stay chain-wide).
    Re-runnable — after re-classifying more suppliers as "main", call this again
    to pull the newly-included suppliers' history.
    """
    from .expiry_audit import run_backfill, MAIN_SUPPLIER_CATEGORIES

    today = _dt.date.today()
    date_to = _parse_date(request.data.get('to')) or today
    date_from = _parse_date(request.data.get('from'))
    if not date_from:
        years = int(request.data.get('years', 3) or 3)
        try:
            date_from = date_to.replace(year=date_to.year - years, day=1)
        except ValueError:
            date_from = date_to.replace(year=date_to.year - years, day=1, month=date_to.month)
    if date_from > date_to:
        return Response({'detail': 'الفترة غير صحيحة.'}, status=status.HTTP_400_BAD_REQUEST)

    categories = request.data.get('categories') or list(MAIN_SUPPLIER_CATEGORIES)
    triggered_by = str(request.user)

    # Always mirror ALL branches — the mirror must stay complete so every branch's
    # report works (purchases are centralized at HQ then distributed). Any 'branch'
    # in the request is intentionally ignored.
    run = PurchaseExpiryAuditRun.objects.create(
        window_from=date_from, window_to=date_to, branch_scope='',
        categories=categories, triggered_by=triggered_by,
    )

    def _run():
        try:
            run_backfill(run, date_from, date_to, branch=None,
                         categories=categories)
            run.finish('success')
        except Exception as e:   # noqa: BLE001
            logger.exception('purchase-expiry backfill failed (run #%s)', run.pk)
            run.finish('failed', error=str(e))

    threading.Thread(target=_run, daemon=True).start()

    return Response({
        'detail': 'بدأت مزامنة صلاحيات الشراء في الخلفية.',
        'run': PurchaseExpiryAuditRunSerializer(run).data,
    }, status=status.HTTP_202_ACCEPTED)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def spawn_expiry_count_session(request):
    """
    Create a physical stock-count session (mode=expiry_audit) from the audit
    report, so staff can pull the items and record their real shelf expiry.

    Body:
      branch      required — one branch per count session
      from / to   required — the ENTERED-EXPIRY window that defined the candidates
      categories  optional
      item_codes  optional — explicit subset; else all candidates for the branch
      name        optional
    """
    from .expiry_audit import audit_candidates, MAIN_SUPPLIER_CATEGORIES
    from .models import PurchaseExpiryEntry
    from apps.stockcount.models import StockCountSession, StockCountSnapshot
    from apps.stockcount.engine import generate_snapshot

    branch = (request.data.get('branch') or '').strip()
    date_from = _parse_date(request.data.get('from'))
    date_to   = _parse_date(request.data.get('to'))
    if not branch:
        return Response({'detail': 'كود الفرع مطلوب.'}, status=status.HTTP_400_BAD_REQUEST)
    if not date_from or not date_to:
        return Response({'detail': 'يجب تحديد الفترة (from / to).'},
                        status=status.HTTP_400_BAD_REQUEST)

    categories = request.data.get('categories') or None
    item_codes = request.data.get('item_codes')
    if item_codes and isinstance(item_codes, (list, tuple)):
        codes = [str(c).strip() for c in item_codes if str(c).strip()]
    else:
        try:
            cands = audit_candidates(
                date_from, date_to, branch_codes=[branch],
                categories=categories, only_in_stock=True,
            )
        except Exception as e:
            logger.exception('spawn_expiry_count_session candidate lookup failed')
            return Response({'detail': f'خطأ أثناء جلب الأصناف: {e}'},
                            status=status.HTTP_503_SERVICE_UNAVAILABLE)
        codes = [c['item_code'] for c in cands]

    if not codes:
        return Response({'detail': 'لا توجد أصناف مطابقة للمعايير في هذا الفرع.'},
                        status=status.HTTP_400_BAD_REQUEST)

    profile = getattr(request.user, 'staff_profile', None)
    name = (request.data.get('name') or
            f'جرد صلاحية — فرع {branch} — {date_from} : {date_to}')

    session = StockCountSession.objects.create(
        name=name, mode='expiry_audit', branch_code=branch,
        date_from=date_from, date_to=date_to,
        item_codes_filter=codes, created_by=profile,
        notes='مولّدة من محرك تدقيق صلاحيات الشراء (أصناف من موردين رئيسيين).',
    )

    # Snapshot current stock for exactly these items (filtered-stock path).
    try:
        count = generate_snapshot(session)
    except ValueError as e:
        session.delete()
        return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
    except Exception as e:
        session.delete()
        logger.exception('spawn_expiry_count_session snapshot failed')
        return Response({'detail': f'خطأ في الاتصال بـ SOFTECH: {e}'},
                        status=status.HTTP_503_SERVICE_UNAVAILABLE)

    # Backfill the entered-expiry hint (earliest keyed expiry within the window)
    # onto each snapshot. Chain-wide (NOT filtered by branch) — purchases are
    # received centrally then distributed, so the expiry entry lives at HQ while
    # the stock sits at this branch. Same semantics as the report.
    from django.db.models import Min
    cats = categories or list(MAIN_SUPPLIER_CATEGORIES)
    hint_qs = (
        PurchaseExpiryEntry.objects
        .filter(entered_expiry__gte=date_from, entered_expiry__lte=date_to,
                supplier_category__in=cats, item_code__in=codes)
        .values('item_code').annotate(hint=Min('entered_expiry'))
    )
    hints = {r['item_code']: r['hint'] for r in hint_qs}
    to_update = []
    for snap in StockCountSnapshot.objects.filter(session=session):
        h = hints.get(snap.item_code)
        if h:
            snap.entered_expiry_hint = h
            to_update.append(snap)
    if to_update:
        StockCountSnapshot.objects.bulk_update(to_update, ['entered_expiry_hint'])

    return Response({
        'detail': f'تم إنشاء جلسة جرد صلاحية لـ {count} صنف.',
        'session_id': session.pk,
        'item_count': count,
        'branch_code': branch,
        'status': session.status,
    }, status=status.HTTP_201_CREATED)


def _store_list(params):
    raw = params.get('stores') or params.get('store') or ''
    if isinstance(raw, (list, tuple)):
        return [str(s).strip() for s in raw if str(s).strip()]
    return [s.strip() for s in str(raw).split(',') if s.strip()]


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def stock_expiry_summary_view(request):
    """KPI cards for the FEFO tabs from the StockExpiryBalance mirror (all nodes)."""
    from .stock_expiry import stock_expiry_summary
    branches = _branch_list(request.query_params) or None
    stores = _store_list(request.query_params) or None
    inc_q = str(request.query_params.get('include_quarantine', '')).lower() == 'true'
    return Response(stock_expiry_summary(branch_codes=branches, store_codes=stores,
                                         include_quarantine=inc_q))


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def stock_expiry_stores_view(request):
    """Distinct stores (HQ/branch warehouses) in the mirror — powers the store filter."""
    from .stock_expiry import stock_expiry_stores
    branches = _branch_list(request.query_params) or None
    return Response({'stores': stock_expiry_stores(branch_codes=branches)})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def stock_expiry_report_view(request):
    """
    Near-expiry / expired report over the mirror.
    Query: mode=near|expired|range|all, within_days, exp_from, exp_to, branches,
           stores, include_quarantine, imported_only, fridge_only, medicine_type,
           sort, limit.
    """
    from .stock_expiry import stock_expiry_report
    p = request.query_params
    branches = _branch_list(p) or None
    stores = _store_list(p) or None

    def _b(name):
        return str(p.get(name, '')).lower() == 'true'

    try:
        within = int(p.get('within_days', 180))
    except (TypeError, ValueError):
        within = 180
    rows = stock_expiry_report(
        mode=p.get('mode', 'near'), within_days=within, branch_codes=branches,
        store_codes=stores, exp_from=_parse_date(p.get('exp_from')),
        exp_to=_parse_date(p.get('exp_to')),
        include_quarantine=_b('include_quarantine'), imported_only=_b('imported_only'),
        fridge_only=_b('fridge_only'), medicine_type=(p.get('medicine_type') or None),
        sort=(p.get('sort') or None),
    )
    return Response({'count': len(rows), 'items': rows})


class StockExpirySyncRunListView(generics.ListAPIView):
    """GET — recent stock-expiry sync runs (per-node status)."""
    permission_classes = [IsAuthenticated]

    def get(self, request, *args, **kwargs):
        from .models import StockExpirySyncRun
        runs = StockExpirySyncRun.objects.all()[:20]
        return Response([{
            'id': r.id, 'started_at': r.started_at, 'finished_at': r.finished_at,
            'status': r.status, 'nodes_total': r.nodes_total, 'nodes_ok': r.nodes_ok,
            'nodes_failed': r.nodes_failed, 'rows_synced': r.rows_synced,
            'detail': r.detail, 'triggered_by': r.triggered_by,
        } for r in runs])


@api_view(['POST'])
@permission_classes([IsAdminUser])
def stock_expiry_sync_trigger(request):
    """Kick off a multi-node stkbalexpiry sync in the background (admin)."""
    from .models import StockExpirySyncRun
    from .stock_expiry import sync_stock_expiry_all
    branch = str(request.data.get('branch') or '').strip() or None
    run = StockExpirySyncRun.objects.create(triggered_by=str(request.user))

    def _run():
        try:
            sync_stock_expiry_all(run=run, only_branch=branch)
        except Exception as e:   # noqa: BLE001
            logger.exception('stock-expiry sync failed (run #%s)', run.pk)
            run.status = 'failed'; run.detail = {'fatal': str(e)[:300]}; run.finish('failed')

    threading.Thread(target=_run, daemon=True).start()
    return Response({'detail': 'بدأت مزامنة صلاحية المخزون في الخلفية.', 'run_id': run.pk},
                    status=status.HTTP_202_ACCEPTED)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def export_stock_expiry(request):
    """Export the displayed live stock-expiry rows to xlsx (Latin digits, RTL)."""
    from django.http import HttpResponse
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    items = request.data.get('items') or []
    label = str(request.data.get('label') or 'كل الفروع')
    mode = str(request.data.get('mode') or '')
    _TIER_AR = {'expired': 'منتهية', 'critical': '≤30 يوم', 'high': '≤90 يوم',
                'watch': '≤180 يوم', 'ok': '>180 يوم'}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'صلاحية المخزون'
    ws.sheet_view.rightToLeft = True
    cols = [
        ('كود', 'item_code', 12), ('الصنف', 'item_name', 40),
        ('الكمية', 'total_qty', 12), ('تكلفة الوحدة', 'unit_cost', 12),
        ('قيمة معرّضة', 'value_at_risk', 14), ('أقرب صلاحية', 'earliest_expiry', 14),
        ('الحالة', 'tier', 12), ('دفعات', 'batch_count', 8),
        ('التصنيف العام', 'medicine_type', 18), ('المنشأ', 'origin', 14),
        ('مستورد', 'is_imported', 8), ('ثلاجة', 'is_fridge', 8),
        ('الفروع', 'branches', 16),
    ]
    ws.append([f'أرصدة الصلاحية — {label} — {mode}'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([c[0] for c in cols])
    fill = PatternFill('solid', fgColor='022871')
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = c[2]

    def _f(row, key):
        v = row.get(key)
        if key in ('is_imported', 'is_fridge'):
            return 'نعم' if v else ''
        if key == 'tier':
            return _TIER_AR.get(v, v or '')
        if key == 'branches':
            return '، '.join(v or [])
        return v if v is not None else ''

    red = Font(color='C00000')
    for row in items:
        ws.append([_f(row, c[1]) for c in cols])
        if row.get('tier') == 'expired':
            for i in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=i).font = red
    ws.freeze_panes = 'A4'

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="stock_expiry.xlsx"'
    return resp


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def spawn_stock_expiry_count(request):
    """
    Create a physical stock-count session (mode=expiry_audit) from the LIVE
    stock-expiry grid — for one branch + the given item_codes. The entered-expiry
    hint comes from the stkbalexpiry mirror (the actual on-hand expiry).
    Body: { branch (required), item_codes (required), name? }.
    """
    from apps.stockcount.models import StockCountSession, StockCountSnapshot
    from apps.stockcount.engine import generate_snapshot
    from django.db.models import Min
    from .models import StockExpiryBalance

    branch = str(request.data.get('branch') or '').strip()
    raw = request.data.get('item_codes') or []
    codes = [str(c).strip() for c in raw if str(c).strip()] if isinstance(raw, (list, tuple)) else []
    if not branch:
        return Response({'detail': 'كود الفرع مطلوب.'}, status=status.HTTP_400_BAD_REQUEST)
    if not codes:
        return Response({'detail': 'لا توجد أصناف محددة.'}, status=status.HTTP_400_BAD_REQUEST)

    profile = getattr(request.user, 'staff_profile', None)
    name = request.data.get('name') or f'جرد صلاحية (حي) — فرع {branch}'
    session = StockCountSession.objects.create(
        name=name, mode='expiry_audit', branch_code=branch,
        item_codes_filter=codes, created_by=profile,
        notes='مولّدة من أرصدة صلاحية المخزون الحية (stkbalexpiry).',
    )
    try:
        count = generate_snapshot(session)
    except ValueError as e:
        session.delete()
        return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
    except Exception as e:
        session.delete()
        logger.exception('spawn_stock_expiry_count snapshot failed')
        return Response({'detail': f'خطأ في الاتصال بـ SOFTECH: {e}'},
                        status=status.HTTP_503_SERVICE_UNAVAILABLE)

    # Hint = earliest on-hand expiry for this item at this branch (from the mirror).
    hints = {r['item_code']: r['h'] for r in StockExpiryBalance.objects
             .filter(branch_code=branch, item_code__in=codes)
             .values('item_code').annotate(h=Min('expiry_date'))}
    to_update = []
    for snap in StockCountSnapshot.objects.filter(session=session):
        h = hints.get(snap.item_code)
        if h:
            snap.entered_expiry_hint = h
            to_update.append(snap)
    if to_update:
        StockCountSnapshot.objects.bulk_update(to_update, ['entered_expiry_hint'])

    return Response({'detail': f'تم إنشاء جلسة جرد صلاحية لـ {count} صنف.',
                     'session_id': session.pk, 'item_count': count, 'branch_code': branch,
                     'status': session.status}, status=status.HTTP_201_CREATED)


def _disposal_dict(r):
    return {
        'id': r.id, 'item_code': r.item_code, 'item_name': r.item_name,
        'branch_code': r.branch_code, 'store_code': r.store_code, 'batch_no': r.batch_no,
        'expiry_date': r.expiry_date.isoformat() if r.expiry_date else None,
        'qty': float(r.qty), 'unit_cost': float(r.unit_cost), 'value': float(r.value),
        'decision': r.decision, 'decision_display': r.get_decision_display(),
        'supplier_code': r.supplier_code, 'supplier_name': r.supplier_name,
        'reason': r.reason, 'status': r.status, 'status_display': r.get_status_display(),
        'created_by': r.created_by.full_name if r.created_by else None,
        'created_at': r.created_at,
        'decided_by': r.decided_by.full_name if r.decided_by else None,
        'decided_at': r.decided_at,
    }


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def expiry_disposal(request):
    """
    GET  — list disposal/return decisions (filter: status, branch, decision).
    POST — create one from a flagged expired/near-expiry item (status=draft).
           Body: item_code, branch_code, qty, decision, expiry_date?, store_code?,
                 batch_no?, supplier_code?, reason?.
    Read/record only — no SOFTECH write.
    """
    from .models import ExpiryDisposalRecord

    if request.method == 'GET':
        qs = ExpiryDisposalRecord.objects.select_related('created_by', 'decided_by')
        p = request.query_params
        if p.get('status'):
            qs = qs.filter(status=p['status'])
        if p.get('branch'):
            qs = qs.filter(branch_code=p['branch'])
        if p.get('decision'):
            qs = qs.filter(decision=p['decision'])
        rows = [_disposal_dict(r) for r in qs[:1000]]
        return Response({'count': len(rows), 'items': rows})

    # POST — create
    from apps.catalog.models import Item
    d = request.data
    item_code = str(d.get('item_code') or '').strip()
    branch = str(d.get('branch_code') or '').strip()
    decision = str(d.get('decision') or '').strip()
    valid = {c[0] for c in ExpiryDisposalRecord.DECISION_CHOICES}
    if not item_code or not branch:
        return Response({'detail': 'كود الصنف والفرع مطلوبان.'}, status=status.HTTP_400_BAD_REQUEST)
    if decision not in valid:
        return Response({'detail': 'قرار غير صالح.'}, status=status.HTTP_400_BAD_REQUEST)
    from decimal import Decimal, InvalidOperation
    try:
        qty = Decimal(str(d.get('qty')))
    except (InvalidOperation, TypeError, ValueError):
        qty = None
    if not qty or qty <= 0:
        return Response({'detail': 'الكمية يجب أن تكون أكبر من صفر.'}, status=status.HTTP_400_BAD_REQUEST)

    item = Item.objects.filter(softech_id=item_code).only('name', 'cost_price').first()
    unit_cost = float(item.cost_price) if (item and item.cost_price is not None) else 0.0
    rec = ExpiryDisposalRecord.objects.create(
        item_code=item_code, item_name=(item.name if item else '') or str(d.get('item_name') or ''),
        branch_code=branch, store_code=str(d.get('store_code') or ''),
        batch_no=str(d.get('batch_no') or ''), expiry_date=_parse_date(d.get('expiry_date')),
        qty=qty, unit_cost=unit_cost, value=round(float(qty) * unit_cost, 2),
        decision=decision, supplier_code=str(d.get('supplier_code') or ''),
        supplier_name=str(d.get('supplier_name') or ''), reason=str(d.get('reason') or ''),
        created_by=getattr(request.user, 'staff_profile', None),
    )
    return Response(_disposal_dict(rec), status=status.HTTP_201_CREATED)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def expiry_disposal_status(request, pk):
    """Transition a disposal record: approve / done / cancel. approve+done are
    gated to admin / quality_manager / supervisor. No SOFTECH write."""
    from django.utils import timezone
    from .models import ExpiryDisposalRecord
    try:
        rec = ExpiryDisposalRecord.objects.get(pk=pk)
    except ExpiryDisposalRecord.DoesNotExist:
        return Response({'detail': 'غير موجود'}, status=status.HTTP_404_NOT_FOUND)

    new_status = str(request.data.get('status') or '').strip()
    if new_status not in {'approved', 'done', 'cancelled', 'draft'}:
        return Response({'detail': 'حالة غير صالحة.'}, status=status.HTTP_400_BAD_REQUEST)

    profile = getattr(request.user, 'staff_profile', None)
    role = getattr(profile, 'role', None)
    if new_status in {'approved', 'done'} and role not in {'admin', 'quality_manager', 'supervisor'}:
        return Response({'detail': 'يتطلب اعتماد الإتلاف/المرتجع صلاحية مدير/جودة.'},
                        status=status.HTTP_403_FORBIDDEN)
    if new_status == 'done' and rec.status != 'approved':
        return Response({'detail': 'يجب اعتماد القرار قبل تنفيذه.'}, status=status.HTTP_400_BAD_REQUEST)

    rec.status = new_status
    if request.data.get('notes'):
        rec.reason = (rec.reason + f"\n[{new_status}] " + str(request.data['notes'])).strip()
    if new_status in {'approved', 'done'}:
        rec.decided_by = profile
        rec.decided_at = timezone.now()
    rec.save(update_fields=['status', 'reason', 'decided_by', 'decided_at'])
    return Response(_disposal_dict(rec))


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def export_disposal(request):
    """Export the disposal/return ledger (or a filtered subset) to xlsx — the
    printable إتلاف/مرتجع document. Body: items (list of dicts as shown)."""
    from django.http import HttpResponse
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    items = request.data.get('items') or []
    wb = openpyxl.Workbook(); ws = wb.active
    ws.title = 'إتلاف-مرتجع'; ws.sheet_view.rightToLeft = True
    cols = [('كود', 'item_code', 12), ('الصنف', 'item_name', 38), ('الفرع', 'branch_code', 8),
            ('الكمية', 'qty', 10), ('تكلفة الوحدة', 'unit_cost', 12), ('القيمة', 'value', 12),
            ('الصلاحية', 'expiry_date', 12), ('القرار', 'decision_display', 14),
            ('المورد', 'supplier_name', 22), ('الحالة', 'status_display', 12), ('السبب', 'reason', 30)]
    ws.append(['مستند إتلاف / مرتجع الأصناف منتهية الصلاحية'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws['A1'].font = Font(bold=True, size=13); ws.append([])
    ws.append([c[0] for c in cols])
    fill = PatternFill('solid', fgColor='022871')
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=i)
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = c[2]
    for row in items:
        ws.append([row.get(c[1]) if row.get(c[1]) is not None else '' for c in cols])
    ws.freeze_panes = 'A4'
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="disposal_return.xlsx"'
    return resp


class StockBatchViewSet(
    mixins.ListModelMixin,
    mixins.RetrieveModelMixin,
    viewsets.GenericViewSet,
):
    """
    GET /api/batches/                    — list (filter: item, branch, vendor, expiring_in_days)
    GET /api/batches/{id}/               — full detail with movements
    GET /api/batches/fefo/?item=&branch= — FEFO-ordered list for dispatch
    GET /api/batches/near-expiry/        — near-expiry dashboard summary
    GET /api/batches/alerts/             — NearExpiryAlert list
    POST /api/batches/{id}/quarantine/   — quarantine a batch
    """
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = StockBatch.objects.select_related('item', 'branch', 'vendor', 'received_by')
        p  = self.request.query_params

        if p.get('item'):
            qs = qs.filter(item_id=p['item'])
        if p.get('branch'):
            qs = qs.filter(branch_id=p['branch'])
        if p.get('vendor'):
            qs = qs.filter(vendor_id=p['vendor'])
        if p.get('expiring_in_days'):
            from django.utils import timezone
            from datetime import timedelta
            cutoff = timezone.now().date() + timedelta(days=int(p['expiring_in_days']))
            qs = qs.filter(expiry_date__lte=cutoff, is_expired=False, current_qty__gt=0)
        if p.get('is_quarantined'):
            qs = qs.filter(is_quarantined=p['is_quarantined'].lower() == 'true')
        if p.get('is_expired'):
            qs = qs.filter(is_expired=p['is_expired'].lower() == 'true')

        return qs.order_by('expiry_date', 'id')

    def get_serializer_class(self):
        if self.action in ('list', 'fefo'):
            return StockBatchListSerializer
        return StockBatchSerializer

    @action(detail=False, methods=['get'])
    def fefo(self, request):
        """FEFO-ordered active batches for a given item × branch."""
        item_id   = request.query_params.get('item')
        branch_id = request.query_params.get('branch')
        if not item_id or not branch_id:
            return Response({'detail': 'item و branch مطلوبان.'}, status=status.HTTP_400_BAD_REQUEST)

        batches = BatchService.fefo_batches(int(item_id), int(branch_id))
        serializer = FEFORecommendationSerializer(batches, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='near-expiry')
    def near_expiry(self, request):
        """Near-expiry KPI summary for dashboard."""
        branch_id = request.query_params.get('branch')
        days      = int(request.query_params.get('days', 180))
        data      = BatchService.near_expiry_summary(
            branch_id=int(branch_id) if branch_id else None,
            days=days,
        )
        return Response(data)

    @action(detail=False, methods=['get'])
    def alerts(self, request):
        """Active (unresolved) NearExpiryAlert list."""
        qs = NearExpiryAlert.objects.filter(
            resolved_at__isnull=True
        ).select_related('batch', 'batch__item', 'batch__branch', 'batch__vendor')

        if request.query_params.get('branch'):
            qs = qs.filter(batch__branch_id=request.query_params['branch'])
        if request.query_params.get('threshold'):
            qs = qs.filter(threshold_days=request.query_params['threshold'])

        qs = qs.order_by('batch__expiry_date')
        serializer = NearExpiryAlertSerializer(qs, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def quarantine(self, request, pk=None):
        """Put a batch into quarantine (quality hold)."""
        batch  = self.get_object()
        reason = request.data.get('reason', '')
        if not reason:
            return Response({'detail': 'سبب العزل مطلوب.'}, status=status.HTTP_400_BAD_REQUEST)

        # Require approval for quarantine if not admin
        profile = request.user.staff_profile
        if profile.role not in ('admin', 'quality_manager'):
            # Submit approval request first; actual quarantine happens on approval
            from apps.approvals.service import ApprovalService
            ApprovalService.submit(
                workflow_code='batch_quarantine',
                subject_object=batch,
                title=f'عزل دفعة: {batch.item.name} — {batch.batch_number}',
                requested_by=profile,
                context_data={
                    'batch_id': batch.pk,
                    'reason':   reason,
                    'qty':      str(batch.current_qty),
                },
            )
            return Response(
                {'detail': 'تم رفع طلب عزل الدفعة للاعتماد.'},
                status=status.HTTP_202_ACCEPTED,
            )

        BatchService.quarantine(batch=batch, reason=reason, performed_by=profile)
        return Response({'detail': 'تم عزل الدفعة.'}, status=status.HTTP_200_OK)
