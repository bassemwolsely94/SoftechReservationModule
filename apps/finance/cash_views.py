"""
apps/finance/cash_views.py

P-40 «النقدية ١٣ أسبوع» API (docs/product-review/06_SPEC_P-40_13_week_cash_view.md). Company-level finance: the
finance matrix grant AND an all-branch user (owner decision B-1). Read-only toward SOFTECH; the inputs are finance
entries, each change audited in ReconAuditEvent ('cash_input_changed', before/after).

  GET        /api/finance/cash-view/?as_of=&purchase_budget_weekly=
  GET        /api/finance/cash-view/lines/<key>/?week=          rows behind one cell (week 0-based)
  GET/POST   /api/finance/cash-view/opening-balance/            {as_of, amount, note}
  GET/POST/PATCH /api/finance/cash-view/fixed-outflows/         {label, category, amount, frequency, …} · PATCH {id, …}
  GET        /api/finance/cash-view/export/                     xlsx
"""
import datetime as dt
from decimal import Decimal, InvalidOperation

from django.db import transaction
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from . import cash_view as CV
from .receivables_views import _as_of, _j
from .recon_views import FINANCE, FINANCE_EDIT, FINANCE_EXPORT

SOURCE_LABELS = {'softech_mirror': 'SOFTECH', 'finance_input': 'إدخال المالية', 'estimate': 'تقدير'}
_FIXED_FIELDS = ('label', 'category', 'amount', 'frequency', 'day_of_month', 'date', 'start', 'end', 'is_active')


def _budget(request):
    raw = request.query_params.get('purchase_budget_weekly')
    if raw in (None, ''):
        return None
    try:
        v = Decimal(raw)
    except InvalidOperation:
        return None
    return v if v >= 0 else None


def _staff(request):
    return getattr(request.user, 'staff_profile', None)


def _exact(d):
    """Audit snapshot: money as exact strings, dates ISO."""
    return {k: (str(v) if isinstance(v, Decimal) else v.isoformat() if isinstance(v, (dt.date, dt.datetime)) else v)
            for k, v in d.items()}


def _audit(staff, what, before, after):
    from .models import ReconAuditEvent
    ReconAuditEvent.objects.create(action='cash_input_changed', before_state=_exact(before), after_state=_exact(after),
                                   detail=what, performed_by=staff)


def _can_edit(request):
    return FINANCE_EDIT().has_permission(request, None)


def _denied():
    return Response({'detail': 'تعديل مدخلات النقدية يحتاج صلاحية تعديل المالية.'}, status=status.HTTP_403_FORBIDDEN)


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE])
def cash_view(request):
    v = CV.build_cash_view(_as_of(request), purchase_budget_weekly=_budget(request))
    for ln in v['lines']:
        ln['source_label'] = SOURCE_LABELS[ln['source']]
    return Response(_j(v))


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE])
def cash_view_line(request, key):
    try:
        week = int(request.query_params.get('week', 0))
    except ValueError:
        week = 0
    rows = CV.line_rows(_as_of(request), key, week, purchase_budget_weekly=_budget(request))
    return Response(_j({'key': key, 'week': week, 'rows': rows}))


def _opening_dict(o):
    return {'id': o.pk, 'as_of': o.as_of, 'amount': o.amount, 'note': o.note,
            'entered_by': o.entered_by.full_name if o.entered_by_id else None, 'entered_at': o.entered_at}


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated, FINANCE])
def cash_opening_balance(request):
    from .models import CashOpeningBalance
    if request.method == 'GET':
        return Response(_j({'rows': [_opening_dict(o) for o in
                                     CashOpeningBalance.objects.select_related('entered_by')[:26]]}))
    if not _can_edit(request):
        return _denied()
    d = request.data
    try:
        day = dt.date.fromisoformat(str(d.get('as_of')))
        amount = Decimal(str(d.get('amount'))).quantize(Decimal('0.01'))
    except (ValueError, InvalidOperation):
        return Response({'detail': 'التاريخ والرصيد مطلوبان.'}, status=status.HTTP_400_BAD_REQUEST)
    with transaction.atomic():
        o = CashOpeningBalance.objects.create(as_of=day, amount=amount, note=str(d.get('note') or '').strip()[:300],
                                              entered_by=_staff(request))
        _audit(_staff(request), f'opening balance {day}', {}, _opening_dict(o))
    return Response(_j(_opening_dict(o)), status=status.HTTP_201_CREATED)


def _fixed_dict(f):
    return {'id': f.pk, 'label': f.label, 'category': f.category, 'category_label': f.get_category_display(),
            'amount': f.amount, 'frequency': f.frequency, 'frequency_label': f.get_frequency_display(),
            'day_of_month': f.day_of_month, 'date': f.date, 'start': f.start, 'end': f.end, 'is_active': f.is_active,
            'entered_by': f.entered_by.full_name if f.entered_by_id else None}


def _clean_fixed(d, partial):
    from .models import FixedOutflow
    out = {}
    if 'label' in d or not partial:
        out['label'] = str(d.get('label') or '').strip()[:200]
        if not out['label']:
            raise ValueError('البيان مطلوب.')
    for f, choices, msg in (('category', FixedOutflow.CATEGORY_CHOICES, 'فئة غير معروفة.'),
                            ('frequency', FixedOutflow.FREQUENCY_CHOICES, 'التكرار: أسبوعي أو شهري أو مرة واحدة.')):
        if f in d or not partial:
            if d.get(f) not in dict(choices):
                raise ValueError(msg)
            out[f] = d[f]
    if 'amount' in d or not partial:
        try:
            out['amount'] = Decimal(str(d.get('amount'))).quantize(Decimal('0.01'))
        except InvalidOperation:
            raise ValueError('المبلغ رقم.')
        if out['amount'] <= 0:
            raise ValueError('المبلغ أكبر من صفر.')
    if 'day_of_month' in d:
        v = d['day_of_month']
        if v in (None, ''):
            out['day_of_month'] = None
        else:
            try:
                v = int(v)
            except (TypeError, ValueError):
                v = 0
            if not 1 <= v <= 31:
                raise ValueError('يوم الشهر بين 1 و 31.')
            out['day_of_month'] = v
    for f in ('date', 'start', 'end'):
        if f in d:
            if d[f] in (None, ''):
                out[f] = None
            else:
                try:
                    out[f] = dt.date.fromisoformat(str(d[f]))
                except ValueError:
                    raise ValueError('تاريخ غير صحيح.')
    if 'is_active' in d:
        out['is_active'] = d['is_active'] if isinstance(d['is_active'], bool) \
            else str(d['is_active']).lower() in ('true', '1')
    return out


@api_view(['GET', 'POST', 'PATCH'])
@permission_classes([IsAuthenticated, FINANCE])
def cash_fixed_outflows(request):
    from .models import FixedOutflow
    if request.method == 'GET':
        return Response(_j({'rows': [_fixed_dict(f) for f in FixedOutflow.objects.select_related('entered_by')],
                            'categories': FixedOutflow.CATEGORY_CHOICES,
                            'frequencies': FixedOutflow.FREQUENCY_CHOICES}))
    if not _can_edit(request):
        return _denied()
    partial = request.method == 'PATCH'
    try:
        vals = _clean_fixed(request.data, partial)
    except ValueError as e:
        return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
    if partial:
        f = FixedOutflow.objects.filter(pk=request.data.get('id')).first()
        if f is None:
            return Response({'detail': 'البند غير موجود.'}, status=status.HTTP_404_NOT_FOUND)
        before = {k: getattr(f, k) for k in _FIXED_FIELDS}
    else:
        f, before = FixedOutflow(), {}
    for k, v in vals.items():
        setattr(f, k, v)
    if f.frequency == 'one_off' and not f.date:
        return Response({'detail': 'حدد تاريخ المرة الواحدة.'}, status=status.HTTP_400_BAD_REQUEST)
    if f.frequency == 'monthly' and not f.day_of_month:
        return Response({'detail': 'حدد يوم الشهر.'}, status=status.HTTP_400_BAD_REQUEST)
    f.entered_by = _staff(request)
    with transaction.atomic():
        f.save()
        _audit(_staff(request), f'fixed outflow {f.pk} {f.label}', before, {k: getattr(f, k) for k in _FIXED_FIELDS})
    return Response(_j(_fixed_dict(f)), status=status.HTTP_200_OK if partial else status.HTTP_201_CREATED)


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE_EXPORT])
def cash_view_export(request):
    from io import BytesIO
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Font, PatternFill
    as_of = _as_of(request)
    v = CV.build_cash_view(as_of, purchase_budget_weekly=_budget(request))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'النقدية 13 أسبوع'
    ws.sheet_view.rightToLeft = True
    ws.append([f'النقدية 13 أسبوع — حتى {as_of.isoformat()}'
               + (f' — رصيد افتتاحي {v["opening"]["amount"]:,.2f} في {v["opening"]["as_of"]}' if v['opening']
                  else ' — بدون رصيد افتتاحي')])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    head = ['البند', 'المصدر'] + [f'{w["start"]:%d/%m}' for w in v['weeks']]
    ws.append(head)
    fill = PatternFill('solid', fgColor='022871')
    for i in range(1, len(head) + 1):
        c = ws.cell(row=3, column=i)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = fill
        c.alignment = Alignment(horizontal='center')
        ws.column_dimensions[c.column_letter].width = 26 if i == 1 else (14 if i == 2 else 12)
    groups = {'in': 'داخل', 'out': 'خارج', 'memo': 'للعلم (خارج الإجمالي)'}
    for grp, title in groups.items():
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for ln in v['lines']:
            if ln['group'] == grp:
                ws.append([ln['label_ar'], SOURCE_LABELS[ln['source']]] + [float(x) for x in ln['base']])
    ws.append([])
    ws.append(['صافي الأسبوع (أساسي)', ''] + [float(x) for x in v['net']['base']])
    for case, label in (('low', 'الرصيد — متشائم'), ('base', 'الرصيد — أساسي'), ('high', 'الرصيد — متفائل')):
        ws.append([label, ''] + ([float(x) for x in v['closing'][case]] if v['closing'][case] else []))
    ws.freeze_panes = 'C4'
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="cash_view_{as_of.isoformat()}.xlsx"'
    return resp
