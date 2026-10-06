"""
apps/finance/recon_reports.py

Finance-facing views of the reconciliation (doc 23 §14), PostgreSQL only:

  unpaid_invoices()   every purchase invoice still carrying an open balance once the
                      reconstruction is applied, split into what is already covered
                      by an approved match, what is under review (medium/low
                      proposals) and what NO voucher covers — the amount really
                      owed to the supplier, which is what finance must act on.
  build_unpaid_xlsx() the same as an Excel workbook (supplier summary + invoice rows)
  build_review_xlsx() the proposals awaiting human review, grouped by voucher, so a
                      reviewer can check each split before approving or rejecting it.
"""
from __future__ import annotations

import io
from collections import defaultdict
from decimal import Decimal

from django.db import models
from django.db.models import Sum

from . import recon_labels as L
from . import recon_filters as RF

ZERO = Decimal('0')
TOL = Decimal('0.01')
# SOFTECH's own piaster rounding (e.g. 711.56 recorded paid 711.50) is not a debt:
# balances under 1 EGP are treated as settled in the finance reports.
ROUNDING = Decimal('1.00')


def _local(dt):
    """SOFTECH entry timestamp (trans_time) as naive Cairo local time — for Excel cells
    and report rows, so the page shows the time the user saw in SOFTECH."""
    if dt is None:
        return None
    from django.utils import timezone
    return timezone.localtime(dt).replace(tzinfo=None) if timezone.is_aware(dt) else dt


def _hm(dt):
    t = _local(dt)
    return t.strftime('%Y-%m-%d %H:%M') if t else ''


def unpaid_invoices(party_type='supplier', personcode=None, min_outstanding=ROUNDING, f=None):
    from .models import APInvoice, Allocation, MatchCandidate
    inv_qs = APInvoice.objects.filter(party__party_type=party_type, is_return=False)
    if personcode:
        inv_qs = inv_qs.filter(party__softech_personcode=personcode)
    if f is not None:
        inv_qs = RF.apply(inv_qs, f, 'invoice')

    alloc = defaultdict(lambda: defaultdict(lambda: ZERO))
    for r in (Allocation.objects.filter(invoice__in=inv_qs)
              .values('invoice_id', 'origin').annotate(s=Sum('amount'))):
        alloc[r['invoice_id']][r['origin']] += r['s'] or ZERO
    prop = defaultdict(lambda: defaultdict(lambda: ZERO))
    for r in (MatchCandidate.objects.filter(invoice__in=inv_qs,
                                            status=MatchCandidate.STATUS_PROPOSED)
              .exclude(confidence_class=MatchCandidate.CONF_CONFLICT)
              .values('invoice_id', 'confidence_class').annotate(s=Sum('proposed_amount'))):
        prop[r['invoice_id']][r['confidence_class']] += r['s'] or ZERO

    from .recon_returns import open_returned_by_purchase
    returned_open = open_returned_by_purchase(party_type=party_type)
    rows = []
    for inv in inv_qs.select_related('party').order_by('party__softech_personcode', 'docdate'):
        a = alloc[inv.id]
        paid_softech = max(inv.doc_value_pay or ZERO, a['softech'] + a['written'])
        returned = min(returned_open.get(inv.id, ZERO), max((inv.doc_value or ZERO) - paid_softech, ZERO))
        # a returned part is not owed to the supplier — only what was kept
        open_softech = (inv.doc_value or ZERO) - paid_softech - returned
        if open_softech <= min_outstanding:
            continue
        approved = min(a['approved'], open_softech)                   # matched, awaiting write
        after_approved = open_softech - approved
        high = min(prop[inv.id]['high'], after_approved)
        review = min(prop[inv.id]['medium'] + prop[inv.id]['low'], after_approved - high)
        uncovered = after_approved - high - review                    # no voucher covers it
        rows.append({
            'invoice_id': inv.id,
            'personcode': inv.party.softech_personcode, 'supplier': inv.party.name,
            'branchcode': inv.branchcode, 'branch_name': L.branch_name(inv.branchcode),
            'docnumber': str(inv.docnumber).split('.')[0],
            'supplier_docnumber': inv.docnumber2 or '', 'docdate': inv.docdate,
            'entered_at': _local(inv.trans_time),
            'usercode': inv.usercode, 'user_name': L.user_name(inv.usercode),
            'comments': inv.comments or '',
            'doc_value': inv.doc_value, 'paid_in_softech': paid_softech, 'returned_open': returned,
            'open_in_softech': open_softech, 'matched_approved': approved,
            'remaining_calc': APInvoice.remaining_from(inv.doc_value, inv.doc_value_pay,
                                                       a['softech'] + a['written'], a['approved']),
            'returned_open_linked': returned,
            'matched_high_pending': high, 'under_review': review, 'unpaid_uncovered': uncovered,
            'status': ('مغطاة بمطابقة' if uncovered <= TOL and review <= TOL
                       else 'تحت المراجعة' if uncovered <= TOL
                       else 'مسددة جزئياً' if paid_softech + approved + high + review > TOL
                       else 'غير مسددة'),
        })
    return rows


def supplier_summary(rows, party_type='supplier', f=None):
    from .models import ReconParty
    agg = {}
    for r in rows:
        s = agg.setdefault(r['personcode'], {
            'personcode': r['personcode'], 'supplier': r['supplier'], 'invoices': 0,
            'open_in_softech': ZERO, 'matched': ZERO, 'under_review': ZERO,
            'unpaid_uncovered': ZERO, 'unpaid_invoices': 0})
        s['invoices'] += 1
        s['open_in_softech'] += r['open_in_softech']
        s['matched'] += r['matched_approved'] + r['matched_high_pending']
        s['under_review'] += r['under_review']
        s['unpaid_uncovered'] += r['unpaid_uncovered']
        if r['unpaid_uncovered'] > TOL:
            s['unpaid_invoices'] += 1
    from .models import APInvoice
    from .recon_returns import open_return_amounts
    rets_qs = APInvoice.objects.filter(party__party_type=party_type, is_return=True).exclude(source_hash='')
    if f is not None:
        rets_qs = RF.apply(rets_qs, f, 'invoice')
    rets = list(rets_qs.select_related('party'))
    credit = open_return_amounts(rets)
    by_pc = defaultdict(lambda: ZERO)
    for r in rets:
        if r.id in credit:
            by_pc[r.party.softech_personcode] += credit[r.id]
            agg.setdefault(r.party.softech_personcode, {
                'personcode': r.party.softech_personcode, 'supplier': r.party.name, 'invoices': 0,
                'open_in_softech': ZERO, 'matched': ZERO, 'under_review': ZERO,
                'unpaid_uncovered': ZERO, 'unpaid_invoices': 0})
    bal = dict(ReconParty.objects.filter(party_type=party_type,
                                         softech_personcode__in=list(agg))
               .values_list('softech_personcode', 'softech_balance'))
    for pc, s in agg.items():
        s['softech_balance'] = bal.get(pc)
        s['open_returns'] = by_pc.get(pc, ZERO)            # the supplier owes us this
        s['net_payable'] = s['unpaid_uncovered'] - s['open_returns']
    return sorted(agg.values(), key=lambda s: s['unpaid_uncovered'], reverse=True)


# ── Excel ────────────────────────────────────────────────────────────────────

def _sheet(ws, headers, rows, money_cols=()):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='022871')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    import datetime as _dt
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, _dt.datetime):
                c.number_format = 'yyyy-mm-dd hh:mm:ss'
            elif isinstance(c.value, _dt.date):
                c.number_format = 'yyyy-mm-dd'
    for idx in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for c in row:
                c.number_format = '#,##0.00'
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(str(h)) + 6))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def build_unpaid_xlsx(party_type='supplier', personcode=None, f=None) -> bytes:
    from openpyxl import Workbook
    rows = unpaid_invoices(party_type, personcode, f=f)
    summary = supplier_summary(rows, party_type, f=f)
    wb = Workbook()
    ws = wb.active
    ws.title = 'ملخص الموردين'
    _sheet(ws, ['كود المورد', 'المورد', 'رصيد SOFTECH', 'عدد الفواتير المفتوحة',
                'المفتوح في SOFTECH', 'مغطى بمطابقة', 'تحت المراجعة',
                'غير مسدد فعلياً (بلا سند)', 'عدد الفواتير غير المسددة',
                'مرتجعات مستحقة على المورد', 'صافي المستحق للمورد'],
           [[s['personcode'], s['supplier'], s['softech_balance'], s['invoices'],
             s['open_in_softech'], s['matched'], s['under_review'], s['unpaid_uncovered'],
             s['unpaid_invoices'], s['open_returns'], s['net_payable']] for s in summary],
           money_cols=(3, 5, 6, 7, 8, 10, 11))
    ws2 = wb.create_sheet('الفواتير غير المسددة')
    _sheet(ws2, ['كود المورد', 'المورد', 'كود الفرع', 'اسم الفرع', 'رقم الفاتورة', 'مستند المورد',
                 'تاريخ الفاتورة', 'وقت إدخال الفاتورة', 'كود المستخدم', 'اسم المستخدم',
                 'قيمة الفاتورة', 'المسدد في SOFTECH', 'مرتجع لم يُخصم', 'المفتوح في SOFTECH',
                 'مطابقة معتمدة', 'المتبقي (محسوب)', 'الصافي بعد المرتجع (محسوب)', 'مطابقة عالية بانتظار الاعتماد', 'تحت المراجعة',
                 'غير مسدد فعلياً', 'الحالة', 'ملاحظات الفاتورة'],
           [[r['personcode'], r['supplier'], r['branchcode'], r['branch_name'], r['docnumber'],
             r['supplier_docnumber'], r['docdate'], r['entered_at'], r['usercode'], r['user_name'],
             r['doc_value'], r['paid_in_softech'], r['returned_open'],
             r['open_in_softech'], r['matched_approved'], r['remaining_calc'], r['remaining_calc'] - r['returned_open'], r['matched_high_pending'],
             r['under_review'], r['unpaid_uncovered'], r['status'], r['comments']]
            for r in rows if r['unpaid_uncovered'] > TOL or r['under_review'] > TOL],
           money_cols=(11, 12, 13, 14, 15, 16, 17, 18, 19, 20))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_review_xlsx(party_type='supplier', personcode=None, classes=('medium', 'low'), f=None) -> bytes:
    from openpyxl import Workbook
    from .models import MatchCandidate
    qs = (MatchCandidate.objects.filter(status=MatchCandidate.STATUS_PROPOSED,
                                        party__party_type=party_type,
                                        confidence_class__in=classes)
          .select_related('party', 'invoice', 'payment')
          .order_by('party__softech_personcode', 'group_key', 'invoice__docdate'))
    if personcode:
        qs = qs.filter(party__softech_personcode=personcode)
    if f is not None:
        qs = RF.apply(qs, f, 'candidate')
    wb = Workbook()
    ws = wb.active
    ws.title = 'مطابقات للمراجعة'
    _sheet(ws, ['رقم المقترح', 'كود المورد', 'المورد', 'المجموعة', 'الاستراتيجية', 'الثقة %',
                'سبب المراجعة',
                'فرع السند', 'مسلسل السند', 'المسلسل الداخلي', 'رقم الإيصال', 'تاريخ السند',
                'وقت إدخال السند', 'مستخدم السند', 'مبلغ السند', 'ملاحظات السند',
                'فرع الفاتورة', 'رقم الفاتورة', 'تاريخ الفاتورة', 'وقت إدخال الفاتورة',
                'مستخدم الفاتورة', 'قيمة الفاتورة',
                'ملاحظات الفاتورة', 'المبلغ المقترح'],
           [[c.id, c.party.softech_personcode, c.party.name, c.group_key, c.get_strategy_display(),
             c.confidence_score, (c.decision_note or '').replace('يحتاج مراجعة: ', ''),
             L.branch_label(c.payment.branchcode), c.payment.cheqsno, c.payment.cheqno,
             c.payment.ourcheqsno, c.payment.voucher_date, _local(c.payment.trans_time),
             L.user_label(c.payment.usercode), c.payment.amount, c.payment.note,
             L.branch_label(c.invoice.branchcode), str(c.invoice.docnumber).split('.')[0],
             c.invoice.docdate, _local(c.invoice.trans_time), L.user_label(c.invoice.usercode),
             c.invoice.doc_value, c.invoice.comments, c.proposed_amount]
            for c in qs.iterator()], money_cols=(15, 22, 24))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── incomplete payments: invoices started but not finished + part-used vouchers ─

def partial_payments(party_type='supplier', personcode=None, f=None):
    """
    (1) invoices where some payment is applied (in SOFTECH, written by us, or approved
        and awaiting the write) but a balance remains — with every voucher that paid
        it and any match still under review for the remainder;
    (2) vouchers whose money is only partly allocated to invoices.
    """
    import datetime
    from .models import APInvoice, Allocation, MatchCandidate, Payment
    inv_qs = APInvoice.objects.filter(party__party_type=party_type, is_return=False)
    pay_qs = Payment.objects.filter(party__party_type=party_type, direction='out')
    if personcode:
        inv_qs = inv_qs.filter(party__softech_personcode=personcode)
        pay_qs = pay_qs.filter(party__softech_personcode=personcode)
    if f is not None:
        inv_qs = RF.apply(inv_qs, f, 'invoice')
        pay_qs = RF.apply(pay_qs, f, 'payment')

    allocs_by_inv = defaultdict(list)
    used_by_pay = defaultdict(lambda: ZERO)
    invs_by_pay = defaultdict(list)
    for a in Allocation.objects.filter(payment__in=pay_qs).select_related('payment', 'invoice'):
        allocs_by_inv[a.invoice_id].append(a)
        used_by_pay[a.payment_id] += a.payment.allocation_effect(a.amount, a.invoice.is_return)
        invs_by_pay[a.payment_id].append(a)
    review = defaultdict(lambda: ZERO)
    for r in (MatchCandidate.objects.filter(invoice__in=inv_qs, status=MatchCandidate.STATUS_PROPOSED)
              .exclude(confidence_class=MatchCandidate.CONF_CONFLICT)
              .values('invoice_id').annotate(s=Sum('proposed_amount'))):
        review[r['invoice_id']] = r['s'] or ZERO

    today = datetime.date.today()
    inv_rows = []
    for inv in inv_qs.select_related('party').order_by('party__softech_personcode', 'docdate'):
        al = allocs_by_inv.get(inv.id, [])
        via_vouchers = sum((a.amount for a in al), ZERO)
        settled = sum((a.amount for a in al if a.origin != Allocation.ORIGIN_APPROVED), ZERO)
        pending = sum((a.amount for a in al if a.origin == Allocation.ORIGIN_APPROVED), ZERO)
        paid = max(inv.doc_value_pay or ZERO, settled) + pending
        remaining = (inv.doc_value or ZERO) - paid
        if paid <= TOL or remaining < ROUNDING:
            continue
        al_sorted = sorted(al, key=lambda a: (a.payment.voucher_date or datetime.date.min,
                                              _local(a.payment.trans_time) or datetime.datetime.min))
        inv_rows.append({
            'personcode': inv.party.softech_personcode, 'supplier': inv.party.name,
            'branchcode': inv.branchcode, 'branch_name': L.branch_name(inv.branchcode),
            'docnumber': str(inv.docnumber).split('.')[0],
            'supplier_docnumber': inv.docnumber2 or '', 'docdate': inv.docdate,
            'entered_at': _local(inv.trans_time),
            'usercode': inv.usercode, 'user_name': L.user_name(inv.usercode),
            'comments': inv.comments or '', 'invoice_id': inv.id,
            'doc_value': inv.doc_value, 'paid': paid, 'remaining': remaining,
            'paid_pct': (paid / inv.doc_value * 100).quantize(Decimal('0.1')) if inv.doc_value else ZERO,
            'vouchers_count': len({a.payment_id for a in al}),
            'vouchers': ' | '.join(
                f'{a.payment.branchcode}/{a.payment.cheqsno} {_hm(a.payment.trans_time) or a.payment.voucher_date} = {a.amount}'
                f' [{L.user_label(a.payment.usercode)}]'
                + (' (بانتظار الكتابة)' if a.origin == Allocation.ORIGIN_APPROVED else '')
                for a in al_sorted),
            'paid_outside_vouchers': max(ZERO, (inv.doc_value_pay or ZERO) - via_vouchers),
            'last_payment': al_sorted[-1].payment.voucher_date if al_sorted else None,
            'days_open': (today - inv.docdate).days if inv.docdate else None,
            'remaining_under_review': min(review.get(inv.id, ZERO), remaining),
        })

    pay_rows = []
    for p in pay_qs.select_related('party').order_by('party__softech_personcode', 'voucher_date'):
        used = used_by_pay.get(p.id, ZERO)
        left = (p.net_amount or ZERO) - used           # net of refunds bound to it
        if used <= TOL or left < ROUNDING:
            continue
        pay_rows.append({
            'personcode': p.party.softech_personcode, 'supplier': p.party.name,
            'voucher': f'{p.branchcode}/{p.cheqsno}', 'voucher_date': p.voucher_date,
            'entered_at': _local(p.trans_time),
            'branchcode': p.branchcode, 'branch_name': L.branch_name(p.branchcode),
            'cheqsno': p.cheqsno, 'cheqno': p.cheqno or '', 'ourcheqsno': p.ourcheqsno,
            'usercode': p.usercode, 'user_name': L.user_name(p.usercode),
            'amount': p.amount, 'allocated': used, 'unallocated': left, 'note': p.note or '',
            'invoices': ' | '.join(
                f'{a.invoice.branchcode}/{str(a.invoice.docnumber).split(".")[0]} = {a.amount}'
                for a in invs_by_pay[p.id]),
        })
    return inv_rows, pay_rows


def build_partial_xlsx(party_type='supplier', personcode=None, f=None) -> bytes:
    from openpyxl import Workbook
    inv_rows, pay_rows = partial_payments(party_type, personcode, f=f)
    wb = Workbook()
    ws = wb.active
    ws.title = 'فواتير مسددة جزئياً'
    _sheet(ws, ['كود المورد', 'المورد', 'كود الفرع', 'اسم الفرع', 'رقم الفاتورة', 'مستند المورد',
                'تاريخ الفاتورة', 'وقت إدخال الفاتورة', 'كود المستخدم', 'اسم المستخدم',
                'قيمة الفاتورة', 'المسدد', 'المتبقي (محسوب)', '% المسدد', 'عدد السندات',
                'السندات (فرع/مسلسل وقت الإدخال = مبلغ [المستخدم])', 'مسدد خارج السندات', 'آخر سداد',
                'أيام منذ الفاتورة', 'من المتبقي تحت المراجعة', 'ملاحظات الفاتورة'],
           [[r['personcode'], r['supplier'], r['branchcode'], r['branch_name'], r['docnumber'],
             r['supplier_docnumber'], r['docdate'], r['entered_at'], r['usercode'], r['user_name'],
             r['doc_value'], r['paid'], r['remaining'], r['paid_pct'], r['vouchers_count'],
             r['vouchers'], r['paid_outside_vouchers'], r['last_payment'], r['days_open'],
             r['remaining_under_review'], r['comments']] for r in inv_rows],
           money_cols=(11, 12, 13, 17, 20))
    ws.column_dimensions['P'].width = 80
    ws2 = wb.create_sheet('سندات لم تُستكمل')
    _sheet(ws2, ['كود المورد', 'المورد', 'كود الفرع', 'اسم الفرع', 'مسلسل السند', 'المسلسل الداخلي',
                 'رقم الإيصال', 'تاريخ السند', 'وقت إدخال السند', 'كود المستخدم', 'اسم المستخدم', 'مبلغ السند',
                 'المخصص لفواتير', 'المتبقي بلا فاتورة', 'ملاحظات السند', 'الفواتير المسددة به'],
           [[r['personcode'], r['supplier'], r['branchcode'], r['branch_name'], r['cheqsno'],
             r['cheqno'], r['ourcheqsno'], r['voucher_date'], r['entered_at'], r['usercode'], r['user_name'], r['amount'],
             r['allocated'], r['unallocated'], r['note'], r['invoices']] for r in pay_rows],
           money_cols=(12, 13, 14))
    ws2.column_dimensions['P'].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()



def build_returns_chains_xlsx(party_type='supplier', personcode=None, f=None) -> bytes:
    """Open returns the supplier owes us · purchases paid although returned ·
    مدفوعات → مقبوضات → مدفوعات correction chains — each row with supplier,
    branch, date and the SOFTECH user."""
    from openpyxl import Workbook
    from .models import APInvoice, Payment, ReconException, Allocation
    from .recon_returns import open_return_amounts
    wb = Workbook()

    ws = wb.active
    ws.title = 'مرتجعات مستحقة على المورد'
    rets = APInvoice.objects.filter(party__party_type=party_type, is_return=True).exclude(source_hash='')
    if personcode:
        rets = rets.filter(party__softech_personcode=personcode)
    if f is not None:
        rets = RF.apply(rets, f, 'invoice')
    rets = list(rets.select_related('party').order_by('party__softech_personcode', 'docdate'))
    credit = open_return_amounts(rets)
    rows = []
    for r in rets:
        if r.id not in credit:
            continue
        links = ', '.join(f'{l.purchase_branchcode}/{l.purchase_docnumber} ({l.amount})'
                          for l in r.return_links.all())
        rows.append([r.party.softech_personcode, r.party.name, r.branchcode, L.branch_name(r.branchcode),
                     str(r.docnumber).split('.')[0], r.docdate, _local(r.trans_time), r.usercode, L.user_name(r.usercode),
                     r.doc_value, r.doc_value - credit[r.id], credit[r.id], links or '— (مرتجع عام)',
                     r.comments or ''])
    _sheet(ws, ['كود المورد', 'المورد', 'كود الفرع', 'اسم الفرع', 'رقم المرتجع', 'تاريخ المرتجع',
                'وقت إدخال المرتجع', 'كود المستخدم', 'اسم المستخدم', 'قيمة المرتجع', 'ما خُصم/استُرد', 'المتبقي المستحق على المورد',
                'فاتورة الشراء المرتجعة', 'ملاحظات المرتجع'], rows, money_cols=(10, 11, 12))

    ws2 = wb.create_sheet('مسددة رغم إرجاعها')
    ex = ReconException.objects.filter(exception_type=ReconException.TYPE_PAID_RETURNED, status='open',
                                       party__party_type=party_type).select_related('party', 'invoice')
    if personcode:
        ex = ex.filter(party__softech_personcode=personcode)
    if f is not None:
        ex = RF.apply(ex, f, 'exception')
    _sheet(ws2, ['كود المورد', 'المورد', 'كود الفرع', 'اسم الفرع', 'رقم الفاتورة', 'تاريخ الفاتورة',
                 'وقت إدخال الفاتورة', 'مستخدم الفاتورة', 'قيمة الفاتورة', 'التفاصيل'],
           [[e.party.softech_personcode, e.party.name, e.invoice.branchcode, L.branch_name(e.invoice.branchcode),
             str(e.invoice.docnumber).split('.')[0], e.invoice.docdate, _local(e.invoice.trans_time),
             L.user_label(e.invoice.usercode), e.invoice.doc_value, e.detail] for e in ex if e.invoice],
           money_cols=(9,))
    ws2.column_dimensions['J'].width = 90

    ws3 = wb.create_sheet('سلاسل الإلغاء والتصحيح')
    chain = Payment.objects.filter(party_type=party_type).exclude(chain_key='')
    if personcode:
        chain = chain.filter(party__softech_personcode=personcode)
    if f is not None:
        chain = RF.apply(chain, f, 'payment')
    groups = {}
    for p in chain.select_related('party').order_by('chain_key', 'voucher_date', 'trans_time', 'cheqsno'):
        groups.setdefault(p.chain_key, []).append(p)
    linked = {}
    for a in Allocation.objects.filter(payment__in=[p for g in groups.values() for p in g]).select_related('invoice'):
        linked.setdefault(a.payment_id, []).append(f'{a.invoice.docnumber} [{a.get_origin_display()}]')
    rows = []
    for key, ps in groups.items():
        for p in ps:
            rows.append([key, p.get_chain_role_display(), p.party.softech_personcode, p.party.name,
                         p.branchcode, L.branch_name(p.branchcode), p.cheqsno, p.cheqno or '', p.ourcheqsno,
                         'مقبوضات' if p.is_receipt else 'مدفوعات', p.voucher_date, _local(p.trans_time),
                         L.user_label(p.usercode), p.amount, p.note, p.chain_note,
                         ', '.join(linked.get(p.id, [])) or '—'])
    _sheet(ws3, ['رقم السلسلة', 'الدور', 'كود المورد', 'المورد', 'كود الفرع', 'اسم الفرع', 'مسلسل السند',
                 'المسلسل الداخلي', 'رقم الإيصال', 'النوع', 'تاريخ السند', 'وقت إدخال السند', 'المستخدم',
                 'المبلغ', 'ملاحظات السند', 'التفسير', 'مربوط بفواتير'],
           rows, money_cols=(14,))
    ws3.column_dimensions['P'].width = 60

    # bindings (recon_bindings): a مقبوضات refunding part of a مدفوعات (payment matched NET),
    # or naming a purchase invoice (owner to decide: money back or credit)
    ws4 = wb.create_sheet('ربط المقبوضات')
    bound = Payment.objects.filter(party_type=party_type, cheqtype='10').filter(
        models.Q(bound_payment__isnull=False) | models.Q(bound_invoice__isnull=False))
    if personcode:
        bound = bound.filter(party__softech_personcode=personcode)
    if f is not None:
        bound = RF.apply(bound, f, 'payment')
    rows = []
    for r in bound.select_related('party', 'bound_payment', 'bound_invoice').order_by('party__softech_personcode', 'voucher_date'):
        if r.bound_payment_id:
            t = r.bound_payment
            rows.append([r.party.softech_personcode, r.party.name, f'{r.branchcode}/{r.cheqsno}', r.voucher_date,
                         _local(r.trans_time), L.user_label(r.usercode), r.amount, r.note,
                         'يسترد جزءًا من سند صرف', f'{t.branchcode}/{t.cheqsno}', t.voucher_date, t.amount,
                         t.net_amount, 'لا — يُطابَق صافي السند تلقائيًا'])
        else:
            t = r.bound_invoice
            rows.append([r.party.softech_personcode, r.party.name, f'{r.branchcode}/{r.cheqsno}', r.voucher_date,
                         _local(r.trans_time), L.user_label(r.usercode), r.amount, r.note,
                         'يشير لفاتورة شراء', f'{t.branchcode}/{str(t.docnumber).split(".")[0]}', t.docdate,
                         t.doc_value, t.doc_value_pay, 'نعم — استرداد (مستحقة مجددًا) أم خصم (يقلل المستحق)؟'])
    _sheet(ws4, ['كود المورد', 'المورد', 'سند المقبوضات', 'تاريخه', 'وقت الإدخال', 'المستخدم', 'المبلغ',
                 'ملاحظات', 'نوع الربط', 'المستند المرتبط', 'تاريخه', 'قيمته',
                 'صافي السند / المسدد من الفاتورة', 'يحتاج قرار؟'], rows, money_cols=(7, 12, 13))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
