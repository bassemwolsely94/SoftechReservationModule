"""
apps/finance/supplier_terms_views.py

P-42 «شروط الموردين» API (docs/product-review/06_SPEC_P-42_supplier_terms.md). Company-level finance: the finance
matrix grant AND an all-branch user (owner decision B-1 — recon_views._company_only). Read-only toward SOFTECH.

  GET   /api/finance/supplier-terms/                 terms + real suppliers by open value + terms coverage
  POST  /api/finance/supplier-terms/                 {personcode, category, credit_days, effective_from, …}
  PATCH /api/finance/supplier-terms/                 {id, …changed fields}
  GET   /api/finance/supplier-terms/calendar/?as_of=&weeks=
  GET   /api/finance/supplier-terms/on-time/
  GET   /api/finance/supplier-terms/early-pay/
  GET   /api/finance/supplier-terms/residue/
  GET   /api/finance/supplier-terms/export/?sheet=terms|calendar|early_pay|on_time|residue    xlsx
  POST  /api/finance/supplier-terms/recompute/       background, one at a time (GET → {running})
"""
import datetime as dt
import logging
import threading
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from django.db import IntegrityError, transaction
from django.utils import timezone
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from . import supplier_terms as ST
from .receivables_views import _as_of, _j
from .recon_views import FINANCE, FINANCE_EDIT, FINANCE_EXPORT

logger = logging.getLogger('elrezeiky.finance')

CATEGORY_LABELS = {'any': 'كل الأصناف', 'local': 'محلي', 'imported': 'مستورد', 'unclassified': 'غير مصنف'}
_FIELDS = ('category', 'credit_days', 'days_from', 'early_pay_days', 'early_pay_discount_pct', 'early_pay_bonus_note',
           'effective_from', 'effective_to', 'source_note')


def _term_dict(t):
    return {'id': t.pk, 'personcode': t.party.softech_personcode, 'supplier': t.party.name,
            'category': t.category, 'category_label': CATEGORY_LABELS[t.category],
            'credit_days': t.credit_days, 'days_from': t.days_from, 'early_pay_days': t.early_pay_days,
            'early_pay_discount_pct': t.early_pay_discount_pct, 'early_pay_bonus_note': t.early_pay_bonus_note,
            'effective_from': t.effective_from, 'effective_to': t.effective_to, 'source_note': t.source_note,
            'entered_by': t.entered_by.full_name if t.entered_by_id else None, 'updated_at': t.updated_at}


def _snapshot(t):
    return _j({f: getattr(t, f) for f in _FIELDS})


def _clean(d, partial):
    """Validated field values from the request body, or raise ValueError with an Arabic message."""
    from .models import SupplierTerm
    out = {}

    def days(name, required=False):
        if name not in d:
            if required:
                raise ValueError('أيام الائتمان مطلوبة.')
            return
        v = d[name]
        if v in (None, '') and not required:
            out[name] = None
            return
        try:
            v = int(v)
        except (TypeError, ValueError):
            raise ValueError('الأيام رقم صحيح.')
        if not 0 <= v <= 365:
            raise ValueError('الأيام بين 0 و 365.')
        out[name] = v

    def date(name, required=False):
        if name not in d:
            if required:
                raise ValueError('تاريخ السريان مطلوب.')
            return
        if d[name] in (None, '') and not required:
            out[name] = None
            return
        try:
            out[name] = dt.date.fromisoformat(str(d[name]))
        except ValueError:
            raise ValueError('تاريخ غير صحيح.')

    if 'category' in d or not partial:
        if d.get('category') not in dict(SupplierTerm.CATEGORY_CHOICES):
            raise ValueError('الفئة: كل الأصناف أو محلي أو مستورد.')
        out['category'] = d['category']
    if 'days_from' in d:
        if d['days_from'] not in dict(SupplierTerm.DAYS_FROM_CHOICES):
            raise ValueError('بداية الأيام: من تاريخ الفاتورة أو من نهاية الشهر.')
        out['days_from'] = d['days_from']
    days('credit_days', required=not partial)
    days('early_pay_days')
    date('effective_from', required=not partial)
    date('effective_to')
    if 'early_pay_discount_pct' in d:
        v = d['early_pay_discount_pct']
        if v in (None, ''):
            out['early_pay_discount_pct'] = None
        else:
            try:
                v = Decimal(str(v))
            except InvalidOperation:
                raise ValueError('نسبة الخصم رقم.')
            if not Decimal('0') < v < Decimal('100'):
                raise ValueError('نسبة الخصم أكبر من 0 وأقل من 100.')
            out['early_pay_discount_pct'] = v.quantize(Decimal('0.01'))
    for f in ('early_pay_bonus_note', 'source_note'):
        if f in d:
            out[f] = str(d[f] or '').strip()[:300 if f == 'source_note' else 1000]
    return out


@api_view(['GET', 'POST', 'PATCH'])
@permission_classes([IsAuthenticated, FINANCE])
def supplier_terms(request):
    from .models import ReconAuditEvent, ReconParty, SupplierTerm
    if request.method == 'GET':
        return _terms_overview()
    if not FINANCE_EDIT().has_permission(request, None):
        return Response({'detail': 'تعديل الشروط يحتاج صلاحية تعديل المالية.'}, status=status.HTTP_403_FORBIDDEN)
    d = request.data
    partial = request.method == 'PATCH'
    try:
        vals = _clean(d, partial)
    except ValueError as e:
        return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
    if partial:
        t = SupplierTerm.objects.select_related('party').filter(pk=d.get('id')).first()
        if t is None:
            return Response({'detail': 'الشرط غير موجود.'}, status=status.HTTP_404_NOT_FOUND)
        before = _snapshot(t)
    else:
        party = ReconParty.objects.filter(softech_personcode=str(d.get('personcode') or '').strip(),
                                          party_type='supplier').first()
        if party is None:
            return Response({'detail': 'المورد غير موجود في التسوية.'}, status=status.HTTP_404_NOT_FOUND)
        t, before = SupplierTerm(party=party), {}
    for f, v in vals.items():
        setattr(t, f, v)
    if t.effective_to and t.effective_to < t.effective_from:
        return Response({'detail': 'تاريخ الانتهاء قبل تاريخ السريان.'}, status=status.HTTP_400_BAD_REQUEST)
    staff = getattr(request.user, 'staff_profile', None)
    t.entered_by = staff
    try:
        with transaction.atomic():
            t.save()
            ReconAuditEvent.objects.create(action='supplier_terms_changed', party=t.party, before_state=before,
                                           after_state=_snapshot(t), rules_version=ST.RULES_VERSION,
                                           detail=f'supplier {t.party.softech_personcode} {t.category}',
                                           performed_by=staff)
    except IntegrityError:
        return Response({'detail': 'يوجد شرط لنفس المورد والفئة بنفس تاريخ السريان.'},
                        status=status.HTTP_400_BAD_REQUEST)
    try:
        ST.recompute(t.party.softech_personcode)
    except Exception:
        logger.exception('supplier terms recompute failed for %s', t.party.softech_personcode)
    return Response(_j(_term_dict(t)), status=status.HTTP_200_OK if partial else status.HTTP_201_CREATED)


def _terms_overview():
    from .models import SupplierTerm
    internal = ST.internal_codes()
    open_by = defaultdict(lambda: Decimal('0'))
    names = {}
    for s in ST._slices():
        pc = s.invoice.party.softech_personcode
        if pc in internal:
            continue
        open_by[pc] += s.open_amount
        names[pc] = s.invoice.party.name
    top = [pc for pc, _ in sorted(open_by.items(), key=lambda x: -x[1])[:10]]
    terms = list(SupplierTerm.objects.select_related('party', 'entered_by'))
    for t in terms:
        names.setdefault(t.party.softech_personcode, t.party.name)
    with_terms = {t.party.softech_personcode for t in terms}
    suppliers = [{'personcode': pc, 'supplier': names.get(pc, ''), 'open': open_by.get(pc, Decimal('0')),
                  'is_top10': pc in top, 'has_terms': pc in with_terms}
                 for pc in sorted(set(open_by) | with_terms, key=lambda pc: -open_by.get(pc, 0))]
    return Response(_j({'terms': [_term_dict(t) for t in terms], 'suppliers': suppliers,
                        'coverage': ST.terms_coverage(), 'internal_codes': sorted(internal),
                        'category_labels': CATEGORY_LABELS}))


def _weeks(request):
    try:
        return max(1, min(int(request.query_params.get('weeks', 13)), 26))
    except ValueError:
        return 13


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE])
def supplier_terms_calendar(request):
    return Response(_j(ST.calendar(_as_of(request), _weeks(request))))


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE])
def supplier_terms_on_time(request):
    as_of = _as_of(request)
    rows = ST.on_time(as_of)
    paid = sum((r['paid_value'] for r in rows), Decimal('0'))
    linked = sum((r['linked_value'] for r in rows), Decimal('0'))
    return Response(_j({'as_of': as_of, 'window_days': ST.ON_TIME_WINDOW_DAYS, 'rows': rows,
                        'paid_value': paid, 'linked_value': linked,
                        'coverage_pct': (linked * 100 / paid).quantize(Decimal('0.1')) if paid else None}))


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE])
def supplier_terms_early_pay(request):
    return Response(_j(ST.early_pay(_as_of(request))))


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE])
def supplier_terms_residue(request):
    r = ST.residue(_as_of(request))
    keep = ('invoice_id', 'personcode', 'supplier', 'branch_name', 'docnumber', 'supplier_docnumber', 'docdate',
            'doc_value', 'open_in_softech', 'matched_approved', 'matched_high_pending', 'under_review',
            'unpaid_uncovered', 'status', 'account_class', 'age_days', 'comments')
    return Response(_j({**r, 'rows': [{k: x.get(k) for k in keep} for x in r['rows']]}))


# ── export ───────────────────────────────────────────────────────────────────

def _sheet(as_of, sheet):
    """(title, [(header, width)], rows) for one export sheet."""
    from .recon_labels import branch_name
    f = lambda v: float(v) if isinstance(v, Decimal) else v          # noqa: E731
    if sheet == 'terms':
        from .models import SupplierTerm
        return ('شروط الموردين',
                [('الحساب', 10), ('المورد', 34), ('الفئة', 12), ('أيام الائتمان', 10), ('تبدأ', 16),
                 ('أيام الدفع المبكر', 10), ('خصم %', 8), ('بونص', 24), ('ساري من', 12), ('ساري حتى', 12),
                 ('مرجع', 24), ('أدخله', 18)],
                [[t.party.softech_personcode, t.party.name, CATEGORY_LABELS[t.category], t.credit_days,
                  t.get_days_from_display(), t.early_pay_days, f(t.early_pay_discount_pct), t.early_pay_bonus_note,
                  t.effective_from, t.effective_to, t.source_note, t.entered_by.full_name if t.entered_by_id else '']
                 for t in SupplierTerm.objects.select_related('party', 'entered_by')])
    if sheet == 'calendar':
        c = ST.calendar(as_of)
        cols = ([('الحساب', 10), ('المورد', 34), ('الفئة', 12), ('المتأخر', 13), ('بدون شروط', 13)]
                + [(f'أسبوع {w["start"]:%d/%m}', 12) for w in c['weeks']] + [('بعد ذلك', 12), ('المفتوح', 14)])
        rows = [[r['personcode'], r['supplier'], CATEGORY_LABELS[r['category']], f(r['overdue']), f(r['no_terms'])]
                + [f(v) for v in r['weeks']] + [f(r['later']), f(r['open'])] for r in c['rows']]
        rows += [[]] + [['', 'حسابات داخلية (منفصلة)']] + [
            [r['personcode'], r['supplier'], '', '', '', *([''] * len(c['weeks'])), '', f(r['open'])]
            for r in c['internal_rows']]
        return 'تقويم السداد', cols, rows
    if sheet == 'early_pay':
        e = ST.early_pay(as_of)
        return ('فرص الدفع المبكر',
                [('الحساب', 10), ('المورد', 30), ('الفرع', 16), ('الفاتورة', 12), ('التاريخ', 12), ('المفتوح', 14),
                 ('آخر موعد', 12), ('الاستحقاق', 12), ('أيام مكسوبة', 10), ('خصم %', 8), ('قيمة الخصم', 12),
                 ('العائد السنوي %', 12), ('بونص', 24)],
                [[r['personcode'], r['supplier'], branch_name(r['branchcode']), r['docnumber'], r['docdate'],
                  f(r['open_amount']), r['deadline'], r['due_date'], r['days_gained'], f(r['discount_pct']),
                  f(r['discount_value']), f(r['annualised_return_pct']), r['bonus_note']] for r in e['open']])
    if sheet == 'on_time':
        return ('السداد في الموعد',
                [('الحساب', 10), ('المورد', 34), ('أيام مقابل الشروط', 12), ('قيمة مقاسة', 14),
                 ('بدون شروط', 14), ('المدفوع', 14), ('المربوط', 14), ('تغطية الربط %', 12)],
                [[r['personcode'], r['supplier'], f(r['days_vs_terms']), f(r['value']), f(r['no_terms_value']),
                  f(r['paid_value']), f(r['linked_value']), f(r['coverage_pct'])] for r in ST.on_time(as_of)])
    if sheet == 'residue':
        r = ST.residue(as_of)
        return ('البواقي القديمة',
                [('النوع', 10), ('الحساب', 10), ('المورد', 30), ('الفرع', 16), ('الفاتورة', 12), ('التاريخ', 12),
                 ('العمر (يوم)', 10), ('القيمة', 14), ('المفتوح', 14), ('الحالة', 16)],
                [['داخلي' if x['account_class'] == 'internal' else 'مورد فعلي', x['personcode'], x['supplier'],
                  x['branch_name'], x['docnumber'], x['docdate'], x['age_days'], f(x['doc_value']),
                  f(x['open_in_softech']), x['status']] for x in r['rows']])
    return None


@api_view(['GET'])
@permission_classes([IsAuthenticated, FINANCE_EXPORT])
def supplier_terms_export(request):
    from io import BytesIO
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Alignment, Font, PatternFill
    as_of = _as_of(request)
    sheet = request.query_params.get('sheet', 'calendar')
    got = _sheet(as_of, sheet)
    if got is None:
        return Response({'detail': 'ورقة غير معروفة.'}, status=status.HTTP_400_BAD_REQUEST)
    title, cols, rows = got
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True
    ws.append([f'{title} — حتى {as_of.isoformat()}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([c[0] for c in cols])
    fill = PatternFill('solid', fgColor='022871')
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = c[1]
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'A4'
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="supplier_terms_{sheet}_{as_of.isoformat()}.xlsx"'
    return resp


_recompute_lock = threading.Lock()


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated, FINANCE_EDIT])
def supplier_terms_recompute(request):
    """Rebuild every open invoice's category slices and due dates in the background (PostgreSQL only).
    GET → {running} for the screen's progress bar."""
    if request.method == 'GET':
        return Response({'running': _recompute_lock.locked()})
    if _recompute_lock.locked():
        return Response({'detail': 'إعادة الحساب جارية بالفعل.', 'running': True}, status=status.HTTP_409_CONFLICT)

    def _work():
        from django.db import connections
        with _recompute_lock:
            try:
                n = ST.recompute()
                logger.info('supplier terms recompute: %s slices at %s', n, timezone.now())
            except Exception:
                logger.exception('supplier terms recompute failed')
            finally:
                connections.close_all()
    threading.Thread(target=_work, name='supplier-terms-recompute', daemon=True).start()
    return Response({'detail': 'بدأت إعادة حساب شرائح الفواتير ومواعيد الاستحقاق.', 'running': True},
                    status=status.HTTP_202_ACCEPTED)
