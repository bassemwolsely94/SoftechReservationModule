"""
apps/gamification/views.py — mounted at /api/gamification/

Everyone (staff JWT):
  GET  me/                     my level, XP, streak, ranks, badges, recent points, open items
  GET  leaderboard/            ?period=&scope=branch|network&role=&branch=
  GET  branches/               branch-vs-branch ranking (average per active player)
  GET  rules/ levels/ badges/  how points are earned (transparent to everyone)
Managers (RBAC module `gamification`):
  GET  reports/overview/       view    — executive report
  GET  reports/export/         export  — Excel
  PATCH rules/<key>/ levels/<n>/ badges/<key>/   edit
  POST adjust/                 edit    — manual recognition / deduction with a reason
  POST run/                    edit    — re-run the engine now
  GET  changes/                view    — audit of every edit above
"""
from datetime import date

from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from . import engine, services
from .models import Badge, GamificationChange, Level, PointRule

MODULE = 'gamification'
MAX_MANUAL_POINTS = 500


def _profile(request):
    return getattr(request.user, 'staff_profile', None)


def _can(profile, action):
    return bool(profile and (profile.role == 'admin' or profile.can_do(MODULE, action)))


def _deny():
    return Response({'detail': 'ليس لديك صلاحية — Not allowed'}, status=status.HTTP_403_FORBIDDEN)


def _parse_date(v):
    try:
        return date.fromisoformat(v) if v else None
    except ValueError:
        return None


def _range(request):
    period = request.query_params.get('period') or 'month'
    if period not in services.PERIODS:
        period = 'month'
    return services.period_range(period, date_from=_parse_date(request.query_params.get('from')),
                                 date_to=_parse_date(request.query_params.get('to'))), period


def _int(v):
    try:
        return int(v) if v not in (None, '') else None
    except (TypeError, ValueError):
        return None


def _log(profile, action, target, before=None, after=None, reason=''):
    GamificationChange.objects.create(actor=profile, action=action, target=target,
                                      before=before or {}, after=after or {}, reason=reason)


# ── everyone ──────────────────────────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def me(request):
    profile = _profile(request)
    if not profile:
        return Response({'detail': 'no staff profile'}, status=404)
    data = services.me(profile)
    data['can_manage'] = _can(profile, 'view')
    data['can_edit'] = _can(profile, 'edit')
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def leaderboard(request):
    """Visibility (agreed): full ranking inside your own branch; network = top 10 only.
    Managers with gamification/view see every branch and the full network list."""
    profile = _profile(request)
    if not profile:
        return Response({'detail': 'no staff profile'}, status=404)
    (start, end), period = _range(request)
    scope = request.query_params.get('scope') or 'branch'
    role = request.query_params.get('role') or None
    manager = _can(profile, 'view')
    if scope == 'network':
        rows = services.ranking(start, end, role=role,
                                limit=None if manager else services.NETWORK_TOP)
        branch_id = None
    else:
        branch_id = _int(request.query_params.get('branch')) if manager else None
        branch_id = branch_id or profile.branch_id
        if not branch_id:
            return Response({'period': period, 'from': start, 'to': end, 'scope': scope,
                             'rows': [], 'me': None})
        rows = services.ranking(start, end, branch_id=branch_id, role=role)
    mine = next((r for r in rows if r['staff_id'] == profile.id), None)
    if mine is None and scope == 'network':
        mine = {**services.rank_of(profile, start, end, role=role), 'staff_id': profile.id}
    return Response({'period': period, 'from': start, 'to': end, 'scope': scope,
                     'branch_id': branch_id, 'role': role, 'rows': rows, 'me': mine,
                     'limited': scope == 'network' and not manager})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def branches(request):
    (start, end), period = _range(request)
    return Response({'period': period, 'from': start, 'to': end,
                     'rows': services.branch_ranking(start, end)})


def _rule_json(r):
    return {'key': r.key, 'name_ar': r.name_ar, 'name_en': r.name_en, 'desc_ar': r.desc_ar,
            'desc_en': r.desc_en, 'category': r.category, 'points': r.points,
            'unit_value': float(r.unit_value) if r.unit_value is not None else None,
            'cash_multiplier': float(r.cash_multiplier), 'daily_cap': r.daily_cap,
            'grace_hours': r.grace_hours, 'roles': r.roles, 'is_active': r.is_active}


def _level_json(lv):
    return {'number': lv.number, 'min_xp': lv.min_xp, 'title_ar': lv.title_ar,
            'title_en': lv.title_en, 'icon': lv.icon, 'color': lv.color}


def _badge_json(b):
    return {'key': b.key, 'icon': b.icon, 'name_ar': b.name_ar, 'name_en': b.name_en,
            'desc_ar': b.desc_ar, 'desc_en': b.desc_en, 'criteria': b.criteria,
            'rule_key': b.rule_key, 'threshold': b.threshold, 'is_active': b.is_active}


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def rules(request):
    from .defaults import ensure_defaults
    ensure_defaults()
    return Response([_rule_json(r) for r in PointRule.objects.all()])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def levels(request):
    from .defaults import ensure_defaults
    ensure_defaults()
    return Response([_level_json(lv) for lv in Level.objects.all()])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def badges(request):
    profile = _profile(request)
    from .defaults import ensure_defaults
    ensure_defaults()
    if request.query_params.get('all') and _can(profile, 'view'):
        return Response([_badge_json(b) for b in Badge.objects.all()])
    return Response(services.badges_for(profile) if profile else [])


# ── managers ──────────────────────────────────────────────────────────────────

RULE_FIELDS = {'points': int, 'daily_cap': int, 'grace_hours': int, 'is_active': bool,
               'unit_value': float, 'cash_multiplier': float, 'roles': list,
               'name_ar': str, 'name_en': str, 'desc_ar': str, 'desc_en': str}


def _apply(obj, data, fields):
    errors = {}
    changed = {}
    for f, typ in fields.items():
        if f not in data:
            continue
        v = data[f]
        if v is None and f in ('daily_cap', 'grace_hours', 'unit_value'):
            changed[f] = None
            continue
        try:
            if typ is bool:
                v = bool(v)
            elif typ is list:
                if not isinstance(v, list):
                    raise ValueError
                from apps.users.models import ROLE_CHOICES
                valid = {r[0] for r in ROLE_CHOICES}
                if any(x not in valid for x in v):
                    raise ValueError
            elif typ is str:
                v = str(v).strip()
                if not v and f.startswith('name'):
                    raise ValueError
            else:
                v = typ(v)
        except (TypeError, ValueError):
            errors[f] = 'قيمة غير صحيحة — invalid value'
            continue
        changed[f] = v
    return changed, errors


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def rule_detail(request, key):
    profile = _profile(request)
    if not _can(profile, 'edit'):
        return _deny()
    rule = get_object_or_404(PointRule, key=key)
    changed, errors = _apply(rule, request.data, RULE_FIELDS)
    if 'points' in changed and abs(changed['points']) > 1000:
        errors['points'] = 'أقصى قيمة 1000 — max 1000'
    if 'cash_multiplier' in changed and not (0 <= changed['cash_multiplier'] <= 5):
        errors['cash_multiplier'] = 'من 0 إلى 5 — 0 to 5'
    if errors:
        return Response(errors, status=400)
    before = _rule_json(rule)
    for f, v in changed.items():
        setattr(rule, f, v)
    rule.save()
    _log(profile, 'rule_edit', f'rule:{key}', before, _rule_json(rule),
         str(request.data.get('reason', ''))[:300])
    return Response(_rule_json(rule))


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def level_detail(request, number):
    profile = _profile(request)
    if not _can(profile, 'edit'):
        return _deny()
    lv = get_object_or_404(Level, number=number)
    changed, errors = _apply(lv, request.data, {'min_xp': int, 'title_ar': str, 'title_en': str,
                                                'icon': str, 'color': str})
    if 'min_xp' in changed:
        prev = Level.objects.filter(number__lt=number).order_by('-number').first()
        nxt = Level.objects.filter(number__gt=number).order_by('number').first()
        if number == 1 and changed['min_xp'] != 0:
            errors['min_xp'] = 'المستوى الأول يبدأ من صفر — level 1 starts at 0'
        elif (prev and changed['min_xp'] <= prev.min_xp) or (nxt and changed['min_xp'] >= nxt.min_xp):
            errors['min_xp'] = 'يجب أن يكون بين المستوى السابق والتالي — must sit between neighbours'
    for f in ('title_ar', 'title_en'):
        if f in changed and not changed[f]:
            errors[f] = 'مطلوب — required'
    if errors:
        return Response(errors, status=400)
    before = _level_json(lv)
    for f, v in changed.items():
        setattr(lv, f, v)
    lv.save()
    _log(profile, 'level_edit', f'level:{number}', before, _level_json(lv))
    return Response(_level_json(lv))


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def badge_detail(request, key):
    profile = _profile(request)
    if not _can(profile, 'edit'):
        return _deny()
    b = get_object_or_404(Badge, key=key)
    changed, errors = _apply(b, request.data, {'threshold': int, 'is_active': bool, 'icon': str,
                                               'name_ar': str, 'name_en': str,
                                               'desc_ar': str, 'desc_en': str})
    if 'threshold' in changed and changed['threshold'] < 1:
        errors['threshold'] = 'أقل قيمة 1 — min 1'
    if errors:
        return Response(errors, status=400)
    before = _badge_json(b)
    for f, v in changed.items():
        setattr(b, f, v)
    b.save()
    _log(profile, 'badge_edit', f'badge:{key}', before, _badge_json(b))
    return Response(_badge_json(b))


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def adjust(request):
    """Manual recognition (+) or deduction (−) with a written reason. Audited."""
    profile = _profile(request)
    if not _can(profile, 'edit'):
        return _deny()
    from apps.users.models import StaffProfile
    staff = StaffProfile.objects.filter(id=_int(request.data.get('staff_id')),
                                        is_active=True).first()
    points = _int(request.data.get('points'))
    reason = str(request.data.get('reason') or '').strip()
    if not staff:
        return Response({'staff_id': 'موظف غير موجود — unknown staff'}, status=400)
    if staff.id == profile.id:
        return Response({'staff_id': 'لا يمكنك منح نفسك نقاطاً — you cannot award yourself'},
                        status=400)
    if not points or abs(points) > MAX_MANUAL_POINTS:
        return Response({'points': f'من 1 إلى {MAX_MANUAL_POINTS} (أو بالسالب) — 1..{MAX_MANUAL_POINTS}'},
                        status=400)
    if len(reason) < 5:
        return Response({'reason': 'اكتب السبب — a reason is required'}, status=400)
    sc = engine.Scorer()
    rule = sc.rules['manual_award']
    if not rule.is_active:
        return Response({'detail': 'التقدير اليدوي معطّل — manual awards are disabled'}, status=400)
    now = timezone.now()
    with transaction.atomic():
        # manual awards are not limited by role or daily cap
        from .models import PointEvent
        ev = PointEvent.objects.create(
            staff=staff, rule_key='manual_award', category='manual', points=points,
            day=timezone.localdate(now), occurred_at=now,
            source_key=f'manual:{profile.id}:{now.timestamp()}', branch_id=staff.branch_id,
            reason=reason[:300], created_by=profile)
        _log(profile, 'manual_award', f'staff:{staff.id}', {},
             {'points': points, 'event_id': ev.id}, reason[:300])
    engine.refresh_players([staff.id])
    return Response({'id': ev.id, 'points': ev.points}, status=201)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def run_now(request):
    profile = _profile(request)
    if not _can(profile, 'edit'):
        return _deny()
    days = max(0, min(_int(request.data.get('days')) or 0, 31))
    result = engine.run(days=days, finalize=bool(request.data.get('finalize', days > 0)))
    _log(profile, 'run', 'engine', {}, result)
    return Response(result)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def report_overview(request):
    profile = _profile(request)
    if not _can(profile, 'view'):
        return _deny()
    (start, end), period = _range(request)
    data = services.report(start, end, branch_id=_int(request.query_params.get('branch')))
    data['period'] = period
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def report_export(request):
    profile = _profile(request)
    if not _can(profile, 'export'):
        return _deny()
    (start, end), _ = _range(request)
    data = services.report(start, end, branch_id=_int(request.query_params.get('branch')))
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'الموظفون'
    ws.sheet_view.rightToLeft = True
    ws.append(['الترتيب', 'الموظف', 'الدور', 'الفرع', 'المستوى', 'صافي النقاط', 'المكتسب',
               'الخصومات', 'عدد الحركات', 'السلسلة'])
    for r in data['staff']:
        lv = r['level'] or {}
        ws.append([r['rank'], r['name'], r['role_label'], r['branch'],
                   f"{lv.get('number', '')} {lv.get('title_ar', '')}".strip(),
                   r['net'], r['earned'], r['lost'], r['events'], r['streak']])
    if data['branches']:
        wb2 = wb.create_sheet('الفروع')
        wb2.sheet_view.rightToLeft = True
        wb2.append(['الترتيب', 'الفرع', 'اللاعبون', 'صافي النقاط', 'متوسط اللاعب',
                    'نسبة الأيام النظيفة %', 'بنود متأخرة'])
        for r in data['branches']:
            wb2.append([r['rank'], r['branch'], r['players'], r['net'], r['avg_per_player'],
                        r['clean_day_rate'], r['left_behind']])
    ws3 = wb.create_sheet('الفئات')
    ws3.sheet_view.rightToLeft = True
    ws3.append(['الفئة', 'المكتسب', 'الخصومات', 'عدد الحركات'])
    for c in data['categories']:
        ws3.append([c['label_ar'], c['earned'], c['lost'], c['events']])
    ws4 = wb.create_sheet('الترقيات')
    ws4.sheet_view.rightToLeft = True
    ws4.append(['الموظف', 'الفرع', 'المستوى', 'التاريخ'])
    for p in data['promotions']:
        ws4.append([p['name'], p['branch'], f"{p['level']['number']} {p['level']['title_ar']}",
                    timezone.localtime(p['reached_at']).strftime('%Y-%m-%d %H:%M')])
    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="gamification_{start}_{end}.xlsx"'
    wb.save(resp)
    return resp


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def changes(request):
    profile = _profile(request)
    if not _can(profile, 'view'):
        return _deny()
    rows = GamificationChange.objects.select_related('actor__user')[:200]
    return Response([{'id': c.id, 'actor': c.actor.full_name if c.actor else '',
                      'action': c.action, 'target': c.target, 'before': c.before,
                      'after': c.after, 'reason': c.reason, 'created_at': c.created_at}
                     for c in rows])
