"""
Export a SOFTECH-edit report (from `scan_softech_receipt_edits --out report.json`)
to an Excel workbook for revision — every receipt/line where SOFTECH diverges from
our import, with the frozen / HQ / branch values side by side and the flags.

READ-ONLY (consumes the JSON the read-only sweep produced).

Usage:
  manage.py export_softech_edits_excel --in report.json --out softech_edits.xlsx
"""
import json

from django.core.management.base import BaseCommand, CommandError


class Command(BaseCommand):
    help = 'Export the SOFTECH-edit sweep JSON to an Excel workbook.'

    def add_arguments(self, parser):
        parser.add_argument('--in', dest='infile', required=True)
        parser.add_argument('--out', dest='outfile', required=True)

    def handle(self, *args, **o):
        try:
            with open(o['infile'], encoding='utf-8') as f:
                data = json.load(f)
        except Exception as e:
            raise CommandError(f'Could not read {o["infile"]}: {e}')

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        # ── Sheet 1: summary ──
        ws = wb.active
        ws.title = 'ملخص'
        c = data.get('counts', {})
        ws.append(['SOFTECH Edit Scan — ملخص'])
        ws['A1'].font = Font(bold=True, size=14)
        ws.append([])
        ws.append(['إجمالى الإيصالات المفحوصة', c.get('receipts', 0)])
        ws.append(['بنود مفحوصة', c.get('lines_checked', 0)])
        ws.append(['إيصالات مُعدَّلة (SOFTECH ≠ الاستيراد)', c.get('edited', 0)])
        ws.append(['إيصالات غير متطابقة (HQ ≠ الفرع)', c.get('inconsistent', 0)])
        ws.append([])
        ws.append(['حسب الفرع', 'مفحوصة', 'مُعدَّلة', 'غير متطابقة'])
        ws['A8'].font = Font(bold=True)
        for br, s in (data.get('by_branch') or {}).items():
            ws.append([br, s.get('receipts', 0), s.get('edited', 0), s.get('inconsistent', 0)])

        # ── Sheet 2: line-level detail ──
        d = wb.create_sheet('التفاصيل')
        headers = ['المطالبة', 'رقم الإيصال', 'الفرع', 'كود الصنف', 'اسم الصنف',
                   'مُعدَّل', 'غير متطابق',
                   'سعر مُستورد', 'سعر HQ', 'سعر الفرع',
                   'إجمالى مُستورد', 'إجمالى HQ', 'إجمالى الفرع',
                   'حقول HQ≠الفرع']
        d.append(headers)
        for i in range(1, len(headers) + 1):
            d.cell(row=1, column=i).font = Font(bold=True)
            d.cell(row=1, column=i).fill = PatternFill('solid', fgColor='DDDDDD')
        edited_fill = PatternFill('solid', fgColor='FFF3CD')
        incon_fill = PatternFill('solid', fgColor='F8D7DA')
        for rec in data.get('receipts', []):
            for ln in rec.get('lines', []):
                hq = ln.get('hq') or {}
                br = ln.get('branch') or {}
                fr = ln.get('frozen') or {}
                incon = ln.get('inconsistent_fields') or {}
                row = [rec['claim_number'], rec['docnumber'], rec['branch'],
                       ln['itemcode'], ln.get('item_name', ''),
                       'نعم' if ln.get('edited') else '', 'نعم' if incon else '',
                       fr.get('unit_price'), hq.get('itemsaleprice'), br.get('itemsaleprice'),
                       fr.get('line_total'), hq.get('transprice_total'), br.get('transprice_total'),
                       '; '.join(f'{k}:{v[0]}→{v[1]}' for k, v in incon.items())]
                d.append(row)
                r = d.max_row
                if incon:
                    for col in range(1, len(headers) + 1):
                        d.cell(row=r, column=col).fill = incon_fill
                elif ln.get('edited'):
                    for col in range(1, len(headers) + 1):
                        d.cell(row=r, column=col).fill = edited_fill

        for sheet in (ws, d):
            for col in sheet.columns:
                width = max((len(str(cell.value)) for cell in col if cell.value is not None), default=10)
                sheet.column_dimensions[col[0].column_letter].width = min(width + 2, 40)

        wb.save(o['outfile'])
        n = sum(len(r.get('lines', [])) for r in data.get('receipts', []))
        self.stdout.write(self.style.SUCCESS(
            f'Wrote {o["outfile"]} — {len(data.get("receipts", []))} receipts / {n} flagged lines.'))
