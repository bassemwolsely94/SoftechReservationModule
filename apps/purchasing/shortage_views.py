"""
apps/purchasing/shortage_views.py — REST API for the Market-Shortage screen.

  GET  candidates/     — auto-detected candidates (tiered, filterable)
  GET  deltas/         — what changed since last run (new / recovering / re-entered)
  GET  confirmed/      — the sticky confirmed list (+ recovery & auto-detectable flags)
  GET  dismissed/      — items moved out, with reason (retrievable)
  GET  search/?q=      — item lookup for manual add / picking a matching product
  GET  med-types/      — general-classification options for the filter
  POST flag/           — confirm a candidate / add manually  → in_shortage = True
  POST unflag/         — revert a mistaken flag              → in_shortage = False
  POST dismiss/        — move out with a reason (+ matching)  → shortage_dismissed = True
  POST retrieve/       — bring a dismissed item back
  GET  export/         — Excel of a view (candidates|confirmed|dismissed)
  GET  whatsapp/       — plain-text list formatted for WhatsApp
"""
import io
import logging

from django.utils import timezone
from django.db.models import Q
from django.http import HttpResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from apps.catalog.models import Item
from . import shortage as S

logger = logging.getLogger(__name__)


def _staff(request):
    return getattr(request.user, 'staff_profile', None) or \
        getattr(request.user, 'staffprofile', None)


def _cand_json(c, extra=None):
    d = {
        'item_id': c.item_id, 'code': c.softech_id, 'name': c.name,
        'unit_price': c.unit_price, 'med_type': c.med_type, 'med_type_code': c.med_type_code,
        'annual_qty': c.annual_qty, 'monthly_hist': c.monthly_hist, 'stock': c.stock,
        'coverage': c.coverage, 'qty_30d': c.qty_30d, 'sale_months': c.sale_months,
        'suppression': c.suppression, 'tier': c.tier, 'score': c.score,
        'lost_monthly': c.lost_monthly, 'stockout_days': c.stockout_days,
        'already_flagged': c.already_flagged, 'dismissed': c.dismissed,
    }
    if extra:
        d.update(extra)
    return d


def _apply_filters(rows, request):
    """medicine_type (?med=code) + free-text (?q=) over a ShortageCandidate list."""
    med = request.GET.get('med')
    q = (request.GET.get('q') or '').strip().lower()
    if med:
        rows = [c for c in rows if c.med_type_code == med]
    if q:
        rows = [c for c in rows if q in c.name.lower() or q in c.softech_id.lower()]
    return rows


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_candidates(request):
    cands = S.compute_shortage_candidates()
    counts = {k: 0 for k in S.TIER_COUNT_KEYS}
    for c in cands:
        counts[c.tier] = counts.get(c.tier, 0) + 1
    rows = cands
    if request.GET.get('include_flagged', '1') == '0':
        rows = [c for c in rows if not c.already_flagged]
    tier = request.GET.get('tier')
    if tier:
        rows = [c for c in rows if c.tier.startswith(tier)]
    rows = _apply_filters(rows, request)
    total_lost = round(sum(c.lost_monthly for c in rows), 2)
    limit = request.GET.get('limit')
    if limit and limit.isdigit():
        rows = rows[:int(limit)]
    return Response({'total': len(cands), 'tier_counts': counts,
                     'total_lost_monthly': total_lost,
                     'candidates': [_cand_json(c) for c in rows]})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_deltas(request):
    d = S.compute_deltas()
    return Response({
        'has_baseline': d['has_baseline'], 'counts': d['counts'],
        'new':        [_cand_json(c) for c in d['new']],
        'recovering': [_cand_json(c) for c in d['recovering']],
        're_entered': [_cand_json(c) for c in d['re_entered']],
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_confirmed(request):
    S.ensure_snapshot()
    recovery = S.recovery_map()
    autodet  = S.autodetectable_map()
    sig, _run = S.item_signals()
    stockout = S.stockout_map()
    qs = list(Item.objects.filter(in_shortage=True)
              .select_related('shortage_confirmed_by')
              .order_by('-shortage_flagged_at'))
    # #7 back-orders: customers waiting on stock for these items (pending reservations)
    from django.db.models import Count as _Count
    from apps.reservations.models import Reservation
    waiting = {r['item_id']: r['n'] for r in
               Reservation.objects.filter(item_id__in=[i.id for i in qs], status='pending')
               .values('item_id').annotate(n=_Count('id'))}
    rows = []
    for it in qs:
        s = sig.get(it.id, {})
        rows.append({
            'lost_monthly': s.get('lost_monthly', 0.0),
            'stockout_days': stockout.get(it.id, 0),
            'waiting_customers': waiting.get(it.id, 0),
            'item_id': it.id, 'code': it.softech_id, 'name': it.name,
            'unit_price': float(it.unit_price or 0),
            'med_type': it.medicine_type_name_ar or it.medicine_type_name or it.medicine_type or '',
            'med_type_code': it.medicine_type or '',
            'source': it.shortage_source, 'note': it.shortage_note,
            'flagged_at': it.shortage_flagged_at.isoformat() if it.shortage_flagged_at else None,
            'confirmed_by': it.shortage_confirmed_by.full_name if it.shortage_confirmed_by else None,
            'possibly_resolved': it.id in recovery,
            'autodetectable': autodet.get(it.id, 0) if it.shortage_source == 'manual' else 0,
            'softech_synced': it.shortage_softech_synced,
        })
    if request.GET.get('med'):
        rows = [r for r in rows if r['med_type_code'] == request.GET['med']]
    from django.conf import settings
    total_lost = round(sum(r.get('lost_monthly', 0) for r in rows), 2)
    return Response({'count': len(rows), 'items': rows, 'total_lost_monthly': total_lost,
                     'softech_write_enabled': bool(getattr(settings, 'SHORTAGE_SOFTECH_WRITE_ENABLED', False))})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_dismissed(request):
    qs = (Item.objects.filter(shortage_dismissed=True)
          .select_related('shortage_dismissed_by')
          .prefetch_related('shortage_matching_items')
          .order_by('-shortage_dismissed_at'))
    reason_lbl = dict(Item.DISMISS_REASONS)
    qs = list(qs)
    match_ids = {m.id for it in qs for m in it.shortage_matching_items.all()}
    avail = S.branch_availability(list(match_ids))
    rows = []
    for it in qs:
        matches = [{'item_id': m.id, 'code': m.softech_id, 'name': m.name,
                    'availability': avail.get(m.id, {})}
                   for m in it.shortage_matching_items.all()]
        rows.append({
            'item_id': it.id, 'code': it.softech_id, 'name': it.name,
            'med_type': it.medicine_type_name_ar or it.medicine_type_name or it.medicine_type or '',
            'med_type_code': it.medicine_type or '',
            'reason': it.shortage_dismiss_reason,
            'reason_label': reason_lbl.get(it.shortage_dismiss_reason, it.shortage_dismiss_reason),
            'note': it.shortage_dismiss_note,
            'dismissed_at': it.shortage_dismissed_at.isoformat() if it.shortage_dismissed_at else None,
            'dismissed_by': it.shortage_dismissed_by.full_name if it.shortage_dismissed_by else None,
            # a variant is fully covered only if EVERY matching product is in all branches
            'matches_all_branches': bool(matches) and all(m['availability'].get('ok') for m in matches),
            'matching': matches,
        })
    return Response({'count': len(rows), 'reasons': Item.DISMISS_REASONS, 'items': rows})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_med_types(request):
    """Distinct general-classification options among current candidates + confirmed."""
    cands = S.compute_shortage_candidates()
    seen = {}
    for c in cands:
        if c.med_type_code:
            seen[c.med_type_code] = c.med_type
    for it in Item.objects.filter(in_shortage=True).only('medicine_type', 'medicine_type_name_ar', 'medicine_type_name'):
        if it.medicine_type:
            seen[it.medicine_type] = it.medicine_type_name_ar or it.medicine_type_name or it.medicine_type
    opts = sorted(({'code': k, 'label': v} for k, v in seen.items()), key=lambda x: x['label'])
    return Response({'med_types': opts})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_search(request):
    q = (request.GET.get('q') or '').strip()
    if len(q) < 2:
        return Response({'items': []})
    qs = (Item.objects.filter(is_stockable=True)
          .filter(Q(softech_id__icontains=q) | Q(name__icontains=q))
          .only('id', 'softech_id', 'name', 'unit_price', 'in_shortage', 'shortage_dismissed')[:20])
    return Response({'items': [{
        'item_id': i.id, 'code': i.softech_id, 'name': i.name,
        'unit_price': float(i.unit_price or 0),
        'in_shortage': i.in_shortage, 'dismissed': i.shortage_dismissed,
    } for i in qs]})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def shortage_flag(request):
    item = Item.objects.filter(id=request.data.get('item_id')).first()
    if item is None:
        return Response({'detail': 'item not found'}, status=status.HTTP_404_NOT_FOUND)
    item.in_shortage = True
    item.shortage_source = request.data.get('source', 'auto')
    item.shortage_note = (request.data.get('note') or '')[:300]
    item.shortage_flagged_at = timezone.now()
    item.shortage_confirmed_by = _staff(request)
    # confirming clears any prior dismissal
    item.shortage_dismissed = False
    item.shortage_softech_synced = False        # queue the SOFTECH push
    item.save(update_fields=['in_shortage', 'shortage_source', 'shortage_note',
                             'shortage_flagged_at', 'shortage_confirmed_by',
                             'shortage_dismissed', 'shortage_softech_synced'])
    from . import shortage_writer
    synced = shortage_writer.push_item(item)    # best-effort; queued if SOFTECH offline
    logger.info('[SHORTAGE] flag %s (%s) src=%s softech_synced=%s',
                item.softech_id, item.name, item.shortage_source, synced)
    return Response({'ok': True, 'item_id': item.id, 'softech_synced': synced})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def shortage_unflag(request):
    """Revert a mistaken confirmation → in_shortage = False."""
    item = Item.objects.filter(id=request.data.get('item_id')).first()
    if item is None:
        return Response({'detail': 'item not found'}, status=status.HTTP_404_NOT_FOUND)
    item.in_shortage = False
    item.shortage_source = ''
    item.shortage_flagged_at = None
    item.shortage_confirmed_by = None
    item.shortage_softech_synced = False        # queue the SOFTECH clear
    item.save(update_fields=['in_shortage', 'shortage_source', 'shortage_flagged_at',
                             'shortage_confirmed_by', 'shortage_softech_synced'])
    from . import shortage_writer
    synced = shortage_writer.push_item(item)    # clears صنف نواقص + تحذير in SOFTECH
    logger.info('[SHORTAGE] unflag %s (%s) softech_synced=%s', item.softech_id, item.name, synced)
    return Response({'ok': True, 'item_id': item.id, 'softech_synced': synced})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def shortage_dismiss(request):
    """
    Move an item out of the candidate list with a reason (retrievable).
    Body: { item_id, reason, note?, matching_item_id? }
    """
    item = Item.objects.filter(id=request.data.get('item_id')).first()
    if item is None:
        return Response({'detail': 'item not found'}, status=status.HTTP_404_NOT_FOUND)
    reason = request.data.get('reason', 'other')
    valid = {r[0] for r in Item.DISMISS_REASONS}
    if reason not in valid:
        return Response({'detail': f'reason must be one of {sorted(valid)}'},
                        status=status.HTTP_400_BAD_REQUEST)
    was_flagged = item.in_shortage
    item.shortage_dismissed = True
    item.shortage_dismiss_reason = reason
    item.shortage_dismiss_note = (request.data.get('note') or '')[:300]
    item.shortage_dismissed_at = timezone.now()
    item.shortage_dismissed_by = _staff(request)
    # dismissing supersedes a confirmed flag
    item.in_shortage = False
    # push if we must clear a shortage flag OR suspend ordering (on_request / obsolete)
    needs_push = was_flagged or reason in ('on_request', 'obsolete')
    if needs_push:
        item.shortage_softech_synced = False
    item.save(update_fields=['shortage_dismissed', 'shortage_dismiss_reason',
                             'shortage_dismiss_note', 'shortage_dismissed_at',
                             'shortage_dismissed_by', 'in_shortage', 'shortage_softech_synced'])
    if needs_push:
        from . import shortage_writer
        shortage_writer.push_item(item)     # → موقوف for on_request/obsolete; clears shortage flag
    # one OR two matching products (M2M) — accept list or single, cap at 2
    ids = request.data.get('matching_item_ids') or request.data.get('matching_item_id')
    ids = ids if isinstance(ids, list) else ([ids] if ids else [])
    item.shortage_matching_items.set([i for i in ids if i][:2])
    logger.info('[SHORTAGE] dismiss %s (%s) reason=%s', item.softech_id, item.name, reason)
    return Response({'ok': True, 'item_id': item.id})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def shortage_retrieve(request):
    """Bring a dismissed item back into the candidate pool."""
    item = Item.objects.filter(id=request.data.get('item_id')).first()
    if item is None:
        return Response({'detail': 'item not found'}, status=status.HTTP_404_NOT_FOUND)
    item.shortage_dismissed = False
    item.shortage_dismiss_reason = ''
    item.shortage_dismiss_note = ''
    item.shortage_dismissed_at = None
    item.shortage_dismissed_by = None
    restore_supply = item.shortage_supply_suspended   # we had موقوف'd it → restore ordering
    if restore_supply:
        item.shortage_softech_synced = False
    item.save(update_fields=['shortage_dismissed', 'shortage_dismiss_reason',
                             'shortage_dismiss_note', 'shortage_dismissed_at',
                             'shortage_dismissed_by', 'shortage_softech_synced'])
    item.shortage_matching_items.clear()
    if restore_supply:
        from . import shortage_writer
        shortage_writer.push_item(item)     # → أوامر التوريد back to normal (itemnomoreuse='0')
    logger.info('[SHORTAGE] retrieve %s (%s) restore_supply=%s', item.softech_id, item.name, restore_supply)
    return Response({'ok': True, 'item_id': item.id})


# ── Export helpers ────────────────────────────────────────────────────────────
def _rows_for_view(view):
    if view == 'confirmed':
        recovery = S.recovery_map()
        out = []
        for it in Item.objects.filter(in_shortage=True).order_by('-shortage_flagged_at'):
            out.append([it.softech_id, it.name,
                        it.medicine_type_name_ar or it.medicine_type_name or it.medicine_type or '',
                        float(it.unit_price or 0),
                        'يدوي' if it.shortage_source == 'manual' else 'من المحرك',
                        it.shortage_note, 'قد يكون توفّر' if it.id in recovery else '',
                        it.shortage_flagged_at.strftime('%Y-%m-%d') if it.shortage_flagged_at else ''])
        return (['الكود', 'الصنف', 'التصنيف العام', 'سعر', 'المصدر', 'ملاحظة',
                 'حالة التوفّر', 'تاريخ التأكيد'], out)
    if view == 'dismissed':
        lbl = dict(Item.DISMISS_REASONS)
        out = []
        for it in Item.objects.filter(shortage_dismissed=True).prefetch_related('shortage_matching_items'):
            matches = ' | '.join(f'{m.softech_id} — {m.name}' for m in it.shortage_matching_items.all())
            out.append([it.softech_id, it.name,
                        it.medicine_type_name_ar or it.medicine_type_name or it.medicine_type or '',
                        lbl.get(it.shortage_dismiss_reason, ''), it.shortage_dismiss_note,
                        matches,
                        it.shortage_dismissed_at.strftime('%Y-%m-%d') if it.shortage_dismissed_at else ''])
        return (['الكود', 'الصنف', 'التصنيف العام', 'سبب الاستبعاد', 'ملاحظة',
                 'المنتج البديل', 'تاريخ الاستبعاد'], out)
    # default: candidates
    cands = _apply_filters(S.compute_shortage_candidates(), _FakeReq())
    out = [[c.softech_id, c.name, c.med_type, c.tier, c.annual_qty, round(c.monthly_hist, 1),
            c.stock, c.coverage, int(c.suppression * 100), c.unit_price,
            'نعم' if c.already_flagged else ''] for c in cands]
    return (['الكود', 'الصنف', 'التصنيف العام', 'التصنيف', 'مبيعات/سنة', 'معدل/شهر',
             'مخزون', 'تغطية', 'قمع %', 'سعر', 'مؤكَّد؟'], out)


class _FakeReq:
    GET = {}


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_export(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    view = request.GET.get('view', 'candidates')
    header, rows = _rows_for_view(view)
    if view == 'candidates':          # honour filters for the candidate export
        cands = _apply_filters(S.compute_shortage_candidates(), request)
        rows = [[c.softech_id, c.name, c.med_type, c.tier, c.annual_qty, round(c.monthly_hist, 1),
                 c.stock, c.coverage, int(c.suppression * 100), c.unit_price,
                 'نعم' if c.already_flagged else ''] for c in cands]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'نواقص السوق'
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='022871')
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'A2'
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    resp = HttpResponse(buf.read(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="market-shortage-{view}.xlsx"'
    return resp


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_whatsapp(request):
    """Plain-text list formatted for WhatsApp (confirmed by default; ?view=candidates)."""
    view = request.GET.get('view', 'confirmed')
    med = request.GET.get('med')
    lines = ['*قائمة نواقص السوق* 📉', f'_بتاريخ {timezone.now():%Y-%m-%d}_', '']
    if view == 'confirmed':
        qs = Item.objects.filter(in_shortage=True).order_by('name')
        if med:
            qs = qs.filter(medicine_type=med)
        for i, it in enumerate(qs, 1):
            note = f' — {it.shortage_note}' if it.shortage_note else ''
            lines.append(f'{i}. {it.name}{note}')
        if len(lines) == 3:
            lines.append('لا توجد أصناف.')
    else:
        cands = _apply_filters(S.compute_shortage_candidates(), request)
        for i, c in enumerate(cands, 1):
            lines.append(f'{i}. {c.name}')
    lines.append('')
    lines.append(f'_الإجمالي: {max(0, len([l for l in lines if l and l[0].isdigit()]))} صنف_')
    return Response({'text': '\n'.join(lines)})


# ── Tier 2 (#4): auto-substitution suggestions (same molecule, in stock) ──────
_DIRTY_MOLECULES = {'', 'undefined item', 'undefined', 'n/a', 'na', '.', '-'}


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_suggest_matches(request):
    """
    Suggest available substitutes for a shortage item: other stockable/active items
    sharing its scientific name (molecule) that currently have stock. Guards the
    known-dirty name_scientific values so we don't propose garbage.
    """
    import re
    from django.db.models import Sum
    from apps.catalog.models import ItemStock
    item = Item.objects.filter(id=request.GET.get('item_id')).only('id', 'name', 'name_scientific').first()
    if item is None:
        return Response({'detail': 'item not found'}, status=status.HTTP_404_NOT_FOUND)
    mol = (item.name_scientific or '').strip()
    # name_scientific is unreliable (often just the product name), so also match on the
    # item-name prefix (brand + form, up to the first digit) → catches pack-size variants
    # and same-brand alternatives, which is the real substitution need here.
    prefix = re.split(r'\d', item.name or '', 1)[0].strip()
    peer_q = Q()
    if len(mol) >= 4 and mol.lower() not in _DIRTY_MOLECULES:
        peer_q |= Q(name_scientific__iexact=mol)
    if len(prefix) >= 3:
        peer_q |= Q(name__istartswith=prefix)
    if not peer_q:
        return Response({'molecule': mol, 'prefix': prefix, 'suggestions': [], 'reason': 'no_reliable_key'})
    peers = list(Item.objects.filter(peer_q, is_stockable=True, is_active=True, item_archive=False)
                 .exclude(id=item.id).only('id', 'softech_id', 'name', 'unit_price')[:25])
    ids = [p.id for p in peers]
    stock = {r['item_id']: float(r['s'] or 0) for r in
             ItemStock.objects.filter(item_id__in=ids).values('item_id').annotate(s=Sum('quantity_on_hand'))}
    avail = S.branch_availability(ids)
    sug = [{'item_id': p.id, 'code': p.softech_id, 'name': p.name,
            'stock': round(stock.get(p.id, 0), 1), 'availability': avail.get(p.id, {})}
           for p in peers if stock.get(p.id, 0) > 0]
    sug.sort(key=lambda x: -x['stock'])
    return Response({'molecule': mol, 'prefix': prefix, 'suggestions': sug})


# ── Tier 3 (#9): shortage trends from snapshot history ────────────────────────
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def shortage_trends(request):
    from apps.purchasing.models import ShortageSnapshot
    snaps = list(ShortageSnapshot.objects.order_by('created_at')
                 .values('created_at', 'n_candidates', 'n_confirmed')[:90])
    series = [{'date': s['created_at'].date().isoformat(),
               'candidates': s['n_candidates'], 'confirmed': s['n_confirmed']} for s in snaps]
    cands = S.compute_shortage_candidates()
    by_class, lost_by_class, by_tier = {}, {}, {}
    for c in cands:
        k = c.med_type or 'أخرى'
        by_class[k] = by_class.get(k, 0) + 1
        lost_by_class[k] = lost_by_class.get(k, 0) + c.lost_monthly
        by_tier[c.tier] = by_tier.get(c.tier, 0) + 1
    return Response({
        'series': series,
        'by_class': [{'label': k, 'count': v, 'lost': round(lost_by_class.get(k, 0))}
                     for k, v in sorted(by_class.items(), key=lambda x: -x[1])],
        'by_tier': [{'tier': k, 'count': v} for k, v in sorted(by_tier.items())],
        'total_lost': round(sum(c.lost_monthly for c in cands)),
    })
