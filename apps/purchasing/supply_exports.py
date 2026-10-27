"""
apps/purchasing/supply_exports.py — Excel exports for /supply «التوزيعة» and the inter-branch
pairs on «طلبات التوريد / ISR» (owner 2026-10-25). Each file holds exactly what the screen
shows (same filters), branches as «code · name», right-to-left, filters + freeze panes, the
«بيانات الصنف» block after code + name (apps/catalog/item_columns.py — read live at export time),
and a «الشرح» sheet saying where every number comes from. Read-only.
"""
import io

from django.utils import timezone

from apps.catalog.item_columns import ITEM_NOTE

NAVY = '022871'


def _book():
    from openpyxl import Workbook
    return Workbook()


def _sheet(ws, title, heads, rows, widths, freeze='D2', item_block_at=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(name='Arial', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append(r)
    body = Font(name='Arial', size=10)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = freeze
    if item_block_at:
        from apps.catalog.item_columns import style_columns
        style_columns(ws, item_block_at)
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _notes(wb, lines):
    from openpyxl.styles import Font
    ws = wb.create_sheet('الشرح')
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions['A'].width = 120
    for i, t in enumerate(lines, start=1):
        ws.cell(i, 1, t).font = Font(name='Arial', size=13 if i == 1 else 11, bold=(i == 1))


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── «التوزيعة» ──────────────────────────────────────────────────────────────────
CAT_LABELS = {'new_no_rate': 'جديد بلا معدل', 'never_stocked': 'غير مُدرج بفرع', 'over_piled': 'متكدس بفرع',
              'hq_dormant': 'راكد بالرئيسي'}


def distribution_excel(sugs, *, filters=None, run=None) -> bytes:
    """One row per move (suggestion × designated branch)."""
    from apps.catalog.item_columns import ITEM_HEADS, ITEM_WIDTHS, item_cells, load_items
    from apps.finance.recon_labels import branch_label as BL
    items = load_items(codes=[s['itemcode'] for s in sugs])
    rows = []
    for s in sugs:
        fp = s.get('from_position') or {}
        for t in s['to_branches']:
            rows.append([
                CAT_LABELS.get(s['category'], s['category']), s['itemcode'], s['item_name'],
                *item_cells(items.get(str(s['itemcode']))),
                f"{s['sells_at']}/{s['of_branches']}", s.get('age_days'),
                BL(s['from_branch']), fp.get('stock'), fp.get('rate'), fp.get('max'),
                'لحظي' if fp.get('live') else 'تشغيل المحرك',
                BL(t['code']), t.get('qty'), t.get('stock'), t.get('rate'), t.get('max'),
                'لحظي' if t.get('live') else 'تشغيل المحرك',
                s['network_avg'], s['reason']])
    wb = _book()
    _sheet(wb.active, 'التوزيعة',
           ['النوع', 'الكود', 'الصنف', *ITEM_HEADS, 'يُباع بـ (فروع)', 'العمر (يوم)',
            'من (الفرع الحائز)', 'رصيد الحائز', 'بيع الحائز/شهر', 'الحد الأقصى للحائز', 'مصدر رصيد الحائز',
            'إلى (الفرع المستهدف)', 'الكمية المقترحة', 'رصيد المستهدف', 'بيع المستهدف/شهر', 'الحد الأقصى للمستهدف',
            'مصدر رصيد المستهدف', 'بيع الشبكة/شهر', 'السبب'],
           rows, [14, 9, 34, *ITEM_WIDTHS, 9, 9, 22, 10, 10, 10, 12, 22, 10, 10, 10, 10, 12, 10, 40],
           item_block_at=4)
    f = filters or {}
    _notes(wb, [
        f'التوزيعة الاستباقية — {timezone.localtime():%Y-%m-%d %H:%M}',
        '• الفلاتر: ' + (' · '.join(f'{k}: {v}' for k, v in f.items() if v) or 'بدون') + f' — {len(rows)} نقلة.',
        '• الرصيد «لحظي» = من سيرفر الفرع نفسه للأصناف التي لها حد أقصى في SOFTECH؛ وإلا «تشغيل المحرك»'
        + (f' (#{run.pk}، {(run.finished_at or run.started_at):%Y-%m-%d %H:%M})' if run else '') + '.',
        '• معدل البيع الشهري لكل فرع من نفس تشغيل المحرك. الحد الأقصى = حد SOFTECH (المعدل × التغطية).',
        '• هذا الملف للمراجعة فقط — لا يكتب شيئاً في SOFTECH؛ طلبات التوريد تُنشأ من زر «توليد».',
        ITEM_NOTE,
    ])
    return _bytes(wb)


# ── inter-branch pairs (TransferRecommendation) ─────────────────────────────────
def transfer_pairs_excel(*, from_code=None, to_code=None) -> bytes:
    """The latest transfer-recommendation run's lines — one pair (from → to) or all pairs."""
    from apps.catalog.item_columns import ITEM_HEADS, ITEM_WIDTHS, item_cells, load_items
    from apps.finance.recon_labels import branch_label as BL
    from apps.purchasing.models import TransferRecommendation, TransferRecommendationRun
    run = TransferRecommendationRun.objects.filter(status='success').order_by('-id').first()
    qs = (TransferRecommendation.objects.filter(run=run).select_related('item', 'from_branch', 'to_branch')
          .order_by('from_branch__softech_branch_id', 'to_branch__softech_branch_id', '-estimated_value')
          if run else TransferRecommendation.objects.none())
    if from_code:
        qs = qs.filter(from_branch__softech_branch_id=str(from_code))
    if to_code:
        qs = qs.filter(to_branch__softech_branch_id=str(to_code))
    f2 = lambda v: round(float(v or 0), 2)
    qs = list(qs)
    items = load_items(ids=[r.item_id for r in qs])
    rows, pairs = [], {}
    for r in qs:
        a, b = r.from_branch.softech_branch_id, r.to_branch.softech_branch_id
        rows.append([BL(a), BL(b), r.item.softech_id if r.item_id else '', r.item.name if r.item_id else '',
                     *item_cells(items.get(r.item_id)),
                     f2(r.quantity), f2(r.pack_price), f2(r.estimated_value),
                     f2(r.from_stock), f2(r.from_monthly_avg), f2(r.from_surplus),
                     f2(r.to_stock), f2(r.to_monthly_avg), f2(r.to_gap)])
        p = pairs.setdefault((a, b), [0, 0.0, 0.0])
        p[0] += 1
        p[1] += float(r.quantity or 0)
        p[2] += float(r.estimated_value or 0)
    wb = _book()
    _sheet(wb.active, 'الأصناف',
           ['من (الفرع صاحب الفائض)', 'إلى (الفرع صاحب العجز)', 'الكود', 'الصنف', *ITEM_HEADS,
            'الكمية', 'سعر العبوة', 'القيمة',
            'رصيد المُرسِل', 'بيع المُرسِل/شهر', 'فائض المُرسِل', 'رصيد المستلم', 'بيع المستلم/شهر', 'عجز المستلم'],
           rows, [22, 22, 9, 34, *ITEM_WIDTHS, 9, 10, 11, 10, 10, 10, 10, 10, 10], freeze='E2',
           item_block_at=5)
    _sheet(wb.create_sheet(), 'المسارات', ['من', 'إلى', 'عدد الأصناف', 'الكمية', 'القيمة'],
           [[BL(a), BL(b), n, round(q, 1), round(v, 0)] for (a, b), (n, q, v) in
            sorted(pairs.items(), key=lambda kv: -kv[1][2])], [24, 24, 12, 12, 14], freeze='A2')
    _notes(wb, [
        f'التوزيعة بين الفروع (فائض ← الرئيسي ← عجز) — تشغيل #{run.pk if run else "—"}',
        '• ' + (f'المسار: {BL(from_code)} ← {BL(to_code)}' if from_code and to_code else 'كل المسارات') + f' — {len(rows)} صنف.',
        '• الأرصدة ومعدلات البيع والفائض والعجز من تشغيل توصيات التحويل نفسه.',
        '• كل مسار يصبح طلبين مرتبطين عند «توليد»: من الفرع إلى الرئيسي، ثم من الرئيسي إلى الفرع.',
        '• هذا الملف للمراجعة فقط — لا يكتب شيئاً في SOFTECH.',
        ITEM_NOTE,
    ])
    return _bytes(wb)
