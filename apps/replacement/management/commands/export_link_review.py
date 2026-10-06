"""
Owner decision 1b/1d (2026-10-03): Excel REVIEW LIST of supplier-payment links our A/P writer
WROTE into SOFTECH on the six replacement / buy-back accounts — READ-ONLY (Postgres mirror only,
nothing is written anywhere except the .xlsx).

  Sheet «بين فروع»            every written link whose voucher branch ≠ invoice branch
                              (contract-patient payments never cross branches — owner)
  Sheet «نفس الفرع - ثقة منخفضة» written same-branch links made by the low-confidence rule

For each link it shows the invoice (+ patient from the replacement case), the voucher, and two
suggestions computed AS IF the questioned links were removed (only when the amount is within
±1 EGP / ±1 % — otherwise left blank):
  • the same-branch voucher that most likely pays this invoice
  • the same-branch invoice that most likely belongs to this voucher
plus empty «القرار» / «ملاحظات» columns for the reviewer.

    python manage.py export_link_review --out link_review.xlsx
"""
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal

from django.core.management.base import BaseCommand
from django.db.models import F, Sum

ACCOUNTS = ['4469', '4470', '4471', '4472', '3068', '4069']
LINKED_ORIGINS = ['softech', 'written', 'approved']
TOL = Decimal('0.01')
WINDOW = 120   # days


class Command(BaseCommand):
    help = 'Read-only Excel review list of A/P links written on the replacement/buy-back accounts.'

    def add_arguments(self, parser):
        parser.add_argument('--out', default='link_review.xlsx')
        parser.add_argument('--include-medium', action='store_true',
                            help='add a 3rd sheet with same-branch MEDIUM-confidence written links')

    def handle(self, *a, **o):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        from apps.finance.recon_models import Allocation, APInvoice, Payment
        from apps.replacement.models import ReplacementCase

        written = (Allocation.objects.filter(invoice__party__softech_personcode__in=ACCOUNTS, origin='written')
                   .select_related('invoice__party', 'payment', 'candidate'))
        sets = [('بين فروع', written.exclude(payment__branchcode=F('invoice__branchcode'))),
                ('نفس الفرع - ثقة منخفضة', written.filter(payment__branchcode=F('invoice__branchcode'),
                                                          candidate__confidence_class='low'))]
        if o['include_medium']:
            sets.append(('نفس الفرع - ثقة متوسطة', written.filter(payment__branchcode=F('invoice__branchcode'),
                                                                 candidate__confidence_class='medium')))
        flagged = {s: list(qs.order_by('invoice__branchcode', 'invoice__docdate', 'id')) for s, qs in sets}
        flagged_ids = {al.pk for rows in flagged.values() for al in rows}

        # ── in-memory index of every invoice / voucher on the six accounts, with the questioned
        #    links treated as REMOVED (so suggestions show where the money would really go) ──
        inv_linked, pay_linked = defaultdict(Decimal), defaultdict(Decimal)
        for inv_id, pay_id, amt, pk in (Allocation.objects.filter(invoice__party__softech_personcode__in=ACCOUNTS,
                                                                  origin__in=LINKED_ORIGINS)
                                        .values_list('invoice_id', 'payment_id', 'amount', 'pk')):
            if pk in flagged_ids:
                continue
            inv_linked[inv_id] += amt
            pay_linked[pay_id] += amt
        invoices = defaultdict(list)       # (party, branch) → [inv]
        for inv in APInvoice.objects.filter(party__softech_personcode__in=ACCOUNTS, doccode='10').select_related('party'):
            invoices[(inv.party.softech_personcode, inv.branchcode)].append(inv)
        payments = defaultdict(list)
        for p in (Payment.objects.filter(party__softech_personcode__in=ACCOUNTS, cheqtype='20')
                  .exclude(chain_role='reversed').select_related('party')):
            payments[(p.party.softech_personcode, p.branchcode)].append(p)
        patient = {c['purchase_invoice_id']: c for c in ReplacementCase.objects.filter(
            purchase_invoice__party__softech_personcode__in=ACCOUNTS).values(
            'purchase_invoice_id', 'number', 'softech_pic', 'patient_name')}

        def close(amount):
            # only CLOSE suggestions (±1 EGP or ±1 %) — a loose one would mislead the reviewer
            return max(Decimal('1'), amount * Decimal('0.01'))

        def best_voucher(inv, amount):
            cands = []
            for p in payments[(inv.party.softech_personcode, inv.branchcode)]:
                free = p.amount - pay_linked[p.pk]
                if abs(free - amount) > close(amount) or not (inv.docdate - timedelta(days=3) <= p.voucher_date
                                               <= inv.docdate + timedelta(days=WINDOW)):
                    continue
                cands.append((abs(free - amount), abs((p.voucher_date - inv.docdate).days), p, free))
            return min(cands, key=lambda t: (t[0], t[1])) if cands else None

        def best_invoice(pay, amount):
            cands = []
            for inv in invoices[(pay.party.softech_personcode, pay.branchcode)]:
                open_ = inv.doc_value - inv_linked[inv.pk]
                if abs(open_ - amount) > close(amount) or not (pay.voucher_date - timedelta(days=WINDOW) <= inv.docdate
                                                <= pay.voucher_date + timedelta(days=3)):
                    continue
                cands.append((abs(open_ - amount), abs((pay.voucher_date - inv.docdate).days), inv, open_))
            return min(cands, key=lambda t: (t[0], t[1])) if cands else None

        head = ['رقم الربط', 'المورد', 'فرع الفاتورة', 'رقم الفاتورة', 'تاريخ الفاتورة', 'قيمة الفاتورة',
                'الحالة / المريض', 'فرع السند', 'رقم السند', 'تاريخ السند', 'قيمة السند', 'ملاحظة السند',
                'مستخدم السند', 'المبلغ المربوط', 'قاعدة المطابقة', 'الثقة', 'تاريخ الكتابة',
                'سند مقترح للفاتورة (نفس فرعها)', 'تاريخه', 'المتاح فيه',
                'فاتورة مقترحة للسند (نفس فرعه)', 'تاريخها', 'المتبقي فيها', 'مريض الفاتورة المقترحة',
                'القرار', 'ملاحظات المراجع']
        wb = Workbook()
        summary = wb.active
        summary.title = 'ملخص'
        summary.sheet_view.rightToLeft = True
        hdr_fill, hdr_font = PatternFill('solid', fgColor='022871'), Font(bold=True, color='FFFFFF')
        totals = []
        for title, rows in flagged.items():
            ws = wb.create_sheet(title[:31])
            ws.sheet_view.rightToLeft = True
            ws.append(head)
            for c in ws[1]:
                c.fill, c.font, c.alignment = hdr_fill, hdr_font, Alignment(wrap_text=True, vertical='center')
            dv = DataValidation(type='list', formula1='"خاطئ - يُزال,صحيح - يبقى,غير متأكد"', allow_blank=True)
            ws.add_data_validation(dv)
            by_sup = defaultdict(lambda: [0, Decimal('0')])
            for al in rows:
                inv, pay = al.invoice, al.payment
                pt = patient.get(inv.pk) or {}
                bv, bi = best_voucher(inv, al.amount), best_invoice(pay, al.amount)
                pi = patient.get(bi[2].pk) if bi else None
                ws.append([
                    al.pk, f'{inv.party.softech_personcode} {inv.party.name[:40]}', inv.branchcode, inv.docnumber,
                    inv.docdate, float(inv.doc_value),
                    f'{pt.get("number", "")} {pt.get("softech_pic", "")} {pt.get("patient_name", "")}'.strip(),
                    pay.branchcode, pay.cheqsno, pay.voucher_date, float(pay.amount), pay.note, pay.usercode,
                    float(al.amount), al.candidate.strategy if al.candidate_id else '',
                    al.candidate.confidence_class if al.candidate_id else '',
                    al.created_at.strftime('%Y-%m-%d %H:%M') if al.created_at else '',
                    f'{bv[2].branchcode}/{bv[2].cheqsno}' if bv else '', bv[2].voucher_date if bv else None,
                    float(bv[3]) if bv else None,
                    f'{bi[2].branchcode}/{bi[2].docnumber}' if bi else '', bi[2].docdate if bi else None,
                    float(bi[3]) if bi else None,
                    f'{pi["softech_pic"]} {pi["patient_name"]}'.strip() if pi else '', '', ''])
                dv.add(f'Y{ws.max_row}')
                by_sup[inv.party.softech_personcode][0] += 1
                by_sup[inv.party.softech_personcode][1] += al.amount
            ws.freeze_panes = 'A2'
            for col, w in zip('ABCDEFGHIJKLMNOPQRSTUVWXYZ', [9, 26, 7, 10, 11, 11, 30, 7, 9, 11, 11, 28, 8, 11, 14, 8,
                                                              15, 14, 11, 10, 14, 11, 10, 26, 16, 30]):
                ws.column_dimensions[col].width = w
            ws.auto_filter.ref = ws.dimensions
            totals.append((title, len(rows), sum((r.amount for r in rows), Decimal('0')), dict(by_sup)))

        summary.append(['قائمة مراجعة روابط السداد التي كتبها النظام على حسابات البدل (قراءة فقط)'])
        summary.append(['لم يتم تغيير أي شيء في SOFTECH. اختر في عمود «القرار» لكل سطر: خاطئ - يُزال / صحيح - يبقى / غير متأكد.'])
        summary.append(['الاقتراحات محسوبة كأن الروابط محل المراجعة غير موجودة، وفي نفس الفرع فقط.'])
        summary.append([])
        summary.append(['الورقة', 'عدد الروابط', 'المبلغ', 'حسب المورد'])
        for title, n, amt, sup in totals:
            summary.append([title, n, float(amt), ' · '.join(f'{k}: {v[0]} / {float(v[1]):,.0f}' for k, v in sorted(sup.items()))])
        summary.column_dimensions['A'].width = 30
        summary.column_dimensions['D'].width = 90
        summary['A1'].font = Font(bold=True, size=13)
        wb.save(o['out'])
        for title, n, amt, _ in totals:
            self.stdout.write(f'{title}: {n} links, {amt:,.2f}')
        self.stdout.write(self.style.SUCCESS(f'written {o["out"]}'))
