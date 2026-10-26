"""
apps/stock_audit/views.py — mounted at /api/stock-audit/ (doc 32 Phase 3)

  GET  me/                          what this user may do + their branch scope
  GET  summary/                     per check and per branch: open / total / EGP, vs previous period
  GET  findings/                    filtered, paginated list   (?rule=&branch=&status=&severity=&user=&q=&from=&to=&sort=)
  GET  findings/<id>/               one finding
  POST findings/<id>/review/        {status, note}             review (management approve, or own-branch manager)
  POST findings/bulk-review/        {ids, status, note}
  GET  findings/export/             Excel of the current filters
  GET  by-user/                     adjustments per SOFTECH user (check 5 ranking, always shown)
  GET  rules/                       thresholds: defaults, chain values, branch overrides
  PUT  rules/<code>/                {enabled?, params?}        edit thresholds (audited)
  PUT|DELETE rules/<code>/overrides/<branch>/   {enabled, params}
  GET|POST access/                  named branch managers (admin)      POST access/<id>/revoke/
  GET  access/candidates/?q=        staff to grant
  POST refresh/                     re-read the last 3 days from SOFTECH and re-run the checks (background)
  GET  runs/                        last ledger sync + last engine run

Every rule is enforced here; the screen only hides what the user cannot do.
"""
import datetime as _dt
import logging
import threading
from collections import defaultdict
from decimal import Decimal

from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .access import Access
from .models import (StockAuditAccess, StockAuditFinding, StockAuditRule, StockAuditRuleOverride,
                     StockAuditRun, StockLedgerLine, StockLedgerSyncRun)
from .params import ParamError, ensure_rules, validate_params
from .rules import RULES, adjustments_by_user

logger = logging.getLogger('elrezeiky.stock_audit')

GROUPS = {'sales': ['sale_below_cost', 'duplicate_line', 'price_variance_invoice', 'cost_fluctuation'],
          'adjustments': ['stock_adjustment', 'adjustment_after_count'],
          'availability': ['frequent_stockout'],
          'returns': ['return_without_invoice'],
          'timing': ['after_hours_entry']}
SORTS = {'value': ('value_egp', 'id'), '-value': ('-value_egp', '-id'),
         'date': ('occurred_at', 'id'), '-date': ('-occurred_at', '-id')}
BULK_MAX = 500


def _deny(msg='ليس لديك صلاحية — Not allowed'):
    return Response({'detail': msg}, status=status.HTTP_403_FORBIDDEN)


def _bad(msg):
    return Response({'detail': msg}, status=status.HTTP_400_BAD_REQUEST)


def _date(v):
    try:
        return _dt.date.fromisoformat(v) if v else None
    except ValueError:
        return None


def _list(v):
    return [x.strip() for x in (v or '').split(',') if x.strip()]


def _period(request, default_days=30):
    today = timezone.localdate()
    d_to = _date(request.query_params.get('to')) or today
    d_from = _date(request.query_params.get('from')) or (d_to - _dt.timedelta(days=default_days - 1))
    return min(d_from, d_to), max(d_from, d_to)


def _branch_names():
    from apps.branches.models import Branch
    return {b.softech_branch_id: {'ar': b.name_ar or b.name, 'en': b.name or b.name_ar} for b in Branch.objects.all()}


def _user_names(codes):
    from apps.users.models import ERPUser, StaffProfile
    codes = {c for c in codes if c}
    # ERPUser fields are swapped in this project: username = SOFTECH code, user_id = person name
    names = {str(c): (n or '').strip() for c, n in ERPUser.objects.filter(username__in=codes)
             .values_list('username', 'user_id')}
    for p in StaffProfile.objects.filter(softech_user_id__in=codes).select_related('user'):
        names[p.softech_user_id] = p.full_name or names.get(p.softech_user_id, '')
    return names


def _rule_meta(code):
    s = RULES.get(code, {})
    return {'code': code, 'number': s.get('number'), 'name_ar': s.get('name_ar', code), 'name_en': s.get('name_en', code)}


def _finding_json(f, bnames, unames):
    return {
        'id': f.id, 'rule': _rule_meta(f.rule_code), 'severity': f.severity, 'status': f.status,
        'branchcode': f.branchcode, 'branch_name': bnames.get(f.branchcode, {'ar': f.branchcode, 'en': f.branchcode}),
        'itemcode': f.itemcode, 'item_name': f.item.name if f.item_id else '',
        'doccode': f.doccode, 'docnumber': f.docnumber, 'docdate': f.docdate,
        'occurred_at': f.occurred_at, 'softech_user': f.softech_user,
        'user_name': unames.get(f.softech_user, ''), 'staff_id': f.staff_id,
        'value_egp': f.value_egp, 'metric': f.metric, 'threshold': f.threshold,
        'message_ar': f.message_ar, 'message_en': f.message_en, 'evidence': f.evidence,
        'review_note': f.review_note, 'reviewed_by': f.reviewed_by.full_name if f.reviewed_by_id else '',
        'reviewed_at': f.reviewed_at, 'first_seen': f.first_seen, 'last_seen': f.last_seen,
        'next_statuses': sorted(StockAuditFinding.TRANSITIONS.get(f.status, ())),
    }


def _filtered(request, acc):
    qp = request.query_params
    d_from, d_to = _period(request, default_days=int(qp.get('days') or 30))
    qs = acc.scope(StockAuditFinding.objects.all()).filter(docdate__gte=d_from, docdate__lte=d_to)
    rules = _list(qp.get('rule')) or GROUPS.get(qp.get('group') or '', [])
    if rules:
        qs = qs.filter(rule_code__in=rules)
    for param, field in (('branch', 'branchcode'), ('status', 'status'), ('severity', 'severity'),
                         ('user', 'softech_user')):
        vals = _list(qp.get(param))
        if vals:
            qs = qs.filter(**{f'{field}__in': vals})
    q = (qp.get('q') or '').strip()
    if q:
        cond = Q(itemcode=q) | Q(item__name__icontains=q) | Q(message_ar__icontains=q) | Q(personcode=q)
        if q.isdigit():
            cond |= Q(docnumber=int(q))
        qs = qs.filter(cond)
    return qs.order_by(*SORTS.get(qp.get('sort') or '-date', SORTS['-date'])), d_from, d_to


# ── me / summary ──────────────────────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def me(request):
    acc = Access(request)
    data = acc.as_dict()
    data['rules'] = [_rule_meta(c) for c in RULES]
    data['groups'] = GROUPS
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def summary(request):
    acc = Access(request)
    if not acc.can_view:
        return _deny()
    d_from, d_to = _period(request)
    span = (d_to - d_from).days + 1
    p_from, p_to = d_from - _dt.timedelta(days=span), d_from - _dt.timedelta(days=1)
    base = acc.scope(StockAuditFinding.objects.all())
    branches = _list(request.query_params.get('branch'))
    if branches:
        base = base.filter(branchcode__in=branches)
    cur = base.filter(docdate__gte=d_from, docdate__lte=d_to)
    prev = dict(base.filter(docdate__gte=p_from, docdate__lte=p_to)
                .values_list('rule_code').annotate(n=Count('id')))
    open_q = Q(status__in=[StockAuditFinding.ST_OPEN, StockAuditFinding.ST_ESCALATED])
    per_rule = {r['rule_code']: r for r in cur.values('rule_code').annotate(
        total=Count('id'), open=Count('id', filter=open_q),
        critical_open=Count('id', filter=open_q & Q(severity='critical')),
        value_open=Sum('value_egp', filter=open_q))}
    rules = []
    for code in RULES:
        r = per_rule.get(code, {})
        rules.append({**_rule_meta(code), 'total': r.get('total', 0), 'open': r.get('open', 0),
                      'critical_open': r.get('critical_open', 0), 'value_open': r.get('value_open') or Decimal('0'),
                      'previous_total': prev.get(code, 0)})
    bnames = _branch_names()
    matrix = defaultdict(lambda: {'per_rule': {}, 'open': 0, 'value_open': Decimal('0')})
    for row in cur.filter(open_q).values('branchcode', 'rule_code').annotate(n=Count('id'), v=Sum('value_egp')):
        m = matrix[row['branchcode']]
        m['per_rule'][row['rule_code']] = row['n']
        m['open'] += row['n']
        m['value_open'] += row['v'] or Decimal('0')
    blist = [{'branchcode': bc, 'branch_name': bnames.get(bc, {'ar': bc, 'en': bc}), **m}
             for bc, m in sorted(matrix.items(), key=lambda x: -x[1]['value_open'])]
    return Response({'from': d_from, 'to': d_to, 'previous': {'from': p_from, 'to': p_to},
                     'rules': rules, 'branches': blist, 'access': acc.as_dict(), **_run_status()})


# ── findings ──────────────────────────────────────────────────────────────────

class _Page(PageNumberPagination):
    page_size = 50
    page_size_query_param = 'page_size'
    max_page_size = 200


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def findings(request):
    acc = Access(request)
    if not acc.can_view:
        return _deny()
    qs, d_from, d_to = _filtered(request, acc)
    qs = qs.select_related('item', 'reviewed_by__user')
    pager = _Page()
    page = pager.paginate_queryset(qs, request)
    bnames, unames = _branch_names(), _user_names({f.softech_user for f in page})
    resp = pager.get_paginated_response([_finding_json(f, bnames, unames) for f in page])
    resp.data['totals'] = qs.aggregate(value=Sum('value_egp'), n=Count('id'))
    resp.data['from'], resp.data['to'] = d_from, d_to
    return resp


def _get_scoped(acc, pk):
    f = StockAuditFinding.objects.select_related('item', 'reviewed_by__user').filter(pk=pk).first()
    if f is None or not acc.allows_branch(f.branchcode):
        return None
    return f


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def finding_detail(request, pk):
    acc = Access(request)
    if not acc.can_view:
        return _deny()
    f = _get_scoped(acc, pk)
    if f is None:
        return Response({'detail': 'غير موجود — Not found'}, status=404)
    return Response(_finding_json(f, _branch_names(), _user_names({f.softech_user})))


def _apply_review(f, new_status, note, acc, request):
    """Returns an error string, or None when applied."""
    if not acc.allows_branch(f.branchcode):
        return 'خارج فروعك — outside your branches'
    if new_status not in StockAuditFinding.TRANSITIONS.get(f.status, ()):
        return f'لا يمكن الانتقال من «{f.get_status_display()}» إلى هذه الحالة — move not allowed'
    if new_status != StockAuditFinding.ST_REVIEWED and not note:
        return 'اكتب سبباً — a note is required'
    old = f.status
    now = timezone.now()
    who = acc.profile.full_name if acc.profile else ''
    label = dict(StockAuditFinding.STATUS_CHOICES)[new_status]
    line = f'[{timezone.localtime(now):%Y-%m-%d %H:%M} {who}] {label}' + (f': {note}' if note else '')
    f.status, f.reviewed_by, f.reviewed_at = new_status, acc.profile, now
    f.review_note = f'{f.review_note}\n{line}'.strip()
    f.save(update_fields=['status', 'reviewed_by', 'reviewed_at', 'review_note'])
    from apps.audit.models import AuditLog
    AuditLog.log('stock_audit_finding_reviewed', user=acc.profile, obj=f,
                 changes={'status': [old, new_status]}, note=note[:255], request=request,
                 extra={'rule': f.rule_code, 'branch': f.branchcode})
    return None


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def review(request, pk):
    acc = Access(request)
    if not acc.can_review:
        return _deny()
    f = _get_scoped(acc, pk)
    if f is None:
        return Response({'detail': 'غير موجود — Not found'}, status=404)
    err = _apply_review(f, str(request.data.get('status') or ''), str(request.data.get('note') or '').strip(),
                        acc, request)
    if err:
        return _bad(err)
    f.refresh_from_db()
    return Response(_finding_json(f, _branch_names(), _user_names({f.softech_user})))


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def bulk_review(request):
    acc = Access(request)
    if not acc.can_review:
        return _deny()
    ids = request.data.get('ids') or []
    if not isinstance(ids, list) or not ids:
        return _bad('اختر ملاحظة واحدة على الأقل — pick at least one finding')
    if len(ids) > BULK_MAX:
        return _bad(f'الحد الأقصى {BULK_MAX} ملاحظة في المرة — at most {BULK_MAX}')
    new_status = str(request.data.get('status') or '')
    note = str(request.data.get('note') or '').strip()
    done, skipped = 0, []
    with transaction.atomic():
        for f in StockAuditFinding.objects.select_for_update().filter(pk__in=ids):
            err = _apply_review(f, new_status, note, acc, request)
            if err:
                skipped.append({'id': f.id, 'reason': err})
            else:
                done += 1
    return Response({'done': done, 'skipped': skipped})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_findings(request):
    acc = Access(request)
    if not (acc.can_view and acc.can_export):
        return _deny()
    import openpyxl
    from io import BytesIO
    from openpyxl.styles import Alignment, Font, PatternFill
    qs, d_from, d_to = _filtered(request, acc)
    rows = list(qs.select_related('item')[:20000])
    bnames, unames = _branch_names(), _user_names({f.softech_user for f in rows})
    status_ar = dict(StockAuditFinding.STATUS_CHOICES)
    sev_ar = dict(StockAuditFinding.SEVERITY_CHOICES)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'تدقيق حركات المخزون'
    ws.sheet_view.rightToLeft = True
    head = ['الفحص', 'الخطورة', 'الحالة', 'الفرع', 'كود الصنف', 'الصنف', 'نوع المستند', 'رقم المستند',
            'تاريخ المستند', 'وقت الإدخال', 'المستخدم', 'القيمة (ج.م)', 'المقياس', 'الحد', 'الوصف', 'ملاحظات المراجعة']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='022871')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for f in rows:
        ws.append([
            RULES.get(f.rule_code, {}).get('name_ar', f.rule_code), sev_ar.get(f.severity, f.severity),
            status_ar.get(f.status, f.status), bnames.get(f.branchcode, {}).get('ar', f.branchcode),
            f.itemcode, f.item.name if f.item_id else '', f.doccode, f.docnumber,
            f.docdate.isoformat() if f.docdate else '',
            timezone.localtime(f.occurred_at).strftime('%Y-%m-%d %H:%M') if f.occurred_at else '',
            f'{f.softech_user} {unames.get(f.softech_user, "")}'.strip(), float(f.value_egp or 0),
            float(f.metric) if f.metric is not None else None, float(f.threshold) if f.threshold is not None else None,
            f.message_ar, f.review_note,
        ])
    for col, width in zip('ABCDEFGHIJKLMNOP', [22, 9, 13, 22, 10, 34, 9, 12, 12, 16, 22, 12, 10, 9, 70, 40]):
        ws.column_dimensions[col].width = width
    for r in ws.iter_rows(min_row=2, min_col=12, max_col=12):
        for c in r:
            c.number_format = '#,##0.00'
    ws.freeze_panes = 'A2'
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="stock-audit-{d_from}-{d_to}.xlsx"'
    return resp


# ── adjustments by user (check 5 ranking) ─────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def by_user(request):
    acc = Access(request)
    if not acc.can_view:
        return _deny()
    d_from, d_to = _period(request)
    rule = ensure_rules()['stock_adjustment']
    codes = tuple((rule.params or {}).get('doccodes') or RULES['stock_adjustment']['defaults']['doccodes'])
    qs = acc.scope(StockLedgerLine.objects.filter(docdate__gte=d_from, docdate__lte=d_to, doccode__in=codes))
    branches = _list(request.query_params.get('branch'))
    if branches:
        qs = qs.filter(branchcode__in=branches)
    rank = adjustments_by_user(qs.only('branchcode', 'usercode', 'doccode', 'transqty', 'newcostprice'), codes)
    bnames, unames = _branch_names(), _user_names({r['usercode'] for r in rank})
    open_by_user = dict(acc.scope(StockAuditFinding.objects.filter(
        status__in=['open', 'escalated'], docdate__gte=d_from, docdate__lte=d_to))
        .values_list('softech_user').annotate(n=Count('id')))
    for r in rank:
        r['user_name'] = unames.get(r['usercode'], '')
        r['branch_name'] = bnames.get(r['branchcode'], {'ar': r['branchcode'], 'en': r['branchcode']})
        r['open_findings'] = open_by_user.get(r['usercode'], 0)
    return Response({'from': d_from, 'to': d_to, 'doccodes': codes, 'rows': rank})


# ── thresholds ────────────────────────────────────────────────────────────────

def _schema_json(schema):
    return {k: ({'type': v[0], 'min': v[1], 'max': v[2]} if v[0] == 'number' else {'type': v[0]})
            for k, v in schema.items()}


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def rules_list(request):
    acc = Access(request)
    if not acc.can_view:
        return _deny()
    rows = ensure_rules()
    bnames = _branch_names()
    overrides = defaultdict(list)
    for o in StockAuditRuleOverride.objects.select_related('branch', 'rule'):
        overrides[o.rule.code].append({'branchcode': o.branch.softech_branch_id, 'enabled': o.enabled,
                                       'params': o.params,
                                       'branch_name': bnames.get(o.branch.softech_branch_id, {})})
    out = []
    for code, spec in RULES.items():
        r = rows[code]
        out.append({**_rule_meta(code), 'enabled': r.enabled, 'params': {**spec['defaults'], **(r.params or {})},
                    'defaults': spec['defaults'], 'schema': _schema_json(spec['schema']),
                    'overrides': sorted(overrides[code], key=lambda x: x['branchcode']),
                    'updated_at': r.updated_at})
    from apps.branches.models import Branch
    branches = [{'branchcode': b.softech_branch_id, 'branch_name': bnames.get(b.softech_branch_id)}
                for b in Branch.objects.operational().order_by('softech_branch_id')]
    return Response({'rules': out, 'branches': branches, 'can_edit': acc.can_edit_thresholds})


@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def rule_update(request, code):
    acc = Access(request)
    if not acc.can_edit_thresholds:
        return _deny()
    if code not in RULES:
        return Response({'detail': 'unknown rule'}, status=404)
    rule = ensure_rules()[code]
    try:
        params = validate_params(code, request.data.get('params') or {})
    except ParamError as e:
        return _bad(str(e))
    before = {'enabled': rule.enabled, 'params': dict(rule.params or {})}
    if 'enabled' in request.data:
        rule.enabled = bool(request.data.get('enabled'))
    rule.params = {**(rule.params or {}), **params}
    rule.updated_by = acc.profile
    rule.save()
    from apps.audit.models import AuditLog
    AuditLog.log('stock_audit_rule_changed', user=acc.profile, obj=rule, old_data=before,
                 new_data={'enabled': rule.enabled, 'params': rule.params},
                 note=str(request.data.get('note') or '')[:255], request=request)
    return Response({'code': code, 'enabled': rule.enabled, 'params': {**RULES[code]['defaults'], **rule.params}})


@api_view(['PUT', 'DELETE'])
@permission_classes([IsAuthenticated])
def override_update(request, code, branch):
    from apps.branches.models import Branch
    acc = Access(request)
    if not acc.can_edit_thresholds:
        return _deny()
    if code not in RULES:
        return Response({'detail': 'unknown rule'}, status=404)
    b = Branch.objects.filter(softech_branch_id=branch).first()
    if b is None:
        return Response({'detail': 'unknown branch'}, status=404)
    rule = ensure_rules()[code]
    o = StockAuditRuleOverride.objects.filter(rule=rule, branch=b).first()
    before = {'enabled': o.enabled, 'params': o.params} if o else None
    from apps.audit.models import AuditLog
    if request.method == 'DELETE':
        if o:
            o.delete()
            AuditLog.log('stock_audit_override_changed', user=acc.profile, old_data=before, request=request,
                         extra={'rule': code, 'branch': branch, 'deleted': True})
        return Response(status=204)
    try:
        params = validate_params(code, request.data.get('params') or {})
    except ParamError as e:
        return _bad(str(e))
    enabled = request.data.get('enabled', None)
    if enabled is not None and not isinstance(enabled, bool):
        return _bad('enabled must be true, false or null')
    if not params and enabled is None:                    # nothing left → back to the chain values
        if o:
            o.delete()
        AuditLog.log('stock_audit_override_changed', user=acc.profile, old_data=before, request=request,
                     extra={'rule': code, 'branch': branch, 'deleted': True})
        return Response(status=204)
    o = o or StockAuditRuleOverride(rule=rule, branch=b)
    o.params, o.enabled, o.updated_by = params, enabled, acc.profile
    o.save()
    AuditLog.log('stock_audit_override_changed', user=acc.profile, obj=o, old_data=before,
                 new_data={'enabled': o.enabled, 'params': o.params}, request=request,
                 extra={'rule': code, 'branch': branch})
    return Response({'branchcode': branch, 'enabled': o.enabled, 'params': o.params})


# ── named branch managers ─────────────────────────────────────────────────────

def _grant_json(g):
    return {'id': g.id, 'staff_id': g.staff_id, 'staff_name': g.staff.full_name,
            'role': g.staff.role, 'branch': g.staff.branch_name, 'level': g.level, 'note': g.note,
            'granted_by': g.granted_by.full_name if g.granted_by_id else '', 'granted_at': g.granted_at,
            'revoked_by': g.revoked_by.full_name if g.revoked_by_id else '', 'revoked_at': g.revoked_at}


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def access_list(request):
    acc = Access(request)
    if not acc.can_manage_access:
        return _deny()
    if request.method == 'GET':
        qs = StockAuditAccess.objects.select_related('staff__user', 'staff__branch', 'granted_by__user',
                                                     'revoked_by__user')[:200]
        return Response({'rows': [_grant_json(g) for g in qs]})
    from apps.users.models import StaffProfile
    staff = StaffProfile.objects.filter(pk=request.data.get('staff_id'), is_active=True).first()
    if staff is None:
        return _bad('موظف غير موجود — unknown staff')
    if StockAuditAccess.objects.filter(staff=staff, revoked_at__isnull=True).exists():
        return _bad('لديه الصلاحية بالفعل — already granted')
    g = StockAuditAccess.objects.create(staff=staff, granted_by=acc.profile,
                                        note=str(request.data.get('note') or '')[:255])
    from apps.audit.models import AuditLog
    AuditLog.log('stock_audit_access_granted', user=acc.profile, obj=g, note=g.note, request=request,
                 extra={'staff_id': staff.pk, 'staff': staff.full_name})
    return Response(_grant_json(g), status=201)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def access_revoke(request, pk):
    acc = Access(request)
    if not acc.can_manage_access:
        return _deny()
    g = StockAuditAccess.objects.filter(pk=pk, revoked_at__isnull=True).select_related('staff__user').first()
    if g is None:
        return Response({'detail': 'غير موجود — Not found'}, status=404)
    g.revoked_at, g.revoked_by = timezone.now(), acc.profile
    g.save(update_fields=['revoked_at', 'revoked_by'])
    from apps.audit.models import AuditLog
    AuditLog.log('stock_audit_access_revoked', user=acc.profile, obj=g, request=request,
                 note=str(request.data.get('note') or '')[:255], extra={'staff_id': g.staff_id})
    return Response(_grant_json(g))


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def access_candidates(request):
    acc = Access(request)
    if not acc.can_manage_access:
        return _deny()
    from apps.users.models import StaffProfile
    q = (request.query_params.get('q') or '').strip()
    granted = set(StockAuditAccess.objects.filter(revoked_at__isnull=True).values_list('staff_id', flat=True))
    qs = StaffProfile.objects.filter(is_active=True).exclude(role='admin').select_related('user', 'branch')
    if q:
        qs = qs.filter(Q(user__first_name__icontains=q) | Q(user__last_name__icontains=q) |
                       Q(user__username__icontains=q) | Q(softech_username__icontains=q) |
                       Q(softech_user_id=q))
    rows = [{'id': p.id, 'name': p.full_name, 'role': p.role, 'branch': p.branch_name,
             'granted': p.id in granted} for p in qs.order_by('branch__softech_branch_id')[:30]]
    return Response({'rows': rows})


# ── refresh / runs ────────────────────────────────────────────────────────────

_refresh_lock = threading.Lock()


def _run_status():
    ls = StockLedgerSyncRun.objects.order_by('-started_at').first()
    ar = StockAuditRun.objects.order_by('-started_at').first()
    return {
        'ledger_sync': ls and {'id': ls.id, 'status': ls.status, 'started_at': ls.started_at,
                               'finished_at': ls.finished_at, 'rows': ls.rows_synced},
        'audit_run': ar and {'id': ar.id, 'status': ar.status, 'started_at': ar.started_at,
                             'finished_at': ar.finished_at, 'counts': ar.counts},
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def runs(request):
    acc = Access(request)
    if not acc.can_view:
        return _deny()
    return Response({**_run_status(), 'refreshing': _refresh_lock.locked()})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def refresh(request):
    """Re-read the last 3 docdate days from SOFTECH HQ (read-only) and re-run the checks."""
    acc = Access(request)
    if not acc.can_refresh:
        return _deny()
    if not _refresh_lock.acquire(blocking=False):
        return Response({'detail': 'التحديث يعمل بالفعل — a refresh is already running'}, status=409)
    sync_run = StockLedgerSyncRun.objects.create(triggered_by=f'refresh:{request.user}'[:50])

    def _work():
        try:
            from .engine import run_stock_audit
            from .ledger import recent_range, sync_ledger
            sync_ledger(*recent_range(3), run=sync_run)
            run_stock_audit(triggered_by=f'refresh:{request.user}')
        except Exception as e:   # noqa: BLE001
            logger.exception('stock-audit refresh failed')
            if sync_run.status == 'running':
                sync_run.detail = {'fatal': str(e)[:300]}
                sync_run.finish('failed')
        finally:
            _refresh_lock.release()

    threading.Thread(target=_work, daemon=True).start()
    return Response({'detail': 'بدأ التحديث في الخلفية.', 'ledger_run_id': sync_run.pk},
                    status=status.HTTP_202_ACCEPTED)
