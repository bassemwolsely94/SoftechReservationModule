"""
apps/supply/code_report.py — how many supplier lines match by the supplier's OWN product
code (owner 2026-10-05: "once the supplier-code copy is filled, how many lines now match").

Looks at every supplier line in the window that came with a supplier:
  • supplier availability lines (AvailabilityLine.supplier_item_code)
  • supplier invoice lines     (InvoiceLine.vendor_item_code)
and resolves each code with the SAME rule live matching uses
(availability.resolve_vendor_codes: confirmed mapping → SOFTECH itemssuppliers mirror).

Per line: no code · code matched (mapping / mirror) · code ambiguous · code unknown. For a
line a person already confirmed, the code's item is compared with the confirmed one:
agrees, or CONFLICTS (a data-quality problem worth fixing in SOFTECH). "Newly matchable" =
the code now resolves but the line was left unmatched. Barcode matches (the other exact
identity) are counted beside it. Read-only.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import timedelta

from django.utils import timezone

SOURCE_LABELS = {'availability': 'قوائم إتاحة الموردين', 'invoices': 'فواتير الموردين'}


def _blank():
    return {'lines': 0, 'with_code': 0, 'matched': 0, 'by_mapping': 0, 'by_mirror': 0,
            'ambiguous': 0, 'unknown': 0, 'newly_matchable': 0, 'confirmed_with_code': 0,
            'agrees': 0, 'conflicts': 0, 'by_barcode': 0}


def _pct(a, b):
    return round(100 * a / b, 1) if b else None


def mirror_status() -> dict:
    from django.db.models import Count, Max, Q
    from apps.catalog.models import ItemSupplierLink
    agg = ItemSupplierLink.objects.aggregate(
        links=Count('id'), with_code=Count('id', filter=~Q(supp_item_code='')),
        suppliers=Count('supp_code', distinct=True), last=Max('synced_at'))
    return {'links': agg['links'], 'with_code': agg['with_code'], 'suppliers': agg['suppliers'],
            'synced_at': agg['last'], 'empty': not agg['links']}


def _rows(since):
    """(source, line_id, vendor, code, current_item_id, confirmed, via, label)."""
    from apps.invoices.models import InvoiceLine
    from .models import AvailabilityLine
    for ln in (AvailabilityLine.objects.filter(batch__created_at__gte=since, batch__supplier__isnull=False)
               .select_related('batch__supplier').only(
                   'id', 'supplier_item_code', 'item_id', 'is_confirmed', 'match_reason', 'raw_text',
                   'batch__supplier__id', 'batch__supplier__name', 'batch__supplier__softech_personcode')):
        yield ('availability', ln.id, ln.batch.supplier, (ln.supplier_item_code or '').strip(),
               ln.item_id, ln.is_confirmed, (ln.match_reason or {}).get('via'), ln.raw_text)
    for ln in (InvoiceLine.objects.filter(invoice__created_at__gte=since, invoice__vendor__isnull=False)
               .select_related('invoice__vendor').only(
                   'id', 'vendor_item_code', 'item_id', 'is_confirmed', 'manual_name', 'raw_text',
                   'invoice__vendor__id', 'invoice__vendor__name', 'invoice__vendor__softech_personcode')):
        yield ('invoices', ln.id, ln.invoice.vendor, (ln.vendor_item_code or '').strip(),
               ln.item_id, ln.is_confirmed, None, ln.manual_name or ln.raw_text)


def build(days: int = 90) -> dict:
    from apps.catalog.models import Item
    from .availability import resolve_vendor_codes
    since = timezone.now() - timedelta(days=days)
    rows = list(_rows(since))

    # one batched resolution per supplier
    codes_by_vendor = defaultdict(set)
    vendors = {}
    for _s, _i, v, code, *_ in rows:
        vendors[v.id] = v
        if code:
            codes_by_vendor[v.id].add(code)
    resolved = {vid: resolve_vendor_codes(vendors[vid], codes) for vid, codes in codes_by_vendor.items()}

    totals, by_source = _blank(), defaultdict(_blank)
    by_vendor = defaultdict(_blank)
    conflicts, newly = [], []
    for src, lid, v, code, item_id, confirmed, via, label in rows:
        for bucket in (totals, by_source[src], by_vendor[v.id]):
            bucket['lines'] += 1
            if via == 'barcode':
                bucket['by_barcode'] += 1
        if not code:
            continue
        iid, how = resolved[v.id].get(code, (None, 'unknown'))
        for bucket in (totals, by_source[src], by_vendor[v.id]):
            bucket['with_code'] += 1
            if iid:
                bucket['matched'] += 1
                bucket['by_mapping' if how == 'mapping' else 'by_mirror'] += 1
            else:
                bucket['ambiguous' if how == 'ambiguous' else 'unknown'] += 1
            if iid and not item_id:
                bucket['newly_matchable'] += 1
            if confirmed and item_id:
                bucket['confirmed_with_code'] += 1
                if iid == item_id:
                    bucket['agrees'] += 1
                elif iid:
                    bucket['conflicts'] += 1
        if iid and confirmed and item_id and iid != item_id:
            conflicts.append({'source': src, 'line_id': lid, 'supplier': v.name, 'code': code,
                              'text': (label or '')[:120], 'confirmed_item_id': item_id, 'code_item_id': iid,
                              'via': how})
        elif iid and not item_id:
            newly.append({'source': src, 'line_id': lid, 'supplier': v.name, 'code': code,
                          'text': (label or '')[:120], 'code_item_id': iid, 'via': how})

    names = dict(Item.objects.filter(id__in={c['confirmed_item_id'] for c in conflicts}
                                     | {c['code_item_id'] for c in conflicts + newly})
                 .values_list('id', 'name'))
    for c in conflicts:
        c['confirmed_item'] = names.get(c['confirmed_item_id'], '')
        c['code_item'] = names.get(c['code_item_id'], '')
    for c in newly:
        c['code_item'] = names.get(c['code_item_id'], '')

    def with_rates(b):
        return {**b, 'code_rate': _pct(b['with_code'], b['lines']),
                'match_rate': _pct(b['matched'], b['with_code']),
                'agree_rate': _pct(b['agrees'], b['agrees'] + b['conflicts'])}

    from apps.catalog.models import ItemSupplierLink
    from django.db.models import Count, Q
    mirror_by_supp = {r['supp_code']: r for r in ItemSupplierLink.objects.values('supp_code').annotate(
        links=Count('id'), with_code=Count('id', filter=~Q(supp_item_code='')))}
    suppliers = []
    for vid, b in by_vendor.items():
        v = vendors[vid]
        m = mirror_by_supp.get((v.softech_personcode or '').strip(), {})
        suppliers.append({'vendor_id': vid, 'supplier': v.name, 'personcode': v.softech_personcode,
                          'mirror_links': m.get('links', 0), 'mirror_codes': m.get('with_code', 0),
                          **with_rates(b)})
    suppliers.sort(key=lambda r: (-r['with_code'], -r['lines']))
    return {
        'days': days, 'generated_at': timezone.now(), 'mirror': mirror_status(),
        'totals': with_rates(totals),
        'by_source': {k: {'label': SOURCE_LABELS[k], **with_rates(by_source[k])} for k in SOURCE_LABELS},
        'suppliers': suppliers,
        'conflicts': conflicts[:200], 'newly_matchable': newly[:200],
    }


# ── Excel ──────────────────────────────────────────────────────────────────────
def to_excel(rep) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    navy, white = '022871', 'FFFFFF'
    hfont, hfill = Font(name='Arial', bold=True, color=white), PatternFill('solid', fgColor=navy)
    wb = Workbook()

    def sheet(ws, title, heads, rows, widths):
        ws.sheet_view.rightToLeft = True
        ws.append([title])
        ws['A1'].font = Font(name='Arial', size=13, bold=True, color=navy)
        ws.append(heads)
        for c in ws[2]:
            c.font, c.fill, c.alignment = hfont, hfill, Alignment(horizontal='center', wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(2, i).column_letter].width = w
        ws.freeze_panes = 'A3'

    t, m = rep['totals'], rep['mirror']
    ws = wb.active
    ws.title = 'ملخص'
    ws.sheet_view.rightToLeft = True
    lines = [
        ['مطابقة أسطر الموردين بكود المورد', ''],
        ['الفترة', f'آخر {rep["days"]} يوم'],
        ['نسخة SOFTECH (itemssuppliers)', 'فارغة — لم تتم المزامنة بعد' if m['empty']
         else f'{m["links"]:,} ربط · {m["with_code"]:,} بكود مورد · {m["suppliers"]} مورد'],
        ['آخر مزامنة للنسخة', str(m['synced_at'] or '—')[:16]],
        ['', ''],
        ['إجمالي الأسطر', t['lines']],
        ['أسطر بها كود مورد', f'{t["with_code"]} ({t["code_rate"] or 0}%)'],
        ['مطابقة بالكود الآن', f'{t["matched"]} ({t["match_rate"] or 0}% من الأسطر ذات الكود)'],
        ['  — من مطابقات سابقة مؤكدة', t['by_mapping']],
        ['  — من نسخة SOFTECH', t['by_mirror']],
        ['كود لأكثر من صنف (لا تخمين)', t['ambiguous']],
        ['كود غير معروف', t['unknown']],
        ['أسطر غير مطابقة يمكن مطابقتها الآن بالكود', t['newly_matchable']],
        ['أسطر مؤكدة: الكود يوافق الصنف المؤكد', f'{t["agrees"]} ({t["agree_rate"] or 0}%)'],
        ['أسطر مؤكدة: الكود يشير لصنف آخر (تعارض)', t['conflicts']],
        ['مطابقة بالباركود', t['by_barcode']],
    ]
    for r in lines:
        ws.append(r)
    ws['A1'].font = Font(name='Arial', size=13, bold=True, color=navy)
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 46, 52

    heads = ['المورد', 'كود SOFTECH', 'الأسطر', 'بها كود', 'مطابقة بالكود', 'من مطابقات سابقة', 'من نسخة SOFTECH',
             'غامض', 'غير معروف', 'يمكن مطابقتها الآن', 'يوافق المؤكد', 'تعارض', 'باركود',
             'روابط في النسخة', 'أكواد في النسخة']
    sheet(wb.create_sheet('الموردون'), 'حسب المورد', heads,
          [[s['supplier'], s['personcode'], s['lines'], s['with_code'], s['matched'], s['by_mapping'],
            s['by_mirror'], s['ambiguous'], s['unknown'], s['newly_matchable'], s['agrees'], s['conflicts'],
            s['by_barcode'], s['mirror_links'], s['mirror_codes']] for s in rep['suppliers']],
          [28, 11] + [12] * 13)
    sheet(wb.create_sheet('تعارضات'), 'الكود يشير لصنف غير الصنف المؤكد — راجعها في SOFTECH',
          ['المصدر', 'رقم السطر', 'المورد', 'الكود', 'نص السطر', 'الصنف المؤكد', 'صنف الكود', 'مصدر الكود'],
          [[SOURCE_LABELS[c['source']], c['line_id'], c['supplier'], c['code'], c['text'], c['confirmed_item'],
            c['code_item'], 'مطابقة سابقة' if c['via'] == 'mapping' else 'نسخة SOFTECH'] for c in rep['conflicts']],
          [18, 10, 24, 12, 40, 36, 36, 14])
    sheet(wb.create_sheet('يمكن مطابقتها'), 'أسطر غير مطابقة يحدد كودها الصنف الآن',
          ['المصدر', 'رقم السطر', 'المورد', 'الكود', 'نص السطر', 'صنف الكود'],
          [[SOURCE_LABELS[c['source']], c['line_id'], c['supplier'], c['code'], c['text'], c['code_item']]
           for c in rep['newly_matchable']],
          [18, 10, 24, 12, 40, 36])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
