"""
/supply «تلبية طلبات الفروع» — ISR fulfilment plan (apps/purchasing/isr_fulfillment.py):
the deterministic plan (ceiling, whole-pack excess, donor cascade → HQ → shortage), the
formula workbook, and the read-only API (access, validation, snapshot reuse, export).
SOFTECH is never touched here: collect() is replaced by a fixed snapshot.
"""
import io
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase, SimpleTestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.purchasing import isr_fulfillment as F

User = get_user_model()


def _snap():
    """ISR 261609 from branch 160; donors 130/140/150; HQ = '100'.
       A: 10 req. 160 has 2 @ rate 4 (max 6 → 12 after = over). 150 has 20 @ rate 2 (max 3,
          excess 17) → gives all 10.
       B: 15 req. 130 has 5, no sales (max 0, excess 5); 140 has 3.6 @ rate 1 (max 1.5,
          excess 2 whole packs); HQ 4.7 → 4 whole packs; shortage 4.
       C: 3 req, nothing anywhere → shortage 3."""
    return {
        'isr_numbers': ['261609'],
        'isrs': [{'isr': '261609', 'branch': '160', 'for_branch': '100', 'date': '2026-10-01',
                  'approved': 1, 'value': 200.0, 'user': '11', 'lines': [
                      {'code': 'A', 'qty': 10.0, 'nowqty_at_request': 2, 'cost': 5.0, 'price': 7},
                      {'code': 'B', 'qty': 15.0, 'nowqty_at_request': 0, 'cost': 10.0, 'price': 12},
                      {'code': 'C', 'qty': 3.0, 'nowqty_at_request': 0, 'cost': 2.0, 'price': 3},
                  ]}],
        'branches': ['130', '140', '150', '160'],
        'branch_names': {'160': 'فرع 160'},
        'stock': {'160': {'A': 2.0}, '150': {'A': 20.0}, '130': {'B': 5.0}, '140': {'B': 3.6},
                  '100': {'B': 4.7}},
        'eng': {'A|160': [4.0, 4.0], 'A|150': [2.0, -17.0], 'B|140': [1.0, -1.0]},
        'names': {'A': 'ITEM A', 'B': 'ITEM B', 'C': 'ITEM C'},
        'down': [], 'run_id': 81, 'taken_at': '2026-10-02T10:00:00',
    }


class CeilingAndParseTests(SimpleTestCase):
    def test_branch_column_order(self):
        self.assertEqual(F.branch_column_order(['150', '130', '170', '140', '160']),
                         ['130', '140', '150', '160', '170'])
        self.assertEqual(F.branch_column_order(['200', '130', 'CC']), ['130', '200', 'CC'])

    def test_ceiling_matches_workbook_formula(self):
        self.assertEqual(F.ceiling(0.9, 1.5), 1.4)       # 1.35 → 1.4 half-up (no float drift)
        self.assertEqual(F.ceiling(0.3, 1.5), 1.0)       # 0.45 → never below one pack
        self.assertEqual(F.ceiling(0, 1.5), 0.0)         # no sales → no ceiling
        self.assertEqual(F.ceiling(12.38, 1.5), 18.6)

    def test_parse_numbers(self):
        self.assertEqual(F.parse_isr_numbers('261609, 2615043 ،261609'), ['261609', '2615043'])
        self.assertEqual(F.parse_isr_numbers(['261609', 2615043]), ['261609', '2615043'])
        for bad in ('', '26A1', '1; DROP TABLE x', '12345678901'):
            with self.assertRaises(ValueError):
                F.parse_isr_numbers(bad)
        with self.assertRaises(ValueError):
            F.parse_isr_numbers(' '.join(str(i) for i in range(F.MAX_ISRS + 1)))


class ComputeTests(SimpleTestCase):
    def test_plan_cascade_hq_and_shortage(self):
        p = F.compute(_snap(), 1.5)
        r = p['isrs'][0]
        self.assertEqual(r['donors'], ['150', '130', '140'])         # most excess first
        a, b, c = r['lines']
        self.assertEqual((a['from_branches'], a['from_hq'], a['shortage']), (10, 0, 0))
        self.assertTrue(a['req']['over_ceiling'])
        self.assertEqual(a['req']['ceiling'], 6.0)
        self.assertEqual({d['branch']: d['give'] for d in b['donors']}, {'150': 0, '130': 5, '140': 2})
        self.assertEqual((b['from_branches'], b['from_hq'], b['shortage']), (7, 4, 4))
        self.assertEqual((c['from_branches'], c['from_hq'], c['shortage']), (0, 0, 3))
        self.assertIsNone(c['req']['stock'])                         # no stkbal row → blank
        t = r['totals']
        self.assertEqual((t['requested'], t['from_branches'], t['from_hq'], t['shortage']), (28, 17, 4, 7))
        self.assertEqual((t['full_branches'], t['full_all'], t['over_ceiling']), (1, 1, 1))
        self.assertEqual((t['value_branches'], t['value_shortage']), (120.0, 46.0))
        self.assertEqual(r['from_by_donor'], {'150': 10, '130': 5, '140': 2})
        self.assertEqual(r['user_name'], '')                           # snapshot carries no name
        self.assertEqual(p['totals']['shortage'], 7)

    def test_hq_fills_alone_when_it_covers_the_whole_request(self):
        """Owner rule 2026-10-04: HQ enough → one shipment from HQ, no branch transfer."""
        s = _snap()
        s['stock']['100'] = {'A': 10.0, 'B': 4.7}                     # A: HQ 10 ≥ 10 requested
        a, b, _ = F.compute(s, 1.5)['isrs'][0]['lines']
        self.assertEqual((a['hq_first'], a['from_hq'], a['from_branches'], a['shortage']), (True, 10, 0, 0))
        self.assertTrue(all(d['give'] == 0 for d in a['donors']))      # branch 150 keeps its excess
        self.assertEqual((b['hq_first'], b['from_branches'], b['from_hq']), (False, 7, 4))   # HQ short → cascade
        s['stock']['100']['A'] = 9.9                                    # 9 whole packs < 10 → cascade
        a = F.compute(s, 1.5)['isrs'][0]['lines'][0]
        self.assertEqual((a['hq_first'], a['from_branches'], a['from_hq']), (False, 10, 0))
        ws = F.build_workbook(s, 1.5)['ISR 261609']
        heads = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
        self.assertIn('الرئيسي يغطي الطلب؟', heads)
        donor_col = heads.index('من فرع 150') + 1
        self.assertTrue(str(ws.cell(6, donor_col).value).startswith('=IF('))

    def test_higher_coverage_donors_keep_more(self):
        b = F.compute(_snap(), 3)['isrs'][0]['lines'][1]
        self.assertEqual((b['from_branches'], b['from_hq'], b['shortage']), (5, 4, 6))   # 140 keeps 3
        a = F.compute(_snap(), 3)['isrs'][0]['lines'][0]
        self.assertFalse(a['req']['over_ceiling'])                   # 12 after vs max 12


class WorkbookTests(SimpleTestCase):
    def test_sheets_coverage_cell_and_formulas(self):
        wb = F.build_workbook(_snap(), 2)
        self.assertEqual(wb.sheetnames, ['ملخص', 'ISR 261609', 'الشرح'])
        self.assertEqual(wb['ملخص']['C4'].value, 2.0)
        ws = wb['ISR 261609']
        groups = [ws.cell(4, c).value for c in range(1, ws.max_column + 1) if ws.cell(4, c).value]
        at = lambda code: next(i for i, g in enumerate(groups) if str(g).startswith(f'فرع {code}'))  # «فرع 150 · name»
        # owner 2026-10-05: the requesting branch first, then the others by code 130 → 140 → 150
        self.assertTrue(str(groups[1]).startswith('الفرع الطالب 160'))
        self.assertLess(at('130'), at('140'))
        self.assertLess(at('140'), at('150'))
        # …while the draw order stays «most excess first» and is written on the sheet
        self.assertIn('150 ← 130 ← 140', ws.cell(2, 1).value)
        formulas = [ws.cell(6, c).value for c in range(1, ws.max_column + 1)]
        self.assertTrue(any(isinstance(v, str) and F.COV_REF in v for v in formulas))
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertEqual(F.workbook_filename(_snap()), 'isr_fulfillment_261609_2026-10-02.xlsx')


def _user(username, *, role='salesperson', group=None):
    from apps.users.models import ERPUser, StaffProfile
    u = User.objects.create_user(username=username, password='x')
    StaffProfile.objects.create(user=u, role=role)
    if group is not None:
        ERPUser.objects.create(username=username, user_group=str(group), is_active=True)
    return u


@override_settings(SUPPLY_ERP_GROUPS='10,19,21,27')
class FulfilmentApiTests(TestCase):
    URL = '/api/purchasing/isr/fulfilment/'

    def setUp(self):
        cache.clear()
        self.c = APIClient()

    def test_needs_supply_access(self):
        self.c.force_authenticate(_user('1399', group=14))            # صيدلي group
        self.assertEqual(self.c.post(self.URL, {'isrs': '261609'}, format='json').status_code, 403)

    def test_bad_input(self):
        self.c.force_authenticate(_user('18', group=19))
        self.assertEqual(self.c.post(self.URL, {'isrs': 'abc'}, format='json').status_code, 400)
        self.assertEqual(self.c.post(self.URL, {'isrs': '261609', 'coverage': 50},
                                     format='json').status_code, 400)

    def test_missing_isr_is_404(self):
        self.c.force_authenticate(_user('18', group=19))
        with mock.patch.object(F, 'collect', side_effect=F.IsrNotFound(['999'])):
            r = self.c.post(self.URL, {'isrs': '999'}, format='json')
        self.assertEqual(r.status_code, 404)
        self.assertEqual(r.data['missing'], ['999'])

    def test_plan_then_replan_reuses_snapshot_then_export(self):
        self.c.force_authenticate(_user('18', group=19))
        with mock.patch.object(F, 'collect', return_value=_snap()) as col:
            r1 = self.c.post(self.URL, {'isrs': '261609', 'coverage': 1.5}, format='json')
            self.assertEqual(r1.status_code, 200)
            self.assertEqual(r1.data['totals']['shortage'], 7)
            tok = r1.data['token']
            r2 = self.c.post(self.URL, {'isrs': '261609', 'coverage': 3, 'token': tok}, format='json')
            self.assertEqual(r2.data['totals']['shortage'], 9)       # B 6 + C 3
            self.assertEqual(col.call_count, 1)                      # no second SOFTECH read
            self.c.post(self.URL, {'isrs': '261609', 'token': tok, 'refresh': True}, format='json')
            self.assertEqual(col.call_count, 2)                      # explicit refresh re-reads
            ex = self.c.post(self.URL + 'export/', {'isrs': '261609', 'coverage': 2,
                                                    'token': r2.data['token']}, format='json')
        self.assertEqual(ex.status_code, 200)
        self.assertIn('spreadsheetml', ex['Content-Type'])
        self.assertIn('isr_fulfillment_261609', ex['Content-Disposition'])
        wb = load_workbook(io.BytesIO(ex.content))
        self.assertEqual(wb['ملخص']['C4'].value, 2)


def _fake_items_plan(node, items):
    """build_items_plan without SOFTECH / catalog: one line per (code, qty)."""
    return [{'itemcode': c, 'item_name': c, 'itemqty': int(qty), 'itemqty_cfarma': float(qty),
             'nowqty': 0.0, 'itemsaleprice': 1.0, 'itemcostprice': 2.0, 'suppcode': '', 'gap': float(qty)}
            for c, qty in items]


@override_settings(SUPPLY_ERP_GROUPS='10,19,21,27')
class DonorReturnTests(TestCase):
    """«from branches» → one donor→HQ proposal per donor, tagged with the branch's ISR;
    never a second HQ→branch ISR; never twice for the same (ISR, donor)."""

    def setUp(self):
        cache.clear()
        p = mock.patch('apps.purchasing.isr_writer.build_items_plan', side_effect=_fake_items_plan)
        p.start()
        self.addCleanup(p.stop)

    def test_one_branch_to_hq_proposal_per_donor(self):
        from apps.purchasing.models import IsrPush
        res = F.create_donor_returns(_snap(), 1.5)
        self.assertEqual([(c['donor'], c['qty']) for c in res['created']], [('150', 10), ('130', 5), ('140', 2)])
        ps = IsrPush.objects.order_by('id')
        self.assertEqual(ps.count(), 3)                                   # no HQ→branch leg
        for p in ps:
            self.assertEqual((p.kind, p.dest_branchcode, p.status, p.origin_isr),
                             (IsrPush.KIND_BRANCH_TO_HQ, '100', IsrPush.STATUS_PROPOSED, 261609))
            self.assertIn('261609', p.notes)
        self.assertEqual(ps.get(branchcode='150').lines_snapshot[0]['itemcode'], 'A')

    def test_second_call_skips_live_and_cancel_frees(self):
        from apps.purchasing.models import IsrPush
        F.create_donor_returns(_snap(), 1.5)
        again = F.create_donor_returns(_snap(), 1.5)
        self.assertEqual(again['created'], [])
        self.assertEqual({s['donor'] for s in again['skipped']}, {'150', '130', '140'})
        IsrPush.objects.filter(branchcode='130').update(status=IsrPush.STATUS_CANCELLED)
        third = F.create_donor_returns(_snap(), 1.5)
        self.assertEqual([c['donor'] for c in third['created']], ['130'])

    def test_db_constraint_blocks_parallel_duplicate(self):
        from django.db import IntegrityError, transaction
        from apps.purchasing.models import IsrPush
        IsrPush.objects.create(branchcode='150', dest_branchcode='100', kind='branch_to_hq', origin_isr=261609)
        with self.assertRaises(IntegrityError), transaction.atomic():
            IsrPush.objects.create(branchcode='150', dest_branchcode='100', kind='branch_to_hq', origin_isr=261609)

    def test_api_needs_supply_write_and_reads_fresh(self):
        c = APIClient()
        c.force_authenticate(_user('ph1', role='pharmacist'))         # may view, may not create
        url = '/api/purchasing/isr/fulfilment/proposals/'
        self.assertEqual(c.post(url, {'isrs': '261609'}, format='json').status_code, 403)
        c.force_authenticate(_user('18', group=19))
        with mock.patch.object(F, 'collect', return_value=_snap()) as col:
            r = c.post(url, {'isrs': '261609', 'coverage': 1.5, 'isr': '261609', 'token': 'old'}, format='json')
        self.assertEqual(r.status_code, 201)
        self.assertEqual(col.call_count, 1)                           # never a cached snapshot
        self.assertTrue(r.data['plan']['isrs'][0]['user'])
        self.assertEqual(len(r.data['created']), 3)
        self.assertEqual(len(r.data['plan']['returns']['261609']), 3)
        self.assertEqual(c.post(url, {'isrs': '261609', 'isr': '777'}, format='json').status_code, 400)


def _conn(rows):
    cur = mock.MagicMock()
    cur.fetchall.return_value = rows
    conn = mock.MagicMock()
    conn.cursor.return_value = cur
    return conn


class UnreachableBranchFallbackTests(SimpleTestCase):
    """A branch server that is down (e.g. 170, JZ006 timeout on 2026-10-04) must not
    leave its stock blank: use server 100's copy, and say so in Arabic."""

    def _run(self, hq_copy):
        from types import SimpleNamespace as NS
        b130, b170 = NS(softech_branch_id='130'), NS(softech_branch_id='170')

        def node(b, target):
            if b.softech_branch_id == '170':
                raise RuntimeError('java.sql.SQLException: JZ006: Caught IOException')
            return _conn([('A', 2.0)])
        with mock.patch('apps.purchasing.rate_writer._read_conn_for_target', side_effect=node), \
                mock.patch('apps.purchasing.rate_writer.resolve_store', return_value='x'), \
                mock.patch('config.sybase.get_sybase_connection', return_value=_conn([('A', 9.0)])), \
                mock.patch.object(F, '_hq_copy_stock', **hq_copy), \
                self.assertLogs('apps.purchasing.isr_fulfillment', level='WARNING'):
            return F._live_stock([b130, b170], ['A'])

    def test_down_branch_uses_hq_copy_and_says_so(self):
        stock, down, fb = self._run({'return_value': ({'170': {'A': 4.0}}, {'170': '2026-10-03 08:19'})})
        self.assertEqual(stock, {'130': {'A': 2.0}, '100': {'A': 9.0}, '170': {'A': 4.0}})
        self.assertEqual((down, fb), ([], {'170': '2026-10-03 08:19'}))
        note = F.fallback_notes(fb)[0]
        self.assertIn('170', note)
        self.assertIn('2026-10-03 08:19', note)
        self.assertNotIn('JZ006', note)                                  # no raw Java errors

    def test_copy_also_failing_is_reported_in_arabic(self):
        stock, down, fb = self._run({'side_effect': RuntimeError('server 100 down')})
        self.assertNotIn('170', stock)
        self.assertEqual(fb, {})
        self.assertIn('سيرفر فرع 170', down[0])

    def test_copy_time_is_last_synced_branch_transaction(self):
        """'as of' = newest stktransm of that branch on server 100 — NOT stkbal.trans_time
        (a row stamp: 170's rows said 08:19 yesterday while its sales had synced to 00:38)."""
        import datetime as dt
        from types import SimpleNamespace as NS
        conn = _conn([('A', 4.0)])
        conn.cursor.return_value.fetchone.return_value = (dt.datetime(2026, 10, 4, 0, 38, 49),)
        with mock.patch('config.sybase.get_sybase_connection', return_value=conn), \
                mock.patch('apps.purchasing.rate_writer.resolve_store', return_value='170'):
            stock, as_of = F._hq_copy_stock([NS(softech_branch_id='170')], ['A'])
        self.assertEqual((stock, as_of), ({'170': {'A': 4.0}}, {'170': '2026-10-04 00:38'}))
        sqls = [c.args[0] for c in conn.cursor.return_value.execute.call_args_list]
        self.assertTrue(any('FROM stktransm' in s for s in sqls))
        self.assertFalse(any('trans_time FROM stkbal' in s for s in sqls))

    def test_plan_flags_requester_from_copy(self):
        s = _snap()
        s['fallback'] = {'160': '2026-10-03 08:17'}
        p = F.compute(s, 1.5)
        self.assertTrue(p['isrs'][0]['requester_from_hq_copy'])
        self.assertEqual(p['isrs'][0]['requester_copy_at'], '2026-10-03 08:17')
        self.assertIn('160', p['notices'][0])
        wb = F.build_workbook(s, 1.5)
        notes = [c.value for c in wb['الشرح']['A'] if c.value]
        self.assertTrue(any('آخر نسخة متزامنة' in v for v in notes))
        ws = wb['ISR 261609']
        self.assertIn('2026-10-03 08:17', ws['A3'].value)                # warning on the ISR sheet
        self.assertIn('الفرع الطالب 160 (نسخة الرئيسي)',
                      [ws.cell(4, c).value for c in range(1, ws.max_column + 1)])


class BuildItemsPlanStockTests(TestCase):
    """Regression: build_items_plan referenced an undefined _READ_CHUNK, so the source
    node's stock read always failed and every transfer / توزيعة leg recorded nowqty 0."""

    def test_reads_source_node_stock(self):
        from apps.branches.models import Branch
        from apps.catalog.models import Item
        from apps.purchasing import isr_writer
        Branch.objects.create(name='B150', softech_branch_id='150')
        Item.objects.create(softech_id='A', name='ITEM A', pack_price=10, cost_price=5)
        cur = mock.MagicMock()
        cur.fetchall.return_value = [('A', 7.0)]
        conn = mock.MagicMock()
        conn.cursor.return_value = cur
        with mock.patch('config.sybase.get_branch_connection', return_value=conn), \
                self.assertNoLogs('apps.purchasing.isr_writer', level='WARNING'):
            lines = isr_writer.build_items_plan('150', [('A', 3)])
        self.assertEqual((lines[0]['itemqty'], lines[0]['nowqty']), (3, 7.0))


class RecentPickerTests(TestCase):
    def test_recent_list_marks_ours_and_names_user(self):
        import datetime as dt
        from apps.purchasing.models import IsrPush
        from apps.users.models import ERPUser
        # synced SOFTECH users: username = usercode, user_id = SOFTECH userid (the name)
        ERPUser.objects.create(username='11', user_id='Store Keeper', full_name='', user_group='19')
        IsrPush.objects.create(branchcode='170', dest_branchcode='170', isrdocnumber=2617031)
        cur = mock.MagicMock()
        cur.fetchall.return_value = [
            (261609.0, '160', '100', dt.datetime(2026, 10, 2), 1, 36567.589, '11', 45, 345.0),
            (2617031.0, '170', '170', dt.datetime(2026, 9, 29), 1, 27120.7, '1509', 72, 267.0)]
        conn = mock.MagicMock()
        conn.cursor.return_value = cur
        with mock.patch('config.sybase.get_sybase_connection', return_value=conn):
            rows = F.recent_branch_isrs(500)                           # clamped to 90 days
        sql = cur.execute.call_args[0][0]
        self.assertIn('-90', sql)
        self.assertIn("branchcode <> '100'", sql)                      # varchar column
        self.assertEqual(rows[0], {'isr': '261609', 'branch': '160', 'for_branch': '100', 'date': '2026-10-02',
                                   'approved': True, 'value': 36567.589, 'user': '11',
                                   'user_name': 'Store Keeper', 'lines': 45, 'qty': 345.0, 'ours': False})
        self.assertTrue(rows[1]['ours'])
        self.assertEqual(rows[1]['user_name'], '')                     # unknown code → code only
        self.assertEqual(F.user_label('11', {'11': 'Store Keeper'}), '11 · Store Keeper')
        self.assertEqual(F.user_label('1509', {}), '1509')

    def test_isr_list_shows_code_and_softech_name(self):
        from apps.purchasing.models import IsrPush
        from apps.users.models import ERPUser
        ERPUser.objects.create(username='1509', user_id='BASSEM', user_group='10')
        u = _user('1509', role='admin')
        IsrPush.objects.create(branchcode='130', dest_branchcode='130', created_by=u)
        c = APIClient()
        c.force_authenticate(u)
        r = c.get('/api/purchasing/isr/pushes/')
        self.assertEqual(r.data['results'][0]['created_by'], '1509 · BASSEM')

    @override_settings(SUPPLY_ERP_GROUPS='10,19,21,27')
    def test_recent_api_access_and_outage(self):
        c = APIClient()
        url = '/api/purchasing/isr/fulfilment/recent/'
        c.force_authenticate(_user('1399', group=14))
        self.assertEqual(c.get(url).status_code, 403)
        c.force_authenticate(_user('18', group=19))
        with mock.patch.object(F, 'recent_branch_isrs', side_effect=RuntimeError('JZ006')), \
                self.assertLogs('apps.purchasing.isr_views', level='WARNING'):
            self.assertEqual(c.get(url).status_code, 503)
        with mock.patch.object(F, 'recent_branch_isrs', return_value=[]):
            r = c.get(url, {'days': 7})
        self.assertEqual((r.status_code, r.data['days']), (200, 7))


class FillCoverageTests(SimpleTestCase):
    """Owner 2026-10-05: a separate, deterministic fill coverage for the REQUESTING branch —
    show requested + recommended, ship per the chosen basis (default: the smaller)."""

    def test_recommended_qty_rule(self):
        self.assertEqual(F.recommended_qty(4.0, 2.0, 10.0, 1.5), 4.0)     # ceil(6 − 2)
        self.assertEqual(F.recommended_qty(4.0, 2.0, 1.0, 3.0), 10.0)     # may exceed the request
        self.assertEqual(F.recommended_qty(4.0, 9.0, 10.0, 1.5), 0.0)     # already above target
        self.assertEqual(F.recommended_qty(0.9, 0.3, 5.0, 1.5), 2.0)      # 1.35 − 0.3 = 1.05 → up → 2
        self.assertEqual(F.recommended_qty(0, 0, 7.0, 1.5), 7.0)          # no sales → as requested

    def test_basis_and_plan(self):
        r = F.compute(_snap(), 1.5, 1.5, 'min')['isrs'][0]
        a, b, c = r['lines']
        self.assertEqual((a['qty'], a['recommended'], a['plan_qty']), (10.0, 4.0, 4.0))
        self.assertEqual((a['from_branches'], a['shortage']), (4, 0.0))     # 150 gives 4, not 10
        self.assertTrue(b['no_sales'] and b['plan_qty'] == 15.0)             # no sales → requested
        self.assertEqual(r['totals']['requested'], 28.0)
        self.assertEqual((r['totals']['recommended'], r['totals']['planned']), (22.0, 22.0))
        req = F.compute(_snap(), 1.5, 1.5, 'requested')['isrs'][0]['lines'][0]
        self.assertEqual((req['plan_qty'], req['from_branches']), (10.0, 10))   # the original plan
        rec = F.compute(_snap(), 1.5, 3.0, 'recommended')['isrs'][0]['lines'][0]
        self.assertEqual((rec['recommended'], rec['plan_qty']), (10.0, 10.0))  # 4 × 3 − 2
        self.assertEqual(F.compute(_snap(), 1.5)['isrs'][0]['lines'][0]['plan_qty'], 10.0)  # fn default

    def test_donor_coverage_is_independent(self):
        p = F.compute(_snap(), 1.5, 1.5, 'min')
        q = F.compute(_snap(), 1.5, 2.5, 'min')
        donor = lambda plan: plan['isrs'][0]['lines'][0]['donors']
        self.assertEqual(donor(p)[0]['ceiling'], donor(q)[0]['ceiling'])    # 150 keeps the same
        self.assertEqual(q['isrs'][0]['lines'][0]['recommended'], 8.0)      # 4 × 2.5 − 2

    def test_settings_validation(self):
        self.assertEqual(F.plan_settings({}), {'months': 1.0, 'fill': 1.5, 'basis': 'min'})   # donors 1 month
        self.assertEqual(F.plan_settings({'coverage': 2, 'fill_coverage': '1', 'basis': 'recommended'}),
                         {'months': 2.0, 'fill': 1.0, 'basis': 'recommended'})
        for bad in ({'fill_coverage': 0.1}, {'coverage': 20}, {'basis': 'all'}, {'fill_coverage': 'x'}):
            with self.assertRaises(ValueError):
                F.plan_settings(bad)

    def test_workbook_has_the_three_inputs_and_plan_columns(self):
        wb = F.build_workbook(_snap(), 1.5, 2.0, 'min')
        s = wb['ملخص']
        self.assertEqual((s['C4'].value, s['C5'].value, s['C6'].value), (1.5, 2.0, 'الأقل من الاثنين'))
        self.assertTrue(any('C6' in str(dv.sqref) for dv in s.data_validations.dataValidation))
        ws = wb['ISR 261609']
        heads = [c.value for c in ws[ws.max_row - 3] if c.value] + \
                [ws.cell(r, c).value for r in range(1, 6) for c in range(1, 10)]
        self.assertIn('الكمية الموصى بها (تغطية الطلب)', heads)
        self.assertIn(F.plan_label_formula(), heads)          # header follows C6
        formulas = ' '.join(str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str))
        self.assertIn(F.FILL_REF, formulas)
        self.assertIn(F.BASIS_REF, formulas)


class CoverageNeedTests(SimpleTestCase):
    """Owner 2026-10-05: «الاحتياج» follows the coverage settings (requester: fill coverage,
    donors: donor coverage) — never the engine's own gap, never negative."""

    def test_rule(self):
        # owner's rule: months × rate − stock − in transit; positive → rounded UP (≥ 1 pack), else 0
        self.assertEqual(F.coverage_need(5.17, 0, 0.8), 5.0)       # 4.136 → 5
        self.assertEqual(F.coverage_need(0.92, 0, 1.0), 1.0)       # 0.92 → 1
        self.assertEqual(F.coverage_need(0.28, 0, 0.8), 1.0)       # slow mover, positive 0.224 → 1 pack
        self.assertEqual(F.coverage_need(0.28, 0.5, 0.8), 0.0)     # 0.224 − 0.5 < 0 → 0 (pack only if positive)
        self.assertEqual(F.coverage_need(16.7, 4, 0.8), 10.0)      # 13.36 − 4 = 9.36 → 10
        self.assertEqual(F.coverage_need(5, 4, 0.8), 0.0)          # exactly 0 → 0
        self.assertEqual(F.coverage_need(2.0, 20, 1.0), 0.0)       # well above → 0, not negative
        self.assertEqual(F.coverage_need(0, 3, 1.0), 0.0)          # doesn't sell → no need

    def test_plan_uses_settings_not_engine_gap(self):
        snap = _snap()
        snap['eng']['A|150'] = [2.0, -17.0]                         # engine says −17 (above its target)
        line = F.compute(snap, 1.0, 0.8, 'min')['isrs'][0]['lines'][0]
        d150 = next(d for d in line['donors'] if d['branch'] == '150')
        self.assertEqual(d150['need'], 0.0)                         # 20 on hand ≥ 2 × 1 month
        self.assertEqual(line['req']['need'], 2.0)                  # 160: 0.8 × 4 − 2 = 1.2 → up → 2
        self.assertGreaterEqual(min(d['need'] for d in line['donors']), 0)

    def test_workbook_need_columns_are_formulas(self):
        ws = F.build_workbook(_snap(), 1.0, 0.8, 'min')['ISR 261609']
        heads = {ws.cell(5, c).value: c for c in range(1, ws.max_column + 1)}
        self.assertIn('الاحتياج (حتى تغطية الطالب)', heads)
        self.assertNotIn('احتياج المحرك', heads)
        f = ws.cell(6, heads['الاحتياج (حتى تغطية الطالب)']).value
        self.assertTrue(f.startswith('=MAX(0,ROUNDUP(') and F.FILL_REF in f)


class InTransitTests(SimpleTestCase):
    """Owner 2026-10-05: count what is already on the way (doccode-125 transfers not yet
    received, issued within the engine's window — 14 days) so we never buy / move it twice."""

    def _snap(self, transit):
        s = _snap()
        s['transit'] = transit
        return s

    def test_requester_transit_lowers_recommended_and_need(self):
        base = F.compute(_snap(), 1.0, 1.5, 'min')['isrs'][0]['lines'][0]          # A @160
        self.assertEqual((base['recommended'], base['req']['need']), (4.0, 4.0))  # ceil(6 − 2)
        line = F.compute(self._snap({'160': {'A': 3.0}}), 1.0, 1.5, 'min')['isrs'][0]['lines'][0]
        self.assertEqual(line['req']['transit'], 3.0)
        self.assertEqual((line['recommended'], line['req']['need'], line['plan_qty']), (1.0, 1.0, 1.0))
        self.assertEqual(line['req']['after'], 2.0 + 3.0 + 1.0)                   # stock + transit + supplied
        self.assertEqual(line['shortage'] + line['from_branches'] + line['from_hq'], 1.0)

    def test_donor_transit_counts_for_its_need_but_never_moves(self):
        line = F.compute(self._snap({'150': {'A': 5.0}}), 1.0, 1.5, 'requested')['isrs'][0]['lines'][0]
        d150 = next(d for d in line['donors'] if d['branch'] == '150')
        self.assertEqual(d150['transit'], 5.0)
        self.assertEqual(d150['excess'], 18)                                       # 20 on shelf − 2, transit excluded
        self.assertEqual(line['network']['transit'], 5.0)

    def test_old_snapshot_without_transit_still_works(self):
        s = _snap()
        s.pop('transit', None)
        self.assertEqual(F.compute(s, 1.0, 1.5, 'min')['isrs'][0]['lines'][0]['req']['transit'], 0.0)

    def test_workbook_columns_and_formulas(self):
        ws = F.build_workbook(self._snap({'160': {'A': 3.0}}), 1.0, 1.5, 'min')['ISR 261609']
        heads = {ws.cell(5, c).value: c for c in range(1, ws.max_column + 1)}
        self.assertEqual(ws.cell(6, heads['بالطريق (لم يُستلم)']).value, 3.0)
        from openpyxl.utils import get_column_letter as L
        need = ws.cell(6, heads['الاحتياج (حتى تغطية الطالب)']).value
        self.assertIn(f'-{L(heads["بالطريق (لم يُستلم)"])}6', need)               # need nets transit
        rec = ws.cell(6, 5).value
        self.assertIn(f'{L(heads["الاحتياج (حتى تغطية الطالب)"])}6', rec)         # recommended = need
        self.assertIn('بالطريق للفروع', heads)


class InTransitWindowTests(TestCase):
    """The shared definition (apps/purchasing/in_transit.py): status in_transit, receiving
    branch known, issued within the window; summed per receiving branch."""

    def test_window_and_mapping(self):
        import datetime
        from django.utils import timezone
        from apps.branches.models import Branch
        from apps.purchasing.in_transit import by_branch_code
        from apps.transits.models import InTransitTransfer
        b160 = Branch.objects.create(softech_branch_id='160', name='Ramsis')
        b130 = Branch.objects.create(softech_branch_id='130', name='Nozha')
        today = timezone.localdate()
        mk = lambda doc, to, days, status='in_transit', items=None: InTransitTransfer.objects.create(
            erp_doc_number=doc, erp_supplying_branch_code='150', receiving_branch=to,
            issue_date=today - datetime.timedelta(days=days), transit_status=status,
            items_snapshot=items or [{'itemcode': 'A', 'qty': 2}])
        mk('1', b160, 1)
        mk('2', b160, 13, items=[{'itemcode': 'A', 'qty': 3}, {'itemcode': 'B', 'qty': 1}])
        mk('3', b160, 20)                                  # older than 14 days → stale, ignored
        mk('4', b130, 2, status='received')                # received → not in transit
        mk('5', b130, 2, items=[{'itemcode': 'Z', 'qty': 9}])  # item not asked for
        data, days = by_branch_code(['A', 'B'], max_age=14)
        self.assertEqual(days, 14)
        self.assertEqual(data, {'160': {'A': 5.0, 'B': 1.0}})
        self.assertEqual(by_branch_code(['A'], max_age=0)[0]['160']['A'], 7.0)   # 0 = no age limit


class RequesterCeilingTests(SimpleTestCase):
    """Owner 2026-10-05: the requesting branch's maximum uses ITS coverage (C5), and
    «يتجاوز الحد الأقصى؟» only when something is supplied and stock + transit + supplied > max."""

    def test_max_uses_fill_coverage(self):
        a = F.compute(_snap(), 1.5, 0.8, 'min')['isrs'][0]['lines'][0]          # A @160, rate 4
        self.assertEqual(a['req']['ceiling'], 3.2)                              # 4 × 0.8, not 4 × 1.5

    def test_flag_only_when_supplied(self):
        s = _snap()
        s['stock']['160']['A'] = 9.0                                            # already above max
        a = F.compute(s, 1.5, 0.8, 'min')['isrs'][0]['lines'][0]
        self.assertEqual(a['plan_qty'], 0.0)
        self.assertFalse(a['req']['over_ceiling'])                              # existing overstock ≠ flag
        r = F.compute(s, 1.5, 0.8, 'requested')['isrs'][0]['lines'][0]
        self.assertTrue(r['req']['over_ceiling'])                               # 9 + 10 supplied > 3.2

    def test_workbook(self):
        ws = F.build_workbook(_snap(), 1.5, 0.8, 'min')['ISR 261609']
        heads = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
        rc = heads.index('الحد الأقصى') + 1                                      # the requester's (first)
        self.assertIn(F.FILL_REF, ws.cell(6, rc).value)
        rx = heads.index('يتجاوز الحد الأقصى؟') + 1
        self.assertTrue(ws.cell(6, rx).value.startswith('=IF(AND(F6>0,'))      # F = supplied qty
