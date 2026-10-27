"""
«بيانات الصنف» in every /supply Excel (owner 2026-10-26) — apps/catalog/item_columns.py: one
definition, read live at export time (an item update shows in the next export), placed right after
the item's code + name; plus the three new tab exports (inbox / cases / rates).
"""
import io
from decimal import Decimal
from unittest import mock

from django.test import TestCase
from openpyxl import load_workbook

from apps.branches.models import Branch
from apps.catalog import item_columns as IC
from apps.catalog.models import Item
from apps.tests.factories import make_admin, make_item


def _rows(data, sheet=None):
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _item(code='70001', name='PANADOL'):
    it = make_item(name, code)
    Item.objects.filter(pk=it.pk).update(
        supplier_code='4570', supplier_name='CHEMIPHARM', producer_code='110', producer_name='AMOUN',
        medicine_type_name_ar='أدوية', shape_name_ar='أقراص', family_name_ar='مسكنات',
        store_classif_name='Med: Local', pack_price=Decimal('58.00'), cost_price=Decimal('43.50'))
    return Item.objects.select_related('category').get(pk=it.pk)


class BlockTests(TestCase):
    def test_values_in_order(self):
        cells = dict(zip(IC.ITEM_HEADS, IC.item_cells(_item())))
        self.assertEqual(cells['التصنيف العام'], 'أدوية')
        self.assertEqual(cells['المورد الرئيسي'], 'CHEMIPHARM')
        self.assertEqual(cells['الشركة المنتجة'], 'AMOUN')
        self.assertEqual((cells['سعر الجمهور'], cells['سعر التكلفة']), (58.0, 43.5))

    def test_without_prices_and_unknown_item(self):
        self.assertNotIn('سعر التكلفة', [c[0] for c in IC.columns(prices=False)])
        self.assertEqual(len(IC.item_cells(_item(), prices=False)), IC.N - 2)
        self.assertEqual(IC.item_cells(None), [None] * IC.N)


class DistributionLiveTests(TestCase):
    def _export(self):
        from apps.purchasing.supply_exports import distribution_excel
        sug = {'category': 'over_piled', 'itemcode': '70001', 'item_name': 'PANADOL', 'sells_at': 2,
               'of_branches': 5, 'age_days': 40, 'from_branch': '130', 'from_position': {},
               'to_branches': [{'code': '140', 'qty': 2}], 'network_avg': 3, 'reason': 'x'}
        with mock.patch('apps.finance.recon_labels.branch_label', side_effect=str):
            return _rows(distribution_excel([sug]))

    def test_block_after_code_and_name_and_follows_item_updates(self):
        _item()
        rows = self._export()
        self.assertEqual(rows[0][1:3], ['الكود', 'الصنف'])
        self.assertEqual(rows[0][3:3 + IC.N], IC.ITEM_HEADS)
        col = 3 + IC.ITEM_HEADS.index('المورد الرئيسي')
        self.assertEqual(rows[1][col], 'CHEMIPHARM')
        Item.objects.filter(softech_id='70001').update(supplier_name='MINAPHARM')   # item table updated
        self.assertEqual(self._export()[1][col], 'MINAPHARM')                          # → next export


class ComparisonAndPricesTests(TestCase):
    def test_supplier_comparison(self):
        from apps.supply.comparison import to_excel
        _item()
        board = {'days': 30, 'better_pct': 1, 'suppliers': [], 'rows': [{
            'item_code': '70001', 'item_name': 'PANADOL', 'public_price': 58, 'monthly_rate': 3, 'stock': 1,
            'in_transit': 0, 'need': 2, 'state': 'buy', 'best': None, 'allocation': [], 'cells': [],
            'bought_elsewhere': None, 'flags': []}]}
        rows = _rows(to_excel(board))
        self.assertEqual(rows[1][2:2 + IC.N], IC.ITEM_HEADS)
        self.assertEqual(rows[2][2 + IC.ITEM_HEADS.index('الشركة المنتجة')], 'AMOUN')

    def test_purchase_prices(self):
        from apps.supply.price_history import to_excel
        _item()
        board = {'days': 90, 'since': '2026-07-28', 'today': '2026-10-26', 'filters': None, 'stale_days': 90,
                 'better_pct': 1, 'suppliers': [], 'items': [{
                     'code': '70001', 'name': 'PANADOL', 'producer': 'AMOUN', 'public_price': 58,
                     'monthly_rate': 3, 'n_suppliers': 1, 'state_label': 'x', 'paid_pct': 10, 'last': None,
                     'best': None, 'best_warehouse': None, 'gap_pct': 0, 'saving_month': 0, 'suppliers': []}]}
        rows = _rows(to_excel(board))
        self.assertEqual(rows[0][2:2 + IC.N], IC.ITEM_HEADS)
        self.assertEqual(rows[1][2 + IC.ITEM_HEADS.index('سعر التكلفة')], 43.5)


class NewTabExportTests(TestCase):
    def setUp(self):
        _u, _p, self.client = make_admin()
        self.item = _item()
        self.branch = Branch.objects.create(softech_branch_id='130', name='ElNozha', is_active=True)

    def _check(self, r, first_col):
        self.assertEqual(r.status_code, 200, getattr(r, 'data', None))
        rows = _rows(r.content)
        self.assertEqual(rows[0][first_col:first_col + IC.N], IC.ITEM_HEADS)
        self.assertEqual(rows[1][first_col + IC.ITEM_HEADS.index('المورد الرئيسي')], 'CHEMIPHARM')

    def test_rates(self):
        from apps.purchasing.models import SalesRatePush, SalesRatePushLine
        p = SalesRatePush.objects.create(status='proposed')
        SalesRatePushLine.objects.create(push=p, itemcode='70001', item_name='PANADOL', branchcode='130',
                                         storecode='130', old_rate=0.0, new_rate=0.9, eligible=True)
        self._check(self.client.get(f'/api/purchasing/rates/pushes/{p.pk}/export/'), 2)

    def test_cases(self):
        from apps.supply.models import SupplyCase
        SupplyCase.objects.create(item=self.item, branch=self.branch)
        self._check(self.client.get('/api/supply/cases/export/', {'include_closed': '1'}), 2)

    def test_availability(self):
        from apps.supply.models import AvailabilityBatch, AvailabilityLine
        b = AvailabilityBatch.objects.create(source='whatsapp', supplier_name='X')
        ln = AvailabilityLine.objects.create(batch=b, raw_text='panadol 10', item=self.item, is_confirmed=True)
        fake = {'lines': [{'line_id': ln.pk, 'raw_text': 'panadol 10', 'item_id': self.item.pk,
                           'item_softech_id': '70001', 'item_name': 'PANADOL', 'state': 'buy',
                           'flags': ['urgent']}], 'scope': {'all': True}}
        with mock.patch('apps.supply.availability.analyze_batch', return_value=fake):
            self._check(self.client.get(f'/api/supply/availability/{b.pk}/export/'), 3)
