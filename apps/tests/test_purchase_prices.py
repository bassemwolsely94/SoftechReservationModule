"""
«مقارنة أسعار الشراء» (apps/supply/price_history.py) and the purchases refresh
(apps/procurement/refresh.py) — owner 2026-10-24.

Board: accredited suppliers compared on saved purchases — discount after bonus off the
public price then; distributors/manufacturers head-to-head, warehouses apart, excluded
classes never; one company's accounts merged; yardstick = our qty-weighted recent discount;
a stale best never drives «يوجد أرخص». Refresh: catch-up when due, never two at once,
literal date window. SOFTECH is never touched here.
"""
import datetime
from decimal import Decimal
from unittest import mock

from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.supply import price_history as H

TODAY = datetime.date(2026, 10, 24)


def _line(item, supplier, days_ago, *, qty=10, value=None, disc=0.20, bonus=0, public=100, n=[0]):
    """A paid purchase line: `qty` packs at public × (1 − disc), VAT 0."""
    from apps.procurement.models import PurchaseLine
    n[0] += 1
    value = value if value is not None else qty * public * (1 - disc)
    return PurchaseLine.objects.create(
        branch_code='100', supplier_code=supplier, doc_number=str(n[0]), doccode='10',
        doc_date=TODAY - datetime.timedelta(days=days_ago), item_code=item.softech_id, item=item,
        net_qty=Decimal(str(qty)), raw_qty=Decimal(str(qty)), net_value=Decimal(str(round(value, 3))),
        raw_value=Decimal(str(round(value, 3))), bonus_qty=Decimal(str(bonus)), vat_value=Decimal('0'),
        public_price=Decimal(str(public)))


class BoardTests(TestCase):
    def setUp(self):
        from apps.catalog.models import Item
        from apps.procurement.models import SupplierProfile
        for code, name, cl in (('D1', 'RAMCO', '60'), ('D2', 'IBN SINA', '60'), ('786', 'EGY DRUG-SHERIF', '60'),
                               ('4327', 'EGY DRUG-ZAYTOUN', '60'), ('W1', 'مخزن', '50'), ('V1', 'virtual', '80')):
            SupplierProfile.objects.create(supplier_code=code, supplier_name=name, classif_code=cl)
        self.a = Item.objects.create(softech_id='A1', name='ITEM A', pack_price=100, producer_name='P1',
                                     supplier_code='D1')
        self.b = Item.objects.create(softech_id='B1', name='ITEM B', pack_price=100)
        self.c = Item.objects.create(softech_id='C1', name='ITEM C', pack_price=120)
        p = mock.patch.object(H, '_rates', return_value={self.a.id: 30.0, self.b.id: 10.0})
        p.start()
        self.addCleanup(p.stop)

    def board(self, **kw):
        return {r['code']: r for r in H.build(today=TODAY, **kw)['items']}

    def test_cheaper_recent_supplier_against_what_we_really_paid(self):
        _line(self.a, 'D1', 5, qty=90, disc=0.20)           # most packs at 20 %
        _line(self.a, 'W1', 2, qty=1, disc=0.05)            # one emergency pack from a warehouse
        _line(self.a, 'D2', 10, qty=10, disc=0.26)          # IBN SINA 26 %
        _line(self.a, 'V1', 3, qty=5, disc=0.60)            # virtual / contract: never a price to beat
        r = self.board()['A1']
        self.assertEqual((r['last']['name'], r['last']['discount_pct']), ('مخزن', 5.0))
        self.assertEqual(r['paid_pct'], 20.45)                # (90×20 + 1×5 + 10×26) / 101 = 20.445
        self.assertEqual((r['best']['name'], r['best']['discount_pct']), ('IBN SINA', 26.0))
        self.assertEqual((r['state'], r['gap_pct']), ('cheaper', 5.55))
        self.assertEqual(r['saving_month'], round(5.55 / 100 * 100 * 30, 2))
        self.assertEqual((r['n_suppliers'], r['n_warehouses']), (2, 1))
        self.assertEqual(r['best_warehouse']['name'], 'مخزن')
        self.assertNotIn('virtual', [c['name'] for c in r['suppliers']])
        self.assertEqual(r['main_supplier'], 'D1')

    def test_bonus_lowers_cost_and_accounts_of_one_company_merge(self):
        _line(self.b, 'D1', 5, qty=10, disc=0.20)
        _line(self.b, '786', 40, qty=10, disc=0.20, bonus=0)
        _line(self.b, '4327', 8, qty=10, value=800, bonus=2)   # 800 / 12 units = 66.67 → 33.33 %
        r = self.board()['B1']
        egy = [c for c in r['suppliers'] if c['key'].startswith('g:')]
        self.assertEqual(len(egy), 1)                           # شريف + زيتون = one column
        self.assertEqual(egy[0]['buys'], 2)
        self.assertEqual(egy[0]['discount_pct'], 33.33)         # the latest account's buy, after bonus
        self.assertEqual(r['best']['key'], egy[0]['key'])

    def test_stale_best_never_drives_cheaper(self):
        _line(self.a, 'D1', 5, qty=10, disc=0.20)
        _line(self.a, 'D2', 120, qty=10, disc=0.35)             # better, but 4 months old
        r = self.board()['A1']
        self.assertEqual((r['state'], r['best']['name'], r['saving_month']), ('best', 'RAMCO', 0.0))
        self.assertEqual(r['old_best']['name'], 'IBN SINA')     # shown as «كان أرخص»
        self.assertTrue(next(c for c in r['suppliers'] if c['name'] == 'IBN SINA')['stale'])

    def test_single_supplier_items_only_on_request_and_old_public_price(self):
        _line(self.c, 'D1', 5, qty=4, disc=0.10, public=100)    # bought when the public price was 100
        self.assertNotIn('C1', self.board())
        r = self.board(include_single=True)['C1']
        self.assertEqual(r['state'], 'single')
        self.assertTrue(r['last']['old_price'])                 # 120 today
        self.assertEqual(r['suppliers'][0]['cost_now'], 108.0)  # 10 % off today's 120

    def test_supplier_summary_and_excel(self):
        _line(self.a, 'D1', 5, qty=90, disc=0.20)
        _line(self.a, 'D2', 10, qty=10, disc=0.26)
        b = H.build(today=TODAY)
        s = {x['name']: x for x in b['suppliers']}
        self.assertEqual((s['IBN SINA']['best_in'], s['IBN SINA']['cheaper_than_last']), (1, 1))
        self.assertEqual(b['totals']['cheaper'], 1)
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(H.to_excel(b)))
        self.assertEqual(wb.sheetnames, ['الأصناف', 'الموردين', 'الشرح'])
        self.assertEqual(wb['الأصناف']['A2'].value, 'A1')

    def test_item_history(self):
        _line(self.a, 'D1', 5, qty=10, disc=0.20)
        _line(self.a, 'V1', 3, qty=5, disc=0.60)
        h = H.item_history('A1', today=TODAY)
        self.assertEqual([(l['company'], l['discount_pct']) for l in h['lines']], [('RAMCO', 20.0)])
        self.assertIsNone(H.item_history('NOPE'))


class RefreshTests(TestCase):
    def _run(self, status, hours_ago, days=8):
        from apps.procurement.models import ProcurementEngineRun
        r = ProcurementEngineRun.objects.create(status=status, period_days=days)
        t = timezone.now() - datetime.timedelta(hours=hours_ago)
        ProcurementEngineRun.objects.filter(pk=r.pk).update(started_at=t, finished_at=t if status == 'success' else None)
        return r

    def test_catch_up_window_and_due(self):
        from apps.procurement import refresh
        self.assertEqual(refresh.catchup_days(), refresh.FIRST_DAYS)
        self.assertTrue(refresh.is_due())
        self._run('success', 24 * 6)                            # the machine slept for 6 days
        self.assertTrue(refresh.is_due())
        self.assertEqual(refresh.catchup_days(), 6 + refresh.MARGIN_DAYS)
        self._run('success', 2)
        self.assertFalse(refresh.is_due())
        self.assertEqual(refresh.catchup_days(), refresh.MIN_DAYS)

    def test_never_two_at_once_and_dead_runs_ignored(self):
        from apps.procurement import refresh
        self._run('running', 1)
        self.assertIsNotNone(refresh.running_run())
        with mock.patch('apps.procurement.engine.run_procurement_engine') as eng:
            self.assertIsNone(refresh.run_now('test'))
            self.assertFalse(refresh.start_background('test'))
        eng.assert_not_called()
        from apps.procurement.models import ProcurementEngineRun
        ProcurementEngineRun.objects.update(started_at=timezone.now() - datetime.timedelta(hours=5))
        self.assertIsNone(refresh.running_run())                # a dead process, not a running one

    def test_scheduled_tick_only_when_due(self):
        from apps.procurement import refresh
        self._run('success', 2)
        with mock.patch.object(refresh, 'run_now') as rn:
            refresh.scheduled_tick()
        rn.assert_not_called()
        self._run('success', 30)
        from apps.procurement.models import ProcurementEngineRun
        ProcurementEngineRun.objects.filter(finished_at__gte=timezone.now() - datetime.timedelta(hours=3)).delete()
        with mock.patch.object(refresh, 'run_now') as rn:
            refresh.scheduled_tick()
        rn.assert_called_once_with('scheduler')


class QueryWindowTests(SimpleTestCase):
    def test_literal_date_window_on_both_tables(self):
        from apps.procurement.engine import purchases_query
        q = purchases_query(8, today=TODAY)
        self.assertNotIn('?', q)
        self.assertIn("sm.docdate >= '2026-10-16' AND sm.docdate < '2026-10-25'", q)
        self.assertIn("st.docdate >= '2026-10-16' AND st.docdate < '2026-10-25'", q)


class ApiTests(TestCase):
    def setUp(self):
        from .factories import make_user
        self.admin, _, _ = make_user('pp_admin', role='admin')
        self.nobody, _, _ = make_user('pp_sales', role='salesperson')
        self.c = APIClient()

    def test_board_access_and_excel(self):
        self.c.force_authenticate(self.nobody)
        self.assertEqual(self.c.get('/api/supply/purchase-prices/').status_code, 403)
        self.c.force_authenticate(self.admin)
        # Django only gzips bodies ≥ 200 bytes, and the board is cached: pad the (empty)
        # test board to a realistic size and bypass the cache, so the result never
        # depends on what the shared cache of the machine running the tests holds.
        from django.core.cache import cache
        real_build = H.build
        padded = lambda **kw: {**real_build(**kw), 'items': [{'code': str(i)} for i in range(50)]}
        with mock.patch.object(cache, 'get', return_value=None), mock.patch.object(cache, 'set'), \
             mock.patch.object(H, 'build', side_effect=padded):
            r = self.c.get('/api/supply/purchase-prices/', {'days': 90}, HTTP_ACCEPT_ENCODING='gzip')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get('Content-Encoding'), 'gzip')
        x = self.c.get('/api/supply/purchase-prices/', {'format': 'xlsx'})
        self.assertIn('spreadsheetml', x['Content-Type'])
        self.assertEqual(self.c.get('/api/supply/purchase-prices/item/NOPE/').status_code, 404)

    def test_refresh_button_and_freshness(self):
        self.c.force_authenticate(self.nobody)
        self.assertEqual(self.c.post('/api/supply/freshness/refresh-purchases/').status_code, 403)
        self.c.force_authenticate(self.admin)
        with mock.patch('apps.procurement.refresh.start_background', return_value=True) as sb:
            r = self.c.post('/api/supply/freshness/refresh-purchases/')
        self.assertEqual((r.status_code, sb.call_count), (202, 1))
        f = self.c.get('/api/supply/freshness/').data
        self.assertIn('purchases', f)
        self.assertTrue(f['can']['refresh_purchases'])
        self.assertFalse(f['purchases_running'])
