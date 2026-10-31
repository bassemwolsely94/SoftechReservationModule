"""
apps/tests/test_receipt_check.py

P-31 (docs/product-review/06_SPEC_P-31_unordered_delivery_detector.md, revised 2026-10-31): the order file is the
demand pivot (`Total.الفرق-مطلوب/راكد`). Each pivot download is kept as a need sheet; doc-10 lines of the pilot
suppliers are flagged `unordered` (on no sheet in the window) or `over` (above the sheet's gap). PostgreSQL only —
nothing reads or writes SOFTECH.
"""
import datetime as dt
import io
import tempfile
from decimal import Decimal
from unittest import mock

from django.core.management import CommandError, call_command
from django.test import SimpleTestCase, TestCase
from django.utils import timezone

from apps.tests.factories import make_admin, make_item, make_user

TODAY = timezone.localdate()
GAP_COL = 'Total.الفرق-مطلوب/راكد'


def day(n):
    return TODAY + dt.timedelta(days=n)


def at(n):
    return timezone.make_aware(dt.datetime.combine(day(n), dt.time(10, 0)))


def _snap(lines, n=-5, source='import'):
    from apps.purchasing.models import PivotSheetSnapshot
    s = PivotSheetSnapshot.objects.create(source=source, lines=lines, line_count=len(lines))
    PivotSheetSnapshot.objects.filter(pk=s.pk).update(exported_at=at(n))
    s.refresh_from_db()
    return s


def _line(item, qty, n=-1, supplier='565', doc='1', doccode='10', foc=False, value=None, branch='100'):
    from apps.procurement.models import PurchaseLine
    return PurchaseLine.objects.create(branch_code=branch, supplier_code=supplier, doc_number=doc, doc_date=day(n),
                                       item_code=item, doccode=doccode, is_return=(doccode == '120'), is_foc=foc,
                                       net_qty=Decimal(qty), raw_qty=Decimal(qty),
                                       raw_value=Decimal(value if value is not None else qty * 10))


class SnapshotTests(TestCase):

    def test_same_sheet_twice_the_same_day_is_kept_once(self):
        from apps.purchasing.models import PivotSheetSnapshot
        from apps.purchasing.pivot_sheets import record_snapshot
        lines = [{'item_code': 'A', 'total_gap': 4.0, 'supplier_code': '565'}]
        record_snapshot(lines, staff=None, run=None, params={})
        record_snapshot(lines, staff=None, run=None, params={})
        self.assertEqual(PivotSheetSnapshot.objects.count(), 1)
        record_snapshot(lines + [{'item_code': 'B', 'total_gap': 1.0, 'supplier_code': ''}], staff=None, run=None,
                        params={})
        self.assertEqual(PivotSheetSnapshot.objects.count(), 2)

    def test_pivot_download_keeps_exactly_the_positive_total_rows(self):
        from openpyxl import load_workbook
        from apps.branches.models import Branch
        from apps.purchasing.models import DemandCalculationRun, ItemDemandMetrics, PivotSheetSnapshot
        run = DemandCalculationRun.objects.create(status='success')
        br = Branch.objects.create(softech_branch_id='130', name='ElNozha', is_active=True)
        need, rest = make_item('NEED', 'P1'), make_item('REST', 'P2')
        need.supplier_code = '565'
        need.save(update_fields=['supplier_code'])
        ItemDemandMetrics.objects.create(run=run, item=need, branch=br, calc_date=TODAY, monthly_avg=10, gap=4)
        ItemDemandMetrics.objects.create(run=run, item=rest, branch=br, calc_date=TODAY, monthly_avg=1, gap=-3)
        _, staff, client = make_admin('piv_admin')
        r = client.get('/api/purchasing/export/?view=pivot')
        self.assertEqual(r.status_code, 200)
        rows = [list(x) for x in load_workbook(io.BytesIO(r.content), read_only=True).active.iter_rows(values_only=True)]
        hdr = rows[1]
        in_file = {x[hdr.index('itemcode')]: x[hdr.index(GAP_COL)] for x in rows[2:]}
        s = PivotSheetSnapshot.objects.get()
        self.assertEqual((s.source, s.exported_by_id, s.run_id), ('download', staff.pk, run.pk))
        self.assertEqual({l['item_code']: l['total_gap'] for l in s.lines},
                         {c: float(v) for c, v in in_file.items() if v and v > 0})
        self.assertEqual(s.lines[0]['supplier_code'], '565')
        client.get('/api/purchasing/export/?view=pivot')
        self.assertEqual(PivotSheetSnapshot.objects.count(), 1)

    def test_a_recording_failure_still_returns_the_file(self):
        from apps.branches.models import Branch
        from apps.purchasing.models import DemandCalculationRun, ItemDemandMetrics
        run = DemandCalculationRun.objects.create(status='success')
        br = Branch.objects.create(softech_branch_id='130', name='ElNozha', is_active=True)
        ItemDemandMetrics.objects.create(run=run, item=make_item('X', 'P3'), branch=br, calc_date=TODAY,
                                         monthly_avg=10, gap=4)
        _, _, client = make_admin('piv_admin2')
        with mock.patch('apps.purchasing.pivot_sheets.record_snapshot', side_effect=RuntimeError('db down')):
            self.assertEqual(client.get('/api/purchasing/export/?view=pivot').status_code, 200)


class ImportCommandTests(TestCase):

    def _file(self, headers, rows):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['بيانات الصنف'] * len(headers))
        ws.append(headers)
        for r in rows:
            ws.append(r)
        f = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(f.name)
        return f.name

    def test_imports_the_positive_total_gap_rows(self):
        from apps.purchasing.models import PivotSheetSnapshot
        path = self._file(['Items_.كود-المورد', 'itemcode', 'Items_.itemname', GAP_COL],
                          [['565', 'A', 'x', 3], ['30', 'B', 'y', -2], ['', 'C', 'z', 0.5], [None, 'D', 'w', None]])
        call_command('import_pivot_snapshot', path, '--exported-at', '2026-10-19', stdout=io.StringIO())
        s = PivotSheetSnapshot.objects.get()
        self.assertEqual(s.source, 'import')
        self.assertEqual(s.exported_at.date(), dt.date(2026, 10, 19))
        self.assertEqual(s.lines, [{'item_code': 'A', 'total_gap': 3.0, 'supplier_code': '565'},
                                   {'item_code': 'C', 'total_gap': 0.5, 'supplier_code': ''}])

    def test_refuses_a_file_without_the_total_column(self):
        path = self._file(['itemcode', 'Items_.كود-المورد'], [['A', '565']])
        with self.assertRaises(CommandError):
            call_command('import_pivot_snapshot', path, '--exported-at', '2026-10-19', stdout=io.StringIO())


class CheckTests(TestCase):

    def setUp(self):
        _snap([{'item_code': 'ON', 'total_gap': 5.0, 'supplier_code': '30'}], n=-6)
        self.run = lambda: __import__('apps.procurement.receipt_check', fromlist=['x']).run_check(today=TODAY)

    def flags(self):
        from apps.procurement.models import ReceiptCheck
        return {(r.item_code, r.doc_number): r.flag for r in ReceiptCheck.objects.all()}

    def test_unordered_and_on_sheet(self):
        _line('OFF', 2, doc='1')
        _line('ON', 5, doc='2')
        self.run()
        self.assertEqual(self.flags(), {('OFF', '1'): 'unordered'})

    def test_over_the_sheet_gap_by_a_whole_pack(self):
        _line('ON', 4, doc='1', n=-3)
        _line('ON', 2, doc='2', n=-2)                         # 6 received vs 5 asked → over
        self.run()
        self.assertEqual(self.flags(), {('ON', '2'): 'over'})

    def test_not_judged(self):
        _line('OFF', 2, n=-10, doc='1')                       # before the first sheet
        _line('OFF', 2, supplier='999', doc='2')              # not a pilot supplier
        _line('OFF', 2, foc=True, doc='3')                    # a free (bonus) line
        self.run()
        self.assertEqual(self.flags(), {})

    def test_rerun_keeps_a_decision(self):
        from apps.procurement.models import ReceiptCheck
        _line('OFF', 2, doc='1')
        self.run()
        ReceiptCheck.objects.update(status='accepted', decision_reason='طلب هاتفي')
        self.run()
        r = ReceiptCheck.objects.get()
        self.assertEqual((r.status, r.decision_reason, r.flag), ('accepted', 'طلب هاتفي', 'unordered'))

    def test_return_closes_the_row_once(self):
        from apps.procurement.models import ReceiptCheck
        _line('OFF', 2, doc='1', n=-3)
        self.run()
        ReceiptCheck.objects.update(status='return_requested')
        _line('OFF', -2, doc='9', doccode='120', n=-1, value=-20)
        self.run()
        r = ReceiptCheck.objects.get()
        self.assertEqual((r.status, r.returned_value), ('returned', Decimal('20')))
        self.assertEqual(len(r.returned_refs), 1)


class ApiTests(TestCase):

    URL = '/api/supply/receipts/'

    def setUp(self):
        from apps.procurement.models import ReceiptCheck
        self.row = ReceiptCheck.objects.create(branch_code='100', supplier_code='565', doc_number='1', doc_date=day(-1),
                                               item_code='OFF', qty=Decimal('2'), value=Decimal('20'), flag='unordered')
        _, self.admin_staff, self.admin = make_admin('rc_admin')

    def test_list_and_decide(self):
        r = self.admin.get(self.URL)
        self.assertEqual(r.status_code, 200)
        self.assertEqual([x['id'] for x in r.data['rows']], [self.row.pk])
        r = self.admin.post(f'{self.URL}{self.row.pk}/decide/', {'action': 'accept', 'reason': 'طلب هاتفي'},
                            format='json')
        self.assertEqual(r.status_code, 200)
        self.row.refresh_from_db()
        self.assertEqual((self.row.status, self.row.decided_by_id), ('accepted', self.admin_staff.pk))
        self.assertEqual(self.admin.post(f'{self.URL}{self.row.pk}/decide/', {'action': 'accept'},
                                         format='json').status_code, 400)          # reason required

    def test_return_request_is_audited(self):
        from apps.audit.models import AuditLog
        self.admin.post(f'{self.URL}{self.row.pk}/decide/', {'action': 'return'}, format='json')
        self.row.refresh_from_db()
        self.assertEqual(self.row.status, 'return_requested')
        self.assertTrue(AuditLog.objects.filter(action='receipt_check_decided').exists())

    def test_without_supply_rights(self):
        from apps.tests.security.fixtures import seed_code_matrix
        from core.middleware import module_access
        seed_code_matrix()
        module_access.clear_cache()
        self.addCleanup(module_access.clear_cache)
        _, _, sales = make_user('rc_sales', role='salesperson')
        self.assertEqual(sales.post(f'{self.URL}{self.row.pk}/decide/', {'action': 'return'},
                                    format='json').status_code, 403)


class DigestTests(TestCase):

    def test_one_digest_a_day_to_purchasing(self):
        from apps.notifications.models import Notification
        from apps.procurement.models import ReceiptCheck
        from apps.procurement.receipt_check import send_digest
        _, buyer, _ = make_user('rc_buyer', role='purchasing')
        make_admin('rc_adm3')
        ReceiptCheck.objects.create(branch_code='100', supplier_code='565', doc_number='1', doc_date=day(-2),
                                    item_code='OFF', qty=Decimal('2'), value=Decimal('20'), flag='unordered')
        self.assertEqual(send_digest(today=TODAY), 1)
        self.assertEqual(send_digest(today=TODAY), 0)
        n = Notification.objects.get()
        self.assertEqual((n.recipient_id, n.notification_type), (buyer.pk, 'receipt_check_digest'))


class SchedulerTests(SimpleTestCase):

    def test_daily_job_registered(self):
        from apps.sync import tasks
        from .test_scheduler_jobs import _FakeScheduler
        with mock.patch.object(tasks, 'BackgroundScheduler', _FakeScheduler), \
             mock.patch.object(tasks, '_patch_apscheduler_for_python312', lambda: None), \
             mock.patch.object(tasks, '_scheduler', None):
            tasks.start_scheduler()
            ids = {j[1] for j in tasks._scheduler.jobs}
        self.assertIn('receipt_check_daily', ids)
