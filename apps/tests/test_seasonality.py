"""
Seasonality on «لوحة المشتريات» (owner 2026-10-25) — apps/purchasing/seasonality.py.
Bulk invoices kept apart, the index = ratio to a centred 12-month moving average (so growth is
not a «season»), not enough history → no index, the remarks column, and the refresh endpoints.
"""
from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

from django.test import SimpleTestCase, TestCase
from django.utils import timezone

from apps.purchasing import seasonality as S
from apps.purchasing.models import ItemMonthlySales, ItemSeasonality, SeasonalityRun
from apps.tests.factories import make_admin, make_item, make_salesperson

TODAY = date(2026, 10, 26)
MONTHS = S.history_months(TODAY)


def _series(fn):
    return [float(fn(m, i)) for i, m in enumerate(MONTHS)]


class HistoryMonthsTests(SimpleTestCase):
    def test_complete_months_only_oldest_first(self):
        self.assertEqual(len(MONTHS), S.MONTHS_BACK)
        self.assertEqual(MONTHS[-1], date(2026, 9, 1))           # the running month is never used
        self.assertEqual(MONTHS[0], date(2022, 10, 1))
        self.assertEqual(MONTHS, sorted(MONTHS))


class BulkSplitTests(SimpleTestCase):
    def test_bulk_invoice_kept_apart(self):
        s = S.split_invoices([2, 3, 1, 40], threshold=10)
        self.assertEqual(s, {'qty': 46, 'regular': 6, 'bulk': 40, 'invoices': 4, 'bulk_invoices': 1})

    def test_threshold_is_inclusive(self):
        self.assertEqual(S.split_invoices([10], threshold=10)['bulk'], 10)


class ComputeItemTests(SimpleTestCase):
    def test_summer_item_peaks_in_july(self):
        prof = S.compute_item(_series(lambda m, i: 60 if m.month == 7 else 20), MONTHS)
        self.assertIsNotNone(prof)
        self.assertGreater(prof[7]['index'], 2.0)
        self.assertLess(prof[1]['index'], 1.0)
        self.assertTrue(prof[7]['stable'])
        self.assertEqual(prof[7]['years'], 3)
        self.assertAlmostEqual(sum(p['index'] for p in prof.values()) / 12, 1.0, delta=0.02)

    def test_growth_is_not_a_season(self):
        prof = S.compute_item(_series(lambda m, i: 10 + i), MONTHS)
        self.assertIsNotNone(prof)
        for m in range(1, 13):
            self.assertAlmostEqual(prof[m]['index'], 1.0, delta=0.1, msg=m)

    def test_not_enough_history(self):
        self.assertIsNone(S.compute_item(_series(lambda m, i: 20 if i >= 30 else 0), MONTHS))   # 18 months sold
        self.assertIsNone(S.compute_item(_series(lambda m, i: 2), MONTHS))                      # below 4 / month
        self.assertIsNone(S.compute_item([20.0] * 24, MONTHS[-24:]))                            # two years only

    def test_index_clamped(self):
        prof = S.compute_item(_series(lambda m, i: 400 if m.month == 12 else 5), MONTHS)
        self.assertEqual(prof[12]['index'], S.INDEX_MAX)

    def test_years_that_disagree_are_unstable(self):
        # March: ×3 one year, normal the others
        prof = S.compute_item(_series(lambda m, i: 60 if (m.month == 3 and m.year == 2024) else 20), MONTHS)
        self.assertFalse(prof[3]['stable'])
        self.assertTrue(prof[4]['stable'])

    def test_bulk_months_listed(self):
        prof = S.compute_item(_series(lambda m, i: 20), MONTHS,
                              bulk_by_month={5: [(date(2025, 5, 1), 80.0), (date(2024, 5, 1), 0.0)]})
        self.assertEqual(prof[5]['bulk'], [(2025, 80.0)])


class NoteTests(SimpleTestCase):
    def test_no_data(self):
        self.assertIn('لا يمكن تقدير الموسمية', S.note_for(7, None))

    def test_strong_month_with_years_bulk_and_next(self):
        row = {'index': 1.4, 'years': 3, 'stable': False, 'bulk': [(2025, 80.0)],
               'ratios': [{'year': 2023, 'ratio': 1.3}, {'year': 2024, 'ratio': 1.5}, {'year': 2025, 'ratio': 1.4}]}
        note = S.note_for(7, row, {'index': 0.9})
        self.assertIn('يوليو أعلى من المتوسط 40%', note)
        self.assertIn('3 سنوات', note)
        self.assertIn('غير مستقر', note)
        self.assertIn('80 عبوة (يوليو 2025)', note)
        self.assertIn('الشهر التالي ×0.90', note)

    def test_capped_says_so(self):
        row = {'index': S.INDEX_MAX, 'years': 2, 'stable': True, 'bulk': [],
               'ratios': [{'year': 2024, 'ratio': 3.0}, {'year': 2025, 'ratio': 3.0}]}
        self.assertIn('محدود عند', S.note_for(12, row))


class ComputeAllTests(TestCase):
    def _store(self, item, fn):
        ItemMonthlySales.objects.bulk_create([
            ItemMonthlySales(item_code=item.softech_id, item=item, month=m, qty=Decimal(fn(m)),
                             regular_qty=Decimal(fn(m))) for m in MONTHS])

    def test_one_row_per_calendar_month(self):
        item = make_item('SUNBLOCK', 'SB001')
        item.is_stockable, item.pack_price = True, Decimal('100')
        item.save()
        self._store(item, lambda m: 60 if m.month == 7 else 20)
        self.assertEqual(S.compute_all(TODAY), 1)
        rows = {r.month: r for r in ItemSeasonality.objects.filter(item=item)}
        self.assertEqual(len(rows), 12)
        self.assertGreater(rows[7].index, Decimal('2'))
        self.assertIn('يوليو', rows[7].note)

    def test_service_codes_get_no_season(self):
        item = make_item('فرق سعر', 'SV001')
        item.is_stockable, item.pack_price = False, Decimal('0')
        item.save()
        self._store(item, lambda m: 20)
        self.assertEqual(S.compute_all(TODAY), 0)
        self.assertFalse(ItemSeasonality.objects.exists())


class DueAndRunTests(TestCase):
    def test_due_once_a_month_closes(self):
        self.assertTrue(S.is_due(TODAY))
        SeasonalityRun.objects.create(status='success', finished_at=timezone.now())
        self.assertFalse(S.is_due(timezone.localdate()))
        self.assertTrue(S.is_due(timezone.localdate() + timedelta(days=40)))

    def test_stale_running_run_does_not_block(self):
        run = SeasonalityRun.objects.create(status='running')
        self.assertEqual(S.running_run(), run)
        SeasonalityRun.objects.filter(pk=run.pk).update(started_at=timezone.now() - timedelta(hours=S.STALE_RUNNING_HOURS + 1))
        self.assertIsNone(S.running_run())

    def test_months_to_pull_missing_plus_overlap(self):
        ItemMonthlySales.objects.bulk_create([ItemMonthlySales(item_code='X1', month=m) for m in MONTHS])
        self.assertEqual(S.months_to_pull(TODAY), MONTHS[-S.REFRESH_OVERLAP:])


class EndpointTests(TestCase):
    def test_status_shape(self):
        _u, _p, client = make_admin()
        r = client.get('/api/purchasing/seasonality/')
        self.assertEqual(r.status_code, 200)
        for k in ('at', 'running', 'can_refresh', 'months_back'):
            self.assertIn(k, r.data)
        self.assertTrue(r.data['can_refresh'])

    def test_refresh_is_server_side_gated(self):
        _u, _p, client = make_salesperson()          # pharmacist / purchasing may run the engine
        with mock.patch.object(S, 'start_background') as start:
            r = client.post('/api/purchasing/seasonality/refresh/')
        self.assertEqual(r.status_code, 403)
        start.assert_not_called()

    def test_refresh_starts_in_background(self):
        _u, _p, client = make_admin()
        with mock.patch.object(S, 'start_background', return_value=True) as start:
            r = client.post('/api/purchasing/seasonality/refresh/')
        self.assertEqual(r.status_code, 202)
        start.assert_called_once()


class PivotExportColumnsTests(TestCase):
    """The pivot sheet carries this month's index, stable flag and note — info only."""

    def _sheet(self, client):
        import io
        from openpyxl import load_workbook
        r = client.get('/api/purchasing/export/?view=pivot')
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content), read_only=True).active
        return [list(row) for row in ws.iter_rows(values_only=True)]

    def test_season_columns_after_item_info(self):
        from apps.branches.models import Branch
        from apps.purchasing.models import DemandCalculationRun, ItemDemandMetrics
        month = timezone.localdate().month
        nxt = month % 12 + 1
        cur, nn = S.AR_MONTHS[month], S.AR_MONTHS[nxt]
        run = DemandCalculationRun.objects.create(status='success')
        br = Branch.objects.create(softech_branch_id='130', name='ElNozha', is_active=True)
        hot, flat, none = (make_item('HOT', 'S1'), make_item('FLAT', 'S2'), make_item('NONE', 'S3'))
        for it in (hot, flat, none):
            ItemDemandMetrics.objects.create(run=run, item=it, branch=br, calc_date=TODAY,
                                             monthly_avg=10, gap=4)
        ItemSeasonality.objects.create(item=hot, month=month, index=Decimal('1.300'), stable=True, note='ذروة')
        ItemSeasonality.objects.create(item=hot, month=nxt, index=Decimal('0.900'), stable=False)
        ItemSeasonality.objects.create(item=flat, month=month, index=Decimal('0.800'), stable=False)
        _u, _p, client = make_admin()
        rows = self._sheet(client)
        hdr = rows[1]
        col = hdr.index(f'موسمية-{cur}')
        self.assertEqual(hdr[col - 1], 'بونص زائد عن 6 شهور')        # right after the item info
        self.assertEqual(hdr[col:col + 5], [f'موسمية-{cur}', f'ثبات-{cur}', f'موسمية-{nn}', f'ثبات-{nn}',
                                            'ملاحظات-الموسمية'])
        self.assertTrue(hdr[col + 5].startswith('100.') or hdr[col + 5].startswith('130.'))  # then branches
        self.assertEqual(rows[0][col], 'الموسمية')
        by_name = {r[13]: r for r in rows[2:]}
        self.assertEqual(by_name['HOT'][col:col + 5], [1.3, 'مستقر', 0.9, '⚠ غير مستقر', 'ذروة'])
        self.assertEqual(by_name['FLAT'][col:col + 2], [0.8, '⚠ غير مستقر'])
        self.assertEqual(by_name['FLAT'][col + 2], None)                  # no next-month index
        self.assertEqual(by_name['NONE'][col], None)
        self.assertIn(by_name['NONE'][col + 1], (None, ''))
