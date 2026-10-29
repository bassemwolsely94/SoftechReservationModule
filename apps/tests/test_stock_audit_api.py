"""
apps/tests/test_stock_audit_api.py

Doc 32 Phase 3 — /api/stock-audit/: who sees what (management vs NAMED branch managers),
review moves, thresholds, access grants, export, refresh. Permissions are server-side.
"""
import datetime as _dt
from decimal import Decimal
from unittest.mock import patch

from django.test import TestCase
from django.utils import timezone

from apps.audit.models import AuditLog
from apps.stock_audit.models import (StockAuditAccess, StockAuditFinding, StockAuditRuleOverride,
                                     StockLedgerLine)
from apps.tests.factories import make_branch, make_branch2, make_user

TODAY = timezone.localdate()


def finding(branchcode, rule='sale_below_cost', status='open', severity='warning', value='100', user='166', n=1):
    return StockAuditFinding.objects.create(
        rule_code=rule, dedup_key=f'{branchcode}|{rule}|{n}', severity=severity, status=status,
        branchcode=branchcode, itemcode='100001', doccode='115', docnumber=n, docdate=TODAY,
        occurred_at=timezone.now(), softech_user=user, value_egp=Decimal(value),
        message_ar='رسالة', message_en='message', last_seen=timezone.now())


class StockAuditApiBase(TestCase):

    def setUp(self):
        self.b1 = make_branch('فرع أ', 'B01')
        self.b2 = make_branch2('فرع ب', 'B02')
        self.f1 = finding('B01', n=1)
        self.f2 = finding('B02', n=2, value='300', severity='critical')
        _, self.admin_p, self.admin = make_user('sa_admin', role='admin', branch=self.b1)
        _, self.sup_p, self.sup = make_user('sa_sup', role='supervisor', branch=self.b1)
        _, self.qm_p, self.qm = make_user('sa_qm', role='quality_manager', branch=self.b1)
        _, self.ph_p, self.ph = make_user('sa_ph', role='pharmacist', branch=self.b1)
        _, self.view_p, self.viewer = make_user('sa_view', role='viewer', branch=self.b1)

    def grant(self, profile=None):
        return StockAuditAccess.objects.create(staff=profile or self.ph_p, granted_by=self.admin_p)


class AccessTests(StockAuditApiBase):

    def test_management_sees_every_branch(self):
        r = self.sup.get('/api/stock-audit/findings/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual({x['branchcode'] for x in r.json()['results']}, {'B01', 'B02'})
        self.assertEqual(r.json()['results'][0]['branch_name']['ar'] in ('فرع أ', 'فرع ب'), True)

    def test_pharmacist_without_grant_and_viewer_see_nothing(self):
        for c in (self.ph, self.viewer):
            self.assertEqual(c.get('/api/stock-audit/findings/').status_code, 403)
            self.assertEqual(c.get('/api/stock-audit/summary/').status_code, 403)
            self.assertFalse(c.get('/api/stock-audit/me/').json()['can_view'])

    def test_named_branch_manager_sees_only_own_branch(self):
        self.grant()
        r = self.ph.get('/api/stock-audit/findings/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual([x['branchcode'] for x in r.json()['results']], ['B01'])
        self.assertEqual(self.ph.get(f'/api/stock-audit/findings/{self.f2.pk}/').status_code, 404)
        me = self.ph.get('/api/stock-audit/me/').json()
        self.assertEqual((me['branch_manager'], me['branches'], me['can_edit_thresholds']), (True, ['B01'], False))
        s = self.ph.get('/api/stock-audit/summary/').json()
        self.assertEqual([b['branchcode'] for b in s['branches']], ['B01'])

    def test_sidebar_flag_on_auth_me(self):
        self.assertFalse(self.ph.get('/api/auth/me/').json()['can_stock_audit'])
        self.grant()
        self.assertTrue(self.ph.get('/api/auth/me/').json()['can_stock_audit'])
        self.assertTrue(self.qm.get('/api/auth/me/').json()['can_stock_audit'])

    def test_branch_manager_cannot_edit_thresholds_or_refresh(self):
        self.grant()
        self.assertEqual(self.ph.put('/api/stock-audit/rules/sale_below_cost/', {'params': {'min_loss_egp': 1}},
                                     format='json').status_code, 403)
        self.assertEqual(self.ph.post('/api/stock-audit/refresh/').status_code, 403)
        self.assertEqual(self.ph.get('/api/stock-audit/access/').status_code, 403)


class ReviewTests(StockAuditApiBase):

    def test_review_moves_and_notes(self):
        url = f'/api/stock-audit/findings/{self.f1.pk}/review/'
        r = self.sup.post(url, {'status': 'dismissed'}, format='json')
        self.assertEqual(r.status_code, 400)                                # a note is required
        r = self.sup.post(url, {'status': 'reviewed'}, format='json')
        self.assertEqual((r.status_code, r.json()['status']), (200, 'reviewed'))
        self.assertEqual(r.json()['next_statuses'], ['open'])
        r = self.sup.post(url, {'status': 'escalated', 'note': 'x'}, format='json')
        self.assertEqual(r.status_code, 400)                                # reviewed → escalated not allowed
        r = self.sup.post(url, {'status': 'open', 'note': 'wrong call'}, format='json')
        self.assertEqual(r.status_code, 200)
        self.f1.refresh_from_db()
        self.assertEqual(self.f1.reviewed_by, self.sup_p)
        self.assertIn('wrong call', self.f1.review_note)
        self.assertEqual(AuditLog.objects.filter(action='stock_audit_finding_reviewed').count(), 2)

    def test_viewer_cannot_review(self):
        r = self.viewer.post(f'/api/stock-audit/findings/{self.f1.pk}/review/', {'status': 'reviewed'}, format='json')
        self.assertEqual(r.status_code, 403)

    def test_branch_manager_reviews_own_branch_only(self):
        self.grant()
        ok = self.ph.post(f'/api/stock-audit/findings/{self.f1.pk}/review/', {'status': 'reviewed'}, format='json')
        self.assertEqual(ok.status_code, 200)
        other = self.ph.post(f'/api/stock-audit/findings/{self.f2.pk}/review/', {'status': 'reviewed'}, format='json')
        self.assertEqual(other.status_code, 404)
        r = self.ph.post('/api/stock-audit/findings/bulk-review/',
                         {'ids': [self.f2.pk], 'status': 'dismissed', 'note': 'n'}, format='json')
        self.assertEqual((r.json()['done'], len(r.json()['skipped'])), (0, 1))
        self.f2.refresh_from_db()
        self.assertEqual(self.f2.status, 'open')

    def test_bulk_review(self):
        f3 = finding('B01', n=3)
        r = self.sup.post('/api/stock-audit/findings/bulk-review/',
                          {'ids': [self.f1.pk, f3.pk], 'status': 'dismissed', 'note': 'known contract'}, format='json')
        self.assertEqual(r.json(), {'done': 2, 'skipped': []})
        self.assertEqual(StockAuditFinding.objects.filter(status='dismissed').count(), 2)

    def test_filters_and_totals(self):
        r = self.sup.get('/api/stock-audit/findings/?severity=critical')
        self.assertEqual([x['id'] for x in r.json()['results']], [self.f2.pk])
        r = self.sup.get('/api/stock-audit/findings/?group=adjustments')
        self.assertEqual(r.json()['count'], 0)
        r = self.sup.get('/api/stock-audit/findings/?sort=-value')
        self.assertEqual(r.json()['results'][0]['id'], self.f2.pk)
        self.assertEqual(Decimal(str(r.json()['totals']['value'])), Decimal('400'))

    def test_every_grid_column_sorts_both_ways(self):
        from apps.stock_audit.views import SORTS
        for key in ('date', 'value', 'rule', 'branch', 'item', 'status', 'doc', 'user'):
            self.assertIn(key, SORTS)
            asc = [x['id'] for x in self.sup.get(f'/api/stock-audit/findings/?sort={key}').json()['results']]
            desc = [x['id'] for x in self.sup.get(f'/api/stock-audit/findings/?sort=-{key}').json()['results']]
            self.assertEqual(sorted(asc), sorted(desc), key)
        r = self.sup.get('/api/stock-audit/findings/?sort=value')
        self.assertEqual(r.json()['results'][-1]['id'], self.f2.pk)


class ThresholdTests(StockAuditApiBase):

    def test_quality_manager_edits_chain_values(self):
        r = self.qm.put('/api/stock-audit/rules/sale_below_cost/', {'params': {'min_loss_egp': 50}}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['params']['min_loss_egp'], 50)
        self.assertEqual(r.json()['params']['min_loss_pct'], 5)            # untouched default kept
        log = AuditLog.objects.get(action='stock_audit_rule_changed')
        self.assertEqual(log.new_data['params']['min_loss_egp'], 50)

    def test_invalid_values_rejected_and_supervisor_cannot_edit(self):
        r = self.qm.put('/api/stock-audit/rules/sale_below_cost/', {'params': {'min_loss_pct': 500}}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.sup.put('/api/stock-audit/rules/sale_below_cost/', {'params': {'min_loss_pct': 9}}, format='json')
        self.assertEqual(r.status_code, 403)

    def test_branch_override_set_and_cleared(self):
        url = '/api/stock-audit/rules/stock_adjustment/overrides/B02/'
        r = self.qm.put(url, {'params': {'single_egp': 1000}, 'enabled': None}, format='json')
        self.assertEqual(r.status_code, 200)
        rules = self.sup.get('/api/stock-audit/rules/').json()
        adj = next(x for x in rules['rules'] if x['code'] == 'stock_adjustment')
        self.assertEqual(adj['overrides'][0]['branchcode'], 'B02')
        self.assertFalse(rules['can_edit'])
        self.assertEqual(self.qm.put(url, {'params': {}, 'enabled': None}, format='json').status_code, 204)
        self.assertFalse(StockAuditRuleOverride.objects.exists())
        self.assertEqual(AuditLog.objects.filter(action='stock_audit_override_changed').count(), 2)


class AccessGrantTests(StockAuditApiBase):

    def test_admin_grants_and_revokes(self):
        r = self.admin.post('/api/stock-audit/access/', {'staff_id': self.ph_p.pk, 'note': 'مدير فرع أ'}, format='json')
        self.assertEqual(r.status_code, 201)
        self.assertEqual(self.admin.post('/api/stock-audit/access/', {'staff_id': self.ph_p.pk},
                                         format='json').status_code, 400)          # already active
        self.assertEqual(self.ph.get('/api/stock-audit/findings/').status_code, 200)
        cands = self.admin.get('/api/stock-audit/access/candidates/?q=sa_ph').json()['rows']
        self.assertTrue(next(c for c in cands if c['id'] == self.ph_p.pk)['granted'])
        self.assertEqual(self.admin.post(f'/api/stock-audit/access/{r.json()["id"]}/revoke/').status_code, 200)
        self.assertEqual(self.ph.get('/api/stock-audit/findings/').status_code, 403)
        self.assertEqual(AuditLog.objects.filter(action__startswith='stock_audit_access').count(), 2)

    def test_only_admin_manages_access(self):
        self.assertEqual(self.qm.post('/api/stock-audit/access/', {'staff_id': self.ph_p.pk},
                                      format='json').status_code, 403)


class ExportAndRankingTests(StockAuditApiBase):

    def test_export_xlsx_scoped(self):
        r = self.sup.get('/api/stock-audit/findings/export/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.grant()
        import io
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(self.ph.get('/api/stock-audit/findings/export/').content)).active
        self.assertEqual(ws.max_row, 2)                                    # header + own branch only
        self.assertEqual(ws['D2'].value, 'فرع أ')                          # branch name, never a bare code
        self.assertEqual(self.viewer.get('/api/stock-audit/findings/export/').status_code, 403)

    def test_adjustments_by_user(self):
        for bc, user in (('B01', '7'), ('B02', '8')):
            StockLedgerLine.objects.create(branchcode=bc, doccode='150', docnumber=1, docdate=TODAY,
                                           itemcode='100001', usercode=user, transqty=2, newcostprice=50)
        rows = self.sup.get('/api/stock-audit/by-user/').json()['rows']
        self.assertEqual({r['usercode'] for r in rows}, {'7', '8'})
        self.grant()
        rows = self.ph.get('/api/stock-audit/by-user/').json()['rows']
        self.assertEqual([(r['usercode'], float(r['value_egp'])) for r in rows], [('7', 100.0)])


class _SyncThread:
    def __init__(self, target, daemon=None):
        self.target = target

    def start(self):
        self.target()


class RefreshTests(StockAuditApiBase):

    @patch('apps.stock_audit.views.threading.Thread', _SyncThread)
    @patch('apps.stock_audit.engine.run_stock_audit')
    @patch('apps.stock_audit.ledger.sync_ledger')
    def test_refresh_runs_sync_then_checks(self, sync, engine):
        r = self.sup.post('/api/stock-audit/refresh/')
        self.assertEqual(r.status_code, 202)
        sync.assert_called_once()
        engine.assert_called_once()
        self.assertFalse(self.sup.get('/api/stock-audit/runs/').json()['refreshing'])


class SalesContextTests(StockAuditApiBase):
    """Sale findings carry the sales channel, the sales account (e.g. a union / contract) and the
    customer PIC + name — read from the corrected PurchaseHistory and the account placeholders."""

    def setUp(self):
        super().setUp()
        from apps.customers.models import Customer, PurchaseHistory
        from apps.sync.models import SoftechPersonClassif
        SoftechPersonClassif.objects.create(ptcode='10', ptclassifcode='30', ptclassifdescr='عميل دائم')
        SoftechPersonClassif.objects.create(ptcode='10', ptclassifcode='91', ptclassifdescr='عميل نقدى')
        self.union = Customer.objects.create(name='نقابة الزراعين', softech_pic='ACCT-4801', is_account_placeholder=True,
                                             softech_ptclassifcode='30')
        Customer.objects.create(name='Walk-In Customer', softech_pic='ACCT-1510', is_account_placeholder=True,
                                softech_ptclassifcode='91')
        person = Customer.objects.create(name='شركة مكة خصم', softech_pic='B01HD8329')
        StockAuditFinding.objects.filter(pk=self.f1.pk).update(personcode='4801')
        StockAuditFinding.objects.filter(pk=self.f2.pk).update(personcode='1510')
        PurchaseHistory.objects.create(
            customer=person, branch=self.b1, doc_code='115', docnumber='1', sales_channel='30',
            invoice_date=timezone.now(), softech_phcode='B01HD8329', softech_personcode='4801',
            softech_invoice_id=f'B01-115-1-{TODAY:%Y%m%d}')

    def _row(self, rows, pk):
        return next(x for x in rows if x['id'] == pk)

    def test_findings_show_channel_account_and_pic(self):
        rows = self.sup.get('/api/stock-audit/findings/').json()['results']
        s = self._row(rows, self.f1.pk)['sales']
        self.assertEqual((s['channel'], s['channel_label'], s['account_code'], s['account_name']),
                         ('30', 'عميل دائم', '4801', 'نقابة الزراعين'))
        self.assertEqual((s['pic'], s['pic_name']), ('B01HD8329', 'شركة مكة خصم'))
        # no synced invoice: channel falls back to the account's own channel, no PIC
        s2 = self._row(rows, self.f2.pk)['sales']
        self.assertEqual((s2['channel_label'], s2['account_name'], s2['pic']), ('عميل نقدى', 'Walk-In Customer', ''))

    def test_channel_filter_account_sort_and_options(self):
        r = self.sup.get('/api/stock-audit/findings/?channel=30').json()
        self.assertEqual([x['id'] for x in r['results']], [self.f1.pk])
        r = self.sup.get('/api/stock-audit/findings/?sort=account').json()
        self.assertEqual([x['id'] for x in r['results']], [self.f2.pk, self.f1.pk])   # 1510 < 4801
        codes = {c['code'] for c in self.sup.get('/api/stock-audit/me/').json()['channels']}
        self.assertEqual(codes, {'30', '91'})

    def test_export_has_the_sales_columns(self):
        import openpyxl
        from io import BytesIO
        r = self.sup.get('/api/stock-audit/findings/export/')
        ws = openpyxl.load_workbook(BytesIO(r.content)).active
        head = [c.value for c in ws[1]]
        self.assertEqual(head[-5:], ['قناة البيع', 'كود الحساب', 'الحساب', 'PIC العميل', 'اسم العميل'])
        rows = {tuple(c.value for c in row)[-5:] for row in ws.iter_rows(min_row=2)}
        self.assertIn(('عميل دائم', '4801', 'نقابة الزراعين', 'B01HD8329', 'شركة مكة خصم'), rows)
