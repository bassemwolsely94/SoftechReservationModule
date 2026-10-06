"""
Supplier-code coverage report (apps/supply/code_report.py): same resolution rule as live
matching (confirmed mapping → SOFTECH itemssuppliers copy), conflicts vs confirmed items,
lines the code can now match, per-supplier rows, Excel + API.
"""
import io

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.catalog import supplier_links as sl
from apps.supply import code_report
from apps.supply.models import AvailabilityBatch, AvailabilityLine


class CodeReportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from apps.catalog.models import Item
        from apps.invoices.models import InvoiceLine, SupplierInvoice, VendorItemMapping, VendorProfile
        from .factories import make_user
        cls.a = Item.objects.create(softech_id='640001', name='ALPHA 10MG', is_active=True)
        cls.b = Item.objects.create(softech_id='640002', name='BETA 20MG', is_active=True)
        cls.c = Item.objects.create(softech_id='640003', name='GAMMA 5MG', is_active=True)
        Item.objects.create(softech_id='640004', name='DELTA 1MG', is_active=True)
        cls.v = VendorProfile.objects.create(name='PHARMA OVER SEAS', softech_personcode='565')
        VendorItemMapping.objects.create(vendor=cls.v, raw_name_normalized='alpha', vendor_item_code='A1', item=cls.a)
        sl.sync_from_rows([('640002', '565', 'B2', '0'), ('640003', '565', 'C3', '0'),
                           ('640001', '565', 'DUP', '0'), ('640004', '565', 'DUP', '0')])   # one pair = one row
        batch = AvailabilityBatch.objects.create(source='excel', supplier=cls.v)
        mk = lambda **kw: AvailabilityLine.objects.create(batch=batch, raw_text=kw.pop('t', 'x'), **kw)
        mk(supplier_item_code='A1', item=cls.a, is_confirmed=True)          # mapping, agrees
        mk(supplier_item_code='B2', item=cls.c, is_confirmed=True, t='gamma?')  # mirror says BETA → conflict
        mk(supplier_item_code='C3')                                          # unmatched → newly matchable
        mk(supplier_item_code='DUP')                                         # ambiguous in SOFTECH → unknown
        mk(supplier_item_code='ZZ')                                          # unknown
        mk(t='no code line', match_reason={'via': 'barcode'}, item=cls.a)    # barcode, no code
        from apps.branches.models import Branch
        br = Branch.objects.create(softech_branch_id='100', name='HQ')
        inv = SupplierInvoice.objects.create(vendor=cls.v, supplier_name='PHARMA OVER SEAS', branch=br)
        InvoiceLine.objects.create(invoice=inv, vendor_item_code='B2', item=cls.b, is_confirmed=True)
        cls.user, _, _ = make_user('op_codes', role='purchasing')

    def test_totals_and_rule(self):
        rep = code_report.build(30)
        t = rep['totals']
        self.assertEqual((t['lines'], t['with_code']), (7, 6))
        self.assertEqual((t['matched'], t['by_mapping'], t['by_mirror']), (4, 1, 3))
        self.assertEqual((t['ambiguous'], t['unknown']), (0, 2))
        self.assertEqual((t['agrees'], t['conflicts'], t['newly_matchable'], t['by_barcode']), (2, 1, 1, 1))
        self.assertEqual(rep['conflicts'][0]['confirmed_item'], 'GAMMA 5MG')
        self.assertEqual(rep['conflicts'][0]['code_item'], 'BETA 20MG')
        self.assertEqual(rep['newly_matchable'][0]['code_item'], 'GAMMA 5MG')
        self.assertEqual(rep['by_source']['invoices']['agrees'], 1)
        s = rep['suppliers'][0]
        self.assertEqual((s['supplier'], s['mirror_codes']), ('PHARMA OVER SEAS', 4))
        self.assertFalse(rep['mirror']['empty'])

    def test_api_json_and_excel(self):
        c = APIClient()
        c.force_authenticate(self.user)
        r = c.get('/api/supply/reports/supplier-codes/', {'days': 30})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['totals']['matched'], 4)
        x = c.get('/api/supply/reports/supplier-codes/', {'days': 30, 'format': 'xlsx'})
        wb = load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ['ملخص', 'الموردون', 'تعارضات', 'يمكن مطابقتها'])
        self.assertEqual(wb['تعارضات'].max_row, 3)                     # title + header + 1 conflict
