"""
Excel exports on /supply (owner 2026-10-25): each file holds exactly the rows the screen shows
(same filters) — «التوزيعة», the inter-branch pairs on «طلبات التوريد / ISR», and
«مقارنة أسعار الشراء». Plus the shared text rule the exports filter with.
"""
import io
from unittest import mock

from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.catalog.wildcard import text_matches
from apps.purchasing import supply_exports as X


class TextRuleTests(SimpleTestCase):
    def test_same_rules_as_the_screen(self):
        self.assertTrue(text_matches('99863 MILGA ADVANCE 30TAB', 'milga*30'))      # parts in order
        self.assertTrue(text_matches('VOLTAREN 50MG 20TAB', '50 voltaren'))         # any order, all words
        self.assertFalse(text_matches('VOLTAREN 50MG 20TAB', '75*voltaren'))
        self.assertTrue(text_matches('anything', ''))


def _sug(code='4179', cat='over_piled'):
    return {'category': cat, 'itemcode': code, 'item_name': 'PIASCLEDINE 300', 'sells_at': 5, 'of_branches': 5,
            'age_days': 900, 'from_branch': '130', 'network_avg': 45.0, 'reason': 'r', 'total_qty': 6,
            'from_position': {'stock': 7.0, 'rate': 1.33, 'max': 2.0, 'live': True},
            'to_branches': [{'code': '170', 'qty': 5, 'stock': 17.0, 'rate': 43.67, 'max': 65.6, 'live': True},
                            {'code': '160', 'qty': 1, 'stock': 0.0, 'rate': 2.0, 'max': None, 'live': False}]}


class DistributionExcelTests(TestCase):   # the sheet reads «بيانات الصنف» live
    def test_one_row_per_move_with_both_branches(self):
        with mock.patch('apps.finance.recon_labels.branch_name', return_value=''):
            wb = load_workbook(io.BytesIO(X.distribution_excel([_sug()], filters={'بحث': 'piasc*'})))
        ws = wb['التوزيعة']
        self.assertEqual(ws.max_row, 3)                                   # header + 2 moves
        heads = [c.value for c in ws[1]]
        row = dict(zip(heads, [c.value for c in ws[2]]))
        self.assertEqual((row['رصيد الحائز'], row['بيع الحائز/شهر'], row['رصيد المستهدف'], row['بيع المستهدف/شهر']),
                         (7.0, 1.33, 17.0, 43.67))
        self.assertIn('piasc*', wb['الشرح']['A2'].value)


class ApiTests(TestCase):
    def setUp(self):
        from .factories import make_user
        self.user, _, _ = make_user('xl_admin', role='admin')
        self.c = APIClient()
        self.c.force_authenticate(self.user)

    def test_distribution_export_applies_category_and_search(self):
        sugs = [_sug('4179'), _sug('9999', cat='new_no_rate')]
        sugs[1]['item_name'] = 'OTHER'
        with mock.patch('apps.purchasing.distribution.analyze', return_value=sugs), \
                mock.patch('apps.finance.recon_labels.branch_name', return_value=''):
            r = self.c.get('/api/purchasing/isr/distribution/', {'format': 'xlsx', 'category': 'over_piled', 'q': 'piasc*'})
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content))['التوزيعة']
        self.assertEqual({ws.cell(i, 2).value for i in range(2, ws.max_row + 1)}, {'4179'})

    def test_transfer_pair_export(self):
        r = self.c.get('/api/purchasing/isr/transfer-preview/export/', {'from': '150', 'to': '160'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('transfer_150_to_160', r['Content-Disposition'])
        self.assertEqual(load_workbook(io.BytesIO(r.content)).sheetnames, ['الأصناف', 'المسارات', 'الشرح'])

    def test_prices_export_applies_screen_filters(self):
        board = {'items': [{'code': '1', 'name': 'MILGA', 'state': 'cheaper', 'producer': 'EVA', 'suppliers': [{'key': 'D1'}]},
                           {'code': '2', 'name': 'OTHER', 'state': 'best', 'producer': 'EVA', 'suppliers': [{'key': 'D2'}]}],
                 'suppliers': [{'key': 'D1', 'name': 'RAMCO'}, {'key': 'D2', 'name': 'IBN SINA'}],
                 'state_labels': {'cheaper': 'يوجد أرخص', 'best': 'نشتري من الأفضل'}}
        from apps.supply.price_history import filter_board
        f = filter_board(board, state='cheaper', supplier='D1', q='mil*')
        self.assertEqual([r['code'] for r in f['items']], ['1'])
        self.assertEqual([s['name'] for s in f['suppliers']], ['RAMCO'])
        self.assertEqual(f['filters'], {'state': 'يوجد أرخص', 'supplier': 'RAMCO', 'q': 'mil*'})
