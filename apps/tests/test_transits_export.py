"""
apps/tests/test_transits_export.py

Picking-sheet (ورقة التجميع) export + pick-zone management:
  • Ruleset.classify — DB-driven rules, precedence (override → fridge → price → keywords)
  • generate_picking_workbook — sheet layout, consolidation, qty matrix, tags
  • export-picking endpoints — single GET, bulk POST, branch scoping, audit
  • pick-zones / pick-rules / item-overrides management API
"""
import io
from decimal import Decimal

from django.test import TestCase
from rest_framework import status
from rest_framework.test import APIClient

from apps.catalog.models import Item, ItemStock
from apps.transits.export import (
    PRICE_THRESHOLD_DEFAULT,
    build_ruleset,
    generate_picking_workbook,
    seed_default_pick_zones,
)
from apps.transits.models import (
    InTransitAuditEvent,
    InTransitTransfer,
    ItemPickOverride,
    PickZone,
    PickZoneRule,
)
from .factories import make_admin, make_branch, make_branch2, make_item, make_user


def _col(ws, header):
    """
    1-based index of a column by its header text — tests must not depend on
    physical column order (columns get inserted as the sheets evolve).

    Exact matches win, so 'الكمية' picks the table header and not the
    'إجمالي الكمية (عبوات)' label in the document info block above it.
    """
    rows = list(ws.iter_rows(values_only=True))
    for match_exact in (True, False):
        for row in rows:
            for j, cell in enumerate(row, start=1):
                if not cell:
                    continue
                text = str(cell).strip()
                if (text == header) if match_exact else (header in text):
                    return j
    raise AssertionError(f'header {header!r} not found in {ws.title}')


def _cell(ws, row_tuple, header):
    """Value of `header`'s column within an already-fetched values_only row."""
    return row_tuple[_col(ws, header) - 1]


def _make_transfer(doc='900001', supply=None, recv=None, items=None):
    return InTransitTransfer.objects.create(
        erp_doc_number=doc,
        erp_supplying_branch_code=supply.softech_branch_id if supply else '100',
        erp_receiving_branch_code=recv.softech_branch_id if recv else '130',
        supplying_branch=supply,
        receiving_branch=recv,
        issue_date='2026-07-01',
        item_count=len(items or []),
        items_snapshot=items or [],
        transit_status='in_transit',
    )


class RulesetTests(TestCase):
    """DB-driven classification — parity with the Power Query rules + new precedence."""

    def setUp(self):
        seed_default_pick_zones()
        self.rs = build_ruleset()

    def _c(self, name, price=None, fridge=False, itemcode=None):
        zone, _tag = self.rs.classify(name, price, fridge, itemcode=itemcode)
        return zone.name if zone else None

    def test_threshold_is_500(self):
        self.assertEqual(PRICE_THRESHOLD_DEFAULT, 500)
        self.assertEqual(float(self.rs.threshold), 500.0)

    def test_price_over_500_goes_expensive(self):
        self.assertEqual(self._c('ELIQUIS 2.5MG 20TAB (XX)', 532), 'غوالى-ثلاجه')

    def test_price_300_no_longer_expensive(self):
        # under the old 250 bar this was غوالى — now it stays a tablet
        self.assertEqual(self._c('EXFORGE HCT 5/160/12.5MG 14TAB', 300), 'أقراص')

    def test_fridge_flag_beats_price_and_keywords(self):
        self.assertEqual(self._c('SOME TAB 20', 10, fridge=True), 'غوالى-ثلاجه')

    def test_fridge_keyword_in_name(self):
        self.assertEqual(
            self._c('ADRENOCORTINE 1MG 1AMP (FRIDGE)', 50), 'غوالى-ثلاجه')

    def test_keyword_rules(self):
        cases = [
            ('ZYRTEC 10MG 20TAB', 100, 'أقراص'),
            ('DALACIN-C 300MG 10CAP', 114, 'أقراص'),
            ('PANTHENOL 2% CREAM 50GM', 80, 'كريم-مرهم-جل'),
            ('PRISOLINE EYE NASAL DROPS 15ML', 23, 'قطرات-نقط'),
            ('VOLTAREN 75MG/3ML 3AMP', 51, 'حقن'),
            ('GLYCERIN ADULT 5 SUPP PHARCO', 12, 'لبوس'),
            ('OPLEX N SYRUP 125ML', 31, 'شراب'),
            ('CETAL 250MG / 5ML SUSP 60ML', 31, 'شراب'),
            ('ACETYLCYSTEINE 600MG 10SACH', 70, 'فوارات'),
            ('AVAMYS NASAL SPRAY 120 SPRAY', 160, 'بخاخات-سبراى'),
        ]
        for name, price, expected in cases:
            self.assertEqual(self._c(name, price), expected, name)

    def test_rule_order_tab_wins_over_cream(self):
        self.assertEqual(self._c('WEIRD CREAM 10TAB', 10), 'أقراص')

    def test_fallback(self):
        self.assertEqual(self._c('TOTALLY UNKNOWN THING', 10), 'غير مصنف')

    def test_item_override_beats_everything(self):
        item = make_item('TOTALLY UNKNOWN THING', 'IT009')
        zone = PickZone.objects.get(name='ألبان')
        ItemPickOverride.objects.create(item=item, zone=zone, tag='رف A3',
                                        branch=None)
        rs = build_ruleset(item_codes=['IT009'])
        z, tag = rs.classify('TOTALLY UNKNOWN THING', 10, itemcode='IT009')
        self.assertEqual(z.name, 'ألبان')
        self.assertEqual(tag, 'رف A3')

    def test_explain_reasons(self):
        z, _t, reason = self.rs.explain('ZYRTEC 10MG 20TAB', 100)
        self.assertEqual(z.name, 'أقراص')
        self.assertIn('TAB', reason)
        z, _t, reason = self.rs.explain('X (FRIDGE)', 10)
        self.assertIn('ثلاجة', reason)

    def test_pick_path_order_preserved(self):
        """Zones sort by sort_key = the team's historical lexicographic path."""
        names = list(PickZone.objects.order_by('sort_key')
                     .values_list('name', flat=True))
        self.assertEqual(names[:6], [
            'غوالى-ثلاجه', 'أقراص', 'كريم-مرهم-جل',
            'بخاخات-سبراى', 'حقن', 'لبوس',
        ])

    def test_fallback_without_db_zones(self):
        PickZoneRule.objects.all().delete()
        PickZone.objects.all().delete()
        rs = build_ruleset()
        z, _ = rs.classify('ZYRTEC 10MG 20TAB', 100)
        self.assertEqual(z.name, 'أقراص')     # in-memory defaults kick in


class WorkbookTests(TestCase):

    def setUp(self):
        seed_default_pick_zones()
        self.hq  = make_branch('الرئيسي', 'B01')
        self.br1 = make_branch2('م. رمسيس', 'B02')
        self.br2 = make_branch('فرع المعادي', 'B03')

        self.item_tab = make_item('ZYRTEC 10MG 20TAB', 'IT001')
        self.item_tab.pack_price = Decimal('100')
        self.item_tab.producer_name = 'UCB'
        self.item_tab.save()
        ItemStock.objects.create(
            item=self.item_tab, branch=self.hq,
            softech_store_code='1', quantity_on_hand=Decimal('40'),
        )
        # quarantine store must be excluded from رصيد المصدر
        ItemStock.objects.create(
            item=self.item_tab, branch=self.hq,
            softech_store_code='102', quantity_on_hand=Decimal('99'),
        )

        self.item_syr = make_item('OPLEX N SYRUP 125ML', 'IT002')
        self.item_syr.pack_price = Decimal('31')
        self.item_syr.save()

        self.t1 = _make_transfer('900001', self.hq, self.br1, [
            {'itemcode': 'IT001', 'itemname': 'ZYRTEC 10MG 20TAB', 'qty': 5,
             'unit_cost': 70, 'extended_cost': 350,
             'batch': 'L123', 'expiry': '2028-11-11', 'near_expiry': False},
            {'itemcode': 'IT002', 'itemname': 'OPLEX N SYRUP 125ML', 'qty': 10,
             'unit_cost': 20, 'extended_cost': 200,
             'batch': None, 'expiry': None, 'near_expiry': False},
        ])
        self.t2 = _make_transfer('900002', self.hq, self.br2, [
            {'itemcode': 'IT001', 'itemname': 'ZYRTEC 10MG 20TAB', 'qty': 3,
             'unit_cost': 70, 'extended_cost': 210,
             'batch': None, 'expiry': '2027-05-05', 'near_expiry': False},
        ])

    def _load(self, content):
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_single_order_workbook(self):
        content, filename, ctype = generate_picking_workbook([self.t1])
        self.assertEqual(filename, 'picking_900001.xlsx')
        self.assertIn('spreadsheetml', ctype)
        wb = self._load(content)
        self.assertEqual(wb.sheetnames, ['التجميع الموحد', 'مراجعة 900001'])

    def test_multi_order_consolidation(self):
        content, filename, _ = generate_picking_workbook([self.t1, self.t2])
        wb = self._load(content)
        self.assertEqual(
            wb.sheetnames,
            ['التجميع الموحد', 'مراجعة 900001', 'مراجعة 900002'],
        )

        ws = wb['التجميع الموحد']
        rows = list(ws.iter_rows(values_only=True))
        zyrtec = next(r for r in rows if r and r[2] == 'IT001')
        self.assertEqual(zyrtec[1], 'أقراص')           # zone name (clean, no prefix)
        self.assertEqual(zyrtec[3], 'ZYRTEC 10MG 20TAB')
        self.assertEqual(zyrtec[4], '2027-05-05')      # earliest expiry wins
        self.assertEqual(zyrtec[5], 40)                # HQ balance, quarantine excluded
        self.assertEqual(zyrtec[6], 5)                 # qty column order 1 (t1)
        self.assertEqual(zyrtec[7], 3)                 # qty column order 2 (t2)
        self.assertEqual(zyrtec[8], 8)                 # total

    def test_pick_path_order_in_sheet(self):
        """أقراص (sort 20) must appear before شراب (sort 120)."""
        content, _, _ = generate_picking_workbook([self.t1])
        wb = self._load(content)
        ws = wb['التجميع الموحد']
        zones = [r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r[1]]
        self.assertEqual(zones, ['أقراص', 'شراب'])

    def test_zone_stripe_shows_location(self):
        zone = PickZone.objects.get(name='أقراص')
        zone.location = 'ممر 2'
        zone.save()
        content, _, _ = generate_picking_workbook([self.t1])
        wb = self._load(content)
        ws = wb['التجميع الموحد']
        stripes = [r[0] for r in ws.iter_rows(min_row=3, values_only=True)
                   if r[0] and isinstance(r[0], str) and 'أقراص' in r[0]]
        self.assertIn('أقراص — ممر 2', stripes)

    def test_override_tag_lands_in_notes(self):
        zone = PickZone.objects.get(name='ألبان')
        ItemPickOverride.objects.create(item=self.item_syr, zone=zone, tag='رف B7')
        content, _, _ = generate_picking_workbook([self.t1])
        wb = self._load(content)
        ws = wb['مراجعة 900001']
        oplex = next(r for r in ws.iter_rows(values_only=True)
                     if r and r[2] == 'IT002')
        self.assertEqual(_cell(ws, oplex, 'المنطقة'), 'ألبان')     # overridden zone
        self.assertEqual(_cell(ws, oplex, 'ملاحظات'), 'رف B7')     # tag in notes

    def test_revision_sheet_content(self):
        content, _, _ = generate_picking_workbook([self.t1])
        wb = self._load(content)
        ws = wb['مراجعة 900001']
        rows = list(ws.iter_rows(values_only=True))
        zyrtec = next(r for r in rows if r and r[2] == 'IT001')
        self.assertEqual(_cell(ws, zyrtec, 'الكمية'), 5)
        self.assertEqual(_cell(ws, zyrtec, 'ت الصلاحية'), '2028-11-11')
        self.assertEqual(_cell(ws, zyrtec, 'التشغيلة'), 'L123')
        self.assertEqual(_cell(ws, zyrtec, 'سعر الجمهور'), 100)
        self.assertEqual(_cell(ws, zyrtec, 'الإجمالي'), 500)   # 5 × 100
        # whole quantity → no partial-pack note
        self.assertIn(_cell(ws, zyrtec, 'تفصيل الكمية'), (None, ''))

    def test_empty_raises(self):
        with self.assertRaises(ValueError):
            generate_picking_workbook([])


class ExportEndpointTests(TestCase):

    def setUp(self):
        seed_default_pick_zones()
        self.hq  = make_branch('الرئيسي', 'B01')
        self.br1 = make_branch2('م. رمسيس', 'B02')
        make_item('ZYRTEC 10MG 20TAB', 'IT001')
        self.t1 = _make_transfer('900001', self.hq, self.br1, [
            {'itemcode': 'IT001', 'itemname': 'ZYRTEC 10MG 20TAB', 'qty': 5,
             'unit_cost': 70, 'extended_cost': 350,
             'batch': None, 'expiry': None, 'near_expiry': False},
        ])
        _, self.admin_profile, self.admin = make_admin('transit_admin')

    def test_requires_auth(self):
        r = APIClient().get(f'/api/transits/{self.t1.id}/export-picking/')
        self.assertEqual(r.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_single_export(self):
        r = self.admin.get(f'/api/transits/{self.t1.id}/export-picking/')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('picking_900001.xlsx', r['Content-Disposition'])
        self.assertTrue(
            InTransitAuditEvent.objects.filter(
                transfer=self.t1, action='exported',
                actor=self.admin_profile,
            ).exists()
        )

    def test_bulk_export(self):
        t2 = _make_transfer('900002', self.hq, self.br1, [
            {'itemcode': 'IT001', 'itemname': 'ZYRTEC 10MG 20TAB', 'qty': 2,
             'unit_cost': 70, 'extended_cost': 140,
             'batch': None, 'expiry': None, 'near_expiry': False},
        ])
        r = self.admin.post(
            '/api/transits/export-picking/',
            {'ids': [self.t1.id, t2.id]}, format='json',
        )
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(
            InTransitAuditEvent.objects.filter(action='exported').count(), 2,
        )

    def test_bulk_requires_ids(self):
        r = self.admin.post('/api/transits/export-picking/', {}, format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_branch_scoping(self):
        other = make_branch('فرع بعيد', 'B09')
        _, _, client = make_user('faraway_ph', role='pharmacist', branch=other)
        r = client.get(f'/api/transits/{self.t1.id}/export-picking/')
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND)

        r = client.post('/api/transits/export-picking/',
                        {'ids': [self.t1.id]}, format='json')
        self.assertEqual(r.status_code, status.HTTP_404_NOT_FOUND)


class StockingTests(TestCase):
    """
    Destination-branch stocking (ترصيص) — a SEPARATE purpose from picking:
      • resolution chain: branch stocking → default stocking → picking configs
      • export-stocking classifies by the RECEIVING branch config
      • overrides are purpose-scoped
      • stock-count sheets order rows by the counting branch's stocking zones
    """

    def setUp(self):
        seed_default_pick_zones()                     # default PICKING config
        self.hq  = make_branch('الرئيسي', 'B01')
        self.br1 = make_branch2('م. رمسيس', 'B02')

        self.item_tab = make_item('ZYRTEC 10MG 20TAB', 'IT001')
        self.item_tab.pack_price = Decimal('100')
        self.item_tab.save()
        self.item_syr = make_item('OPLEX N SYRUP 125ML', 'IT002')
        self.item_syr.pack_price = Decimal('31')
        self.item_syr.save()

        self.t1 = _make_transfer('900001', self.hq, self.br1, [
            {'itemcode': 'IT001', 'itemname': 'ZYRTEC 10MG 20TAB', 'qty': 5,
             'unit_cost': 70, 'extended_cost': 350,
             'batch': None, 'expiry': None, 'near_expiry': False},
            {'itemcode': 'IT002', 'itemname': 'OPLEX N SYRUP 125ML', 'qty': 10,
             'unit_cost': 20, 'extended_cost': 200,
             'batch': None, 'expiry': None, 'near_expiry': False},
        ])
        _, self.admin_profile, self.admin = make_admin('stocking_admin')

    def _branch_stocking_config(self, branch):
        """Give a branch its own stocking layout: شراب shelf FIRST, then أقراص."""
        z_syr = PickZone.objects.create(branch=branch, purpose='stocking',
                                        name='رف الشراب', sort_key=10)
        z_tab = PickZone.objects.create(branch=branch, purpose='stocking',
                                        name='رف الأقراص', sort_key=20,
                                        is_fallback=True)
        PickZoneRule.objects.create(zone=z_syr, keywords=['SYRUP'], priority=10)
        PickZoneRule.objects.create(zone=z_tab, keywords=['TAB'], priority=20)
        return z_syr, z_tab

    def _load(self, content):
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    # ── resolution chain ──────────────────────────────────────────────────────

    def test_stocking_falls_back_to_picking_config(self):
        rs = build_ruleset(branch=self.br1.id, purpose='stocking')
        z, _ = rs.classify('ZYRTEC 10MG 20TAB', 100)
        self.assertEqual(z.name, 'أقراص')             # from default picking

    def test_branch_stocking_config_wins(self):
        self._branch_stocking_config(self.br1)
        rs = build_ruleset(branch=self.br1.id, purpose='stocking')
        z, _ = rs.classify('OPLEX N SYRUP 125ML', 31)
        self.assertEqual(z.name, 'رف الشراب')
        # picking for the same branch is untouched
        rs_p = build_ruleset(branch=self.br1.id, purpose='picking')
        z, _ = rs_p.classify('OPLEX N SYRUP 125ML', 31)
        self.assertEqual(z.name, 'شراب')

    def test_override_purpose_separation(self):
        z_milk = PickZone.objects.get(name='ألبان')
        ItemPickOverride.objects.create(item=self.item_tab, zone=z_milk,
                                        purpose='picking')
        # picking sees the override, stocking does not
        z, _ = build_ruleset(purpose='picking').classify(
            'ZYRTEC 10MG 20TAB', 100, itemcode='IT001')
        self.assertEqual(z.name, 'ألبان')
        z, _ = build_ruleset(purpose='stocking').classify(
            'ZYRTEC 10MG 20TAB', 100, itemcode='IT001')
        self.assertEqual(z.name, 'أقراص')

    # ── stocking workbook ─────────────────────────────────────────────────────

    def test_stocking_workbook_uses_receiving_branch_layout(self):
        self._branch_stocking_config(self.br1)
        content, filename, _ = generate_picking_workbook([self.t1], mode='stocking')
        self.assertEqual(filename, 'stocking_900001.xlsx')
        wb = self._load(content)
        self.assertEqual(wb.sheetnames, ['الترصيص الموحد', 'ترصيص 900001'])

        ws = wb['الترصيص الموحد']
        zones = [r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r[1]]
        # branch shelf order: الشراب before الأقراص (opposite of picking path)
        self.assertEqual(zones, ['رف الشراب', 'رف الأقراص'])

    def test_picking_workbook_unaffected_by_stocking_config(self):
        self._branch_stocking_config(self.br1)
        content, _, _ = generate_picking_workbook([self.t1], mode='picking')
        wb = self._load(content)
        ws = wb['التجميع الموحد']
        zones = [r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r[1]]
        self.assertEqual(zones, ['أقراص', 'شراب'])    # supplying picking path

    # ── endpoints ─────────────────────────────────────────────────────────────

    def test_export_stocking_endpoints(self):
        r = self.admin.get(f'/api/transits/{self.t1.id}/export-stocking/')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertIn('stocking_900001.xlsx', r['Content-Disposition'])
        self.assertTrue(InTransitAuditEvent.objects.filter(
            transfer=self.t1, action='exported',
            detail__contains='ترصيص').exists())

        r = self.admin.post('/api/transits/export-stocking/',
                            {'ids': [self.t1.id]}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)

    def test_zone_api_purpose_scoping(self):
        self._branch_stocking_config(self.br1)
        r = self.admin.get(f'/api/transits/pick-zones/?branch={self.br1.id}&purpose=stocking')
        self.assertEqual(len(r.data), 2)
        r = self.admin.get(f'/api/transits/pick-zones/?branch={self.br1.id}&purpose=picking')
        self.assertEqual(len(r.data), 0)
        r = self.admin.get('/api/transits/pick-zones/?branch=&purpose=picking')
        self.assertEqual(len(r.data), 15)

    def test_copy_defaults_stocking_clones_picking(self):
        # default stocking is empty → copying clones the default PICKING config
        r = self.admin.post('/api/transits/pick-zones/copy-defaults/',
                            {'branch': self.br1.id, 'purpose': 'stocking'},
                            format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r.data['zones_created'], 15)
        self.assertEqual(
            PickZone.objects.filter(branch=self.br1, purpose='stocking').count(), 15)

    # ── stock count integration ───────────────────────────────────────────────

    def _make_session(self):
        from apps.stockcount.models import StockCountSession, StockCountSnapshot
        session = StockCountSession.objects.create(
            name='جرد اختبار', branch_code=self.br1.softech_branch_id,
            status='snapshot_taken',
        )
        StockCountSnapshot.objects.create(
            session=session, item_code='IT001', item_name='ZYRTEC 10MG 20TAB',
            branch_code=session.branch_code, expected_qty=Decimal('7'),
        )
        StockCountSnapshot.objects.create(
            session=session, item_code='IT002', item_name='OPLEX N SYRUP 125ML',
            branch_code=session.branch_code, expected_qty=Decimal('3'),
        )
        return session

    def test_count_sheet_ordered_by_stocking_zones(self):
        from apps.stockcount.excel_io import generate_count_sheet
        self._branch_stocking_config(self.br1)
        session = self._make_session()
        content, fname, _ = generate_count_sheet(session)
        wb = self._load(content)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        self.assertIn('مرتبة حسب مناطق الترصيص', str(rows[0][0]))
        # stripes present (col A empty, col B carries the zone label)
        stripes = [r[1] for r in rows if r[0] in (None, '') and r[1]]
        self.assertTrue(any('رف الشراب' in str(x) for x in stripes))
        # SYRUP row comes before TAB row (branch shelf order)
        codes = [r[0] for r in rows if r[0] in ('IT001', 'IT002')]
        self.assertEqual(codes, ['IT002', 'IT001'])

    def test_count_sheet_reimport_ignores_stripes(self):
        from apps.stockcount.excel_io import generate_count_sheet, parse_count_sheet
        self._branch_stocking_config(self.br1)
        session = self._make_session()
        content, fname, _ = generate_count_sheet(session)

        # simulate the user filling counted quantities
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        counted_col = _col(ws, 'الكمية المعدودة')
        filled = 0
        for row in ws.iter_rows(min_row=3):
            if row[0].value in ('IT001', 'IT002'):
                row[counted_col - 1].value = 5
                filled += 1
        self.assertEqual(filled, 2)
        buf = io.BytesIO()
        wb.save(buf)

        results = parse_count_sheet(buf.getvalue(), fname)
        self.assertEqual(
            sorted(r['item_code'] for r in results), ['IT001', 'IT002'])

    def test_count_sheet_without_stocking_config_uses_picking(self):
        from apps.stockcount.excel_io import generate_count_sheet
        session = self._make_session()
        content, _, _ = generate_count_sheet(session)
        wb = self._load(content)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        codes = [r[0] for r in rows if r[0] in ('IT001', 'IT002')]
        # default picking path: أقراص (TAB) before شراب (SYRUP)
        self.assertEqual(codes, ['IT001', 'IT002'])


class PrintLayoutTests(TestCase):
    """
    A4 printability of every export (picking / stocking / stock count):
      • all columns fit the declared number of A4 pages across
      • the item-name column wraps, so long names are never clipped
      • side margins are effectively zero (max usable width)
    """

    # A4 landscape 11.69in / portrait 8.27in, minus 2 × 0.1in margins.
    USABLE_LANDSCAPE_IN = 11.69 - 0.2
    USABLE_PORTRAIT_IN  = 8.27 - 0.2

    def setUp(self):
        seed_default_pick_zones()
        self.hq  = make_branch('الرئيسي', 'B01')
        self.br1 = make_branch2('م. رمسيس', 'B02')

        # a deliberately long name (real catalog names reach ~90 chars)
        self.long_name = ('THIOTACID ORIGINAL 600MG 30TAB (3STRIPSX10) BIG NEW SIZE '
                          'ثيوتاسيد اورجينال')
        it = make_item(self.long_name, 'IT001')
        it.pack_price = Decimal('100')
        it.producer_name = 'EVA PHARMA FOR PHARMACEUTICALS'
        it.save()

        self.transfers = []
        for i in range(4):
            self.transfers.append(_make_transfer(
                f'90010{i}', self.hq, self.br1,
                [{'itemcode': 'IT001', 'itemname': self.long_name, 'qty': 5,
                  'unit_cost': 70, 'extended_cost': 350,
                  'batch': 'L123456789', 'expiry': '2028-11-11',
                  'near_expiry': False}],
            ))

    def _sheet_width_in(self, ws):
        from openpyxl.utils import get_column_letter
        cols = range(1, ws.max_column + 1)
        visible, n = 0.0, 0
        for i in cols:
            dim = ws.column_dimensions[get_column_letter(i)]
            if dim.hidden:
                continue
            visible += (dim.width or 8.43)
            n += 1
        return (7 * visible + 5 * n) / 96.0

    def _assert_fits(self, ws, landscape=True):
        pages = ws.page_setup.fitToWidth or 1
        usable = (self.USABLE_LANDSCAPE_IN if landscape
                  else self.USABLE_PORTRAIT_IN) * pages
        width = self._sheet_width_in(ws)
        self.assertLessEqual(
            width, usable + 0.05,
            f'{ws.title}: {width:.2f}in exceeds {usable:.2f}in ({pages} page(s))')
        # margins effectively zero
        self.assertLessEqual(ws.page_margins.left, 0.15)
        self.assertLessEqual(ws.page_margins.right, 0.15)
        # openpyxl round-trips paperSize as an int; the constant is a str
        self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_A4))

    def _name_cell(self, ws, col=4):
        for row in ws.iter_rows(min_row=3):
            if row[col - 1].value and str(row[col - 1].value).startswith('THIOTACID'):
                return row[col - 1]
        return None

    def _load(self, content):
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_picking_sheets_fit_a4_and_wrap(self):
        for count in (1, 4):
            content, _, _ = generate_picking_workbook(
                self.transfers[:count], mode='picking')
            wb = self._load(content)
            for sn in wb.sheetnames:
                ws = wb[sn]
                self._assert_fits(ws)
                cell = self._name_cell(ws)
                self.assertIsNotNone(cell, f'{sn}: item name row not found')
                self.assertTrue(cell.alignment.wrap_text,
                                f'{sn}: item name must wrap')

    def test_stocking_sheets_fit_a4_and_wrap(self):
        content, _, _ = generate_picking_workbook(
            self.transfers[:2], mode='stocking')
        wb = self._load(content)
        for sn in wb.sheetnames:
            ws = wb[sn]
            self._assert_fits(ws)
            self.assertTrue(self._name_cell(ws).alignment.wrap_text)

    def test_many_orders_stay_readable_over_multiple_pages(self):
        """20 branch columns can't fit one page — columns must stay legible."""
        from openpyxl.utils import get_column_letter
        many = [
            _make_transfer(f'9200{i:02d}', self.hq, self.br1,
                           [{'itemcode': 'IT001', 'itemname': self.long_name,
                             'qty': 1, 'unit_cost': 70, 'extended_cost': 70,
                             'batch': None, 'expiry': None, 'near_expiry': False}])
            for i in range(20)
        ]
        content, _, _ = generate_picking_workbook(many, mode='picking')
        ws = self._load(content)['التجميع الموحد']

        self.assertGreater(ws.page_setup.fitToWidth, 1)     # spans pages
        self._assert_fits(ws)
        # identity columns repeat on every printed page
        self.assertEqual(ws.print_title_cols, '$A:$D')
        # nothing collapsed into an unreadable sliver
        name_w = ws.column_dimensions[get_column_letter(4)].width
        self.assertGreaterEqual(name_w, 20)
        for i in range(1, ws.max_column + 1):
            self.assertGreaterEqual(
                ws.column_dimensions[get_column_letter(i)].width or 8.43, 4)

    def test_count_sheet_fits_a4_portrait_and_wraps(self):
        from apps.stockcount.excel_io import generate_count_sheet
        from apps.stockcount.models import StockCountSession, StockCountSnapshot

        session = StockCountSession.objects.create(
            name='جرد طباعة', branch_code=self.br1.softech_branch_id,
            status='snapshot_taken',
        )
        StockCountSnapshot.objects.create(
            session=session, item_code='IT001', item_name=self.long_name,
            branch_code=session.branch_code, expected_qty=Decimal('7'),
        )

        for variance in (False, True):
            content, _, _ = generate_count_sheet(session, include_variance=variance)
            ws = self._load(content).active
            self._assert_fits(ws, landscape=False)
            cell = self._name_cell(ws, col=2)
            self.assertIsNotNone(cell)
            self.assertTrue(cell.alignment.wrap_text)


class QuantityAndMoneyFormatTests(TestCase):
    """
    Reading the printed numbers must be unambiguous:
      • money columns always render with 2 decimals
      • quantities keep their fraction (2.333 never becomes 2.33 or 2)
      • a part-pack quantity is spelled out ('2 علبة + 1 شريط') and highlighted
    """

    def setUp(self):
        seed_default_pick_zones()
        self.hq  = make_branch('الرئيسي', 'B01')
        self.br1 = make_branch2('م. رمسيس', 'B02')

        # 4 strips per pack, priced with an awkward fraction
        self.item = make_item('PROCORALAN 5MG 28TAB (4STRIPSX7)', 'IT001')
        self.item.pack_price = Decimal('123.456')
        self.item.pack_qty = 4
        self.item.save()

        self.t = _make_transfer('900001', self.hq, self.br1, [
            {'itemcode': 'IT001', 'itemname': self.item.name, 'qty': 2.25,
             'unit_cost': 70, 'extended_cost': 157.5,
             'batch': None, 'expiry': None, 'near_expiry': False},
        ])

    def _load(self, content):
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_qty_parts_spells_out_partial_packs(self):
        from apps.transits.export import _qty_parts
        self.assertEqual(_qty_parts(2.25, 4, 'X (4STRIPSX7)'), ('2 علبة + 1 شريط', True))
        self.assertEqual(_qty_parts(0.6, 5, 'NOVOMIX FLEXPEN'), ('3 قلم', True))
        self.assertEqual(_qty_parts(0.5, 2, 'MERALGO 20 CAP (2STRIPSX10)'),
                         ('1 شريط', True))
        self.assertEqual(_qty_parts(0.75, 4, 'JANUMET (4STRIPSX14)'), ('3 شريط', True))
        # whole quantities produce no note at all
        self.assertEqual(_qty_parts(5, 3, 'X'), ('', False))
        self.assertEqual(_qty_parts(0, 3, 'X'), ('', False))
        # odd ERP fraction that is not a clean sub-unit split
        text, partial = _qty_parts(2.00033, 3, 'X')
        self.assertTrue(partial)
        self.assertIn('علبة', text)

    def test_revision_sheet_money_and_qty_formats(self):
        content, _, _ = generate_picking_workbook([self.t])
        ws = self._load(content)['مراجعة 900001']

        row = next(r for r in ws.iter_rows(min_row=3)
                   if r[2].value == 'IT001')
        get = lambda h: row[_col(ws, h) - 1]

        # quantity keeps the fraction and is spelled out; format shows exactly
        # its decimals (2.25 → 2 dp), never a trailing dot
        self.assertAlmostEqual(get('الكمية').value, 2.25)
        self.assertEqual(get('الكمية').number_format, '#,##0.00')
        self.assertEqual(get('تفصيل الكمية').value, '2 علبة + 1 شريط')
        self.assertTrue(get('تفصيل الكمية').font.bold)      # highlighted

        # money columns: capped at 2 dp (conventional), value kept full precision
        self.assertAlmostEqual(get('سعر الجمهور').value, 123.456)
        self.assertEqual(get('سعر الجمهور').number_format, '#,##0.00')
        self.assertAlmostEqual(get('الإجمالي').value, 277.776)   # 2.25 × 123.456
        self.assertEqual(get('الإجمالي').number_format, '#,##0.00')

    def test_whole_numbers_have_no_trailing_dot(self):
        """Integer qty/money render as '5' not '5.' (the reported bug)."""
        whole = _make_transfer('900099', self.hq, self.br1, [
            {'itemcode': 'IT001', 'itemname': self.item.name, 'qty': 5,
             'unit_cost': 100, 'extended_cost': 500,
             'batch': None, 'expiry': None, 'near_expiry': False}])
        content, _, _ = generate_picking_workbook([whole])
        ws = self._load(content)['مراجعة 900099']
        row = next(r for r in ws.iter_rows(min_row=3) if r[2].value == 'IT001')
        get = lambda h: row[_col(ws, h) - 1]
        # 5 packs, price 123.456 → total 617.28
        self.assertEqual(get('الكمية').value, 5)
        self.assertEqual(get('الكمية').number_format, '#,##0')     # no '.'
        # price is a decimal → money capped at 2 dp
        self.assertEqual(get('سعر الجمهور').number_format, '#,##0.00')

    def test_header_block_separates_the_money_bases(self):
        """ERP document value, ERP line cost and retail total are distinct."""
        self.t.doc_value = Decimal('7926.59')
        self.t.save()
        content, _, _ = generate_picking_workbook([self.t])
        ws = self._load(content)['مراجعة 900001']
        text = '\n'.join(str(c) for row in ws.iter_rows(values_only=True)
                         for c in row if c)

        self.assertIn('قيمة المستند (ERP)', text)
        self.assertIn('إجمالي تكلفة البنود (ERP)', text)
        self.assertIn('الإجمالي (سعر الجمهور)', text)

        # the ERP value cell prints its exact 2 decimals (7926.59)
        for row in ws.iter_rows(min_row=2, max_row=6):
            for c in row:
                if c.value == 7926.59:
                    self.assertEqual(c.number_format, '#,##0.00')
                    return
        self.fail('ERP document value not written')

    def test_consolidated_sheet_shows_partial_totals(self):
        content, _, _ = generate_picking_workbook([self.t])
        ws = self._load(content)['التجميع الموحد']
        row = next(r for r in ws.iter_rows(min_row=3) if r[2].value == 'IT001')
        get = lambda h: row[_col(ws, h) - 1]

        self.assertAlmostEqual(get('إجمالي الكمية').value, 2.25)
        self.assertEqual(get('إجمالي الكمية').number_format, '#,##0.00')
        self.assertEqual(get('تفصيل الكمية').value, '2 علبة + 1 شريط')

    def test_count_sheet_shows_partial_expected_qty(self):
        from apps.stockcount.excel_io import generate_count_sheet
        from apps.stockcount.models import StockCountSession, StockCountSnapshot

        session = StockCountSession.objects.create(
            name='جرد كسور', branch_code=self.br1.softech_branch_id,
            status='snapshot_taken',
        )
        StockCountSnapshot.objects.create(
            session=session, item_code='IT001', item_name=self.item.name,
            branch_code=session.branch_code, expected_qty=Decimal('2.25'),
        )
        content, _, _ = generate_count_sheet(session)
        ws = self._load(content).active

        row = next(r for r in ws.iter_rows(min_row=3) if r[0].value == 'IT001')
        get = lambda h: row[_col(ws, h) - 1]
        self.assertAlmostEqual(get('الكمية المتوقعة').value, 2.25)
        self.assertEqual(get('تفصيل المتوقع').value, '2 علبة + 1 شريط')
        # and the sheet still re-imports cleanly with the extra column
        from apps.stockcount.excel_io import parse_count_sheet
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws2 = wb.active
        ws2.cell(row=row[0].row, column=_col(ws2, 'الكمية المعدودة'), value=2)
        buf = io.BytesIO()
        wb.save(buf)
        parsed = parse_count_sheet(buf.getvalue(), 'x.xlsx')
        self.assertEqual(parsed, [{'item_code': 'IT001', 'counted_qty': Decimal('2')}])


class PickZoneApiTests(TestCase):
    """Management API: zones, rules, overrides, settings, preview, uncategorized."""

    def setUp(self):
        seed_default_pick_zones()
        make_branch()
        _, self.admin_profile, self.admin = make_admin('zone_admin')
        _, _, self.pharmacist = make_user('zone_ph', role='pharmacist')

    # ── zones ─────────────────────────────────────────────────────────────────

    def test_list_zones(self):
        r = self.admin.get('/api/transits/pick-zones/')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(len(r.data), 15)
        self.assertEqual(r.data[0]['name'], 'غوالى-ثلاجه')   # sort_key order
        self.assertIn('rule_count', r.data[0])

    def test_create_zone_admin_only(self):
        payload = {'name': 'مستلزمات طبية', 'sort_key': 150, 'location': 'ممر 9'}
        r = self.pharmacist.post('/api/transits/pick-zones/', payload, format='json')
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)
        r = self.admin.post('/api/transits/pick-zones/', payload, format='json')
        self.assertEqual(r.status_code, status.HTTP_201_CREATED)

    def test_pharmacist_can_read(self):
        r = self.pharmacist.get('/api/transits/pick-zones/')
        self.assertEqual(r.status_code, status.HTTP_200_OK)

    def test_special_roles_are_exclusive(self):
        tablets = PickZone.objects.get(name='أقراص')
        r = self.admin.patch(f'/api/transits/pick-zones/{tablets.id}/',
                             {'is_fridge_zone': True}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(PickZone.objects.filter(is_fridge_zone=True).count(), 1)
        self.assertTrue(PickZone.objects.get(name='أقراص').is_fridge_zone)
        self.assertFalse(PickZone.objects.get(name='غوالى-ثلاجه').is_fridge_zone)

    # ── settings ──────────────────────────────────────────────────────────────

    def test_threshold_get_and_update(self):
        r = self.admin.get('/api/transits/pick-zones/settings/')
        self.assertEqual(float(r.data['price_threshold']), 500.0)
        r = self.admin.post('/api/transits/pick-zones/settings/',
                            {'price_threshold': 750}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        r = self.admin.get('/api/transits/pick-zones/settings/')
        self.assertEqual(float(r.data['price_threshold']), 750.0)

    def test_threshold_rejects_garbage(self):
        r = self.admin.post('/api/transits/pick-zones/settings/',
                            {'price_threshold': 'abc'}, format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    # ── preview ───────────────────────────────────────────────────────────────

    def test_preview(self):
        r = self.admin.post('/api/transits/pick-zones/preview/',
                            {'name': 'TEST SYRUP 100ML', 'price': 60},
                            format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r.data['zone_name'], 'شراب')
        self.assertIn('SYRUP', r.data['reason'])

    def test_preview_requires_input(self):
        r = self.admin.post('/api/transits/pick-zones/preview/', {}, format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    # ── rules ─────────────────────────────────────────────────────────────────

    def test_rule_crud_and_reorder(self):
        zone = PickZone.objects.get(name='أقراص')
        r = self.admin.post('/api/transits/pick-rules/',
                            {'zone': zone.id, 'keywords': ['CHEWABLE', ''],
                             'priority': 5}, format='json')
        self.assertEqual(r.status_code, status.HTTP_201_CREATED)
        self.assertEqual(r.data['keywords'], ['CHEWABLE'])   # blanks stripped

        ids = list(PickZoneRule.objects.order_by('priority', 'id')
                   .values_list('id', flat=True))
        ids.reverse()
        r = self.admin.post('/api/transits/pick-rules/reorder/',
                            {'ordered_ids': ids}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        new_order = list(PickZoneRule.objects.order_by('priority', 'id')
                         .values_list('id', flat=True))
        self.assertEqual(new_order, ids)

    def test_rule_rejects_empty_keywords(self):
        zone = PickZone.objects.get(name='أقراص')
        r = self.admin.post('/api/transits/pick-rules/',
                            {'zone': zone.id, 'keywords': []}, format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    # ── overrides ─────────────────────────────────────────────────────────────

    def test_override_create_by_item_code_and_upsert(self):
        make_item('MYSTERY DEVICE', 'IT050')
        zone1 = PickZone.objects.get(name='ألبان')
        zone2 = PickZone.objects.get(name='شامبو')

        r = self.admin.post('/api/transits/item-overrides/',
                            {'item_code': 'IT050', 'zone': zone1.id,
                             'tag': 'رف A1'}, format='json')
        self.assertEqual(r.status_code, status.HTTP_201_CREATED)
        self.assertEqual(r.data['item_name'], 'MYSTERY DEVICE')

        # same item again → upsert, not duplicate
        r = self.admin.post('/api/transits/item-overrides/',
                            {'item_code': 'IT050', 'zone': zone2.id},
                            format='json')
        self.assertEqual(r.status_code, status.HTTP_201_CREATED)
        self.assertEqual(ItemPickOverride.objects.filter(
            item__softech_id='IT050').count(), 1)
        self.assertEqual(ItemPickOverride.objects.get(
            item__softech_id='IT050').zone, zone2)

    def test_override_rejects_unknown_item(self):
        zone = PickZone.objects.get(name='ألبان')
        r = self.admin.post('/api/transits/item-overrides/',
                            {'item_code': 'NOPE99', 'zone': zone.id},
                            format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    # ── uncategorized ─────────────────────────────────────────────────────────

    # ── per-location configs ──────────────────────────────────────────────────

    def test_branch_config_overrides_default(self):
        branch = make_branch2('مخزن فرعي', 'B77')
        # branch gets its own single-zone config
        z = PickZone.objects.create(branch=branch, name='منطقة وحيدة',
                                    sort_key=1, is_fallback=True)
        PickZoneRule.objects.create(zone=z, keywords=['TAB'], priority=10)

        rs_branch  = build_ruleset(branch=branch.id)
        rs_default = build_ruleset()
        zb, _ = rs_branch.classify('ZYRTEC 10MG 20TAB', 100)
        zd, _ = rs_default.classify('ZYRTEC 10MG 20TAB', 100)
        self.assertEqual(zb.name, 'منطقة وحيدة')
        self.assertEqual(zd.name, 'أقراص')

    def test_branch_without_config_falls_back_to_default(self):
        branch = make_branch2('مخزن بلا إعداد', 'B78')
        rs = build_ruleset(branch=branch.id)
        z, _ = rs.classify('ZYRTEC 10MG 20TAB', 100)
        self.assertEqual(z.name, 'أقراص')

    def test_copy_defaults_to_branch(self):
        branch = make_branch2('مخزن منسوخ', 'B79')
        r = self.admin.post('/api/transits/pick-zones/copy-defaults/',
                            {'branch': branch.id}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r.data['zones_created'], 15)
        self.assertEqual(r.data['rules_created'], 15)
        # second call refuses
        r = self.admin.post('/api/transits/pick-zones/copy-defaults/',
                            {'branch': branch.id}, format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        # scoped listing
        r = self.admin.get(f'/api/transits/pick-zones/?branch={branch.id}')
        self.assertEqual(len(r.data), 15)
        r = self.admin.get('/api/transits/pick-zones/?branch=')
        self.assertEqual(len(r.data), 15)   # default set unchanged

    def test_exclusive_roles_scoped_per_config(self):
        branch = make_branch2('مخزن أدوار', 'B80')
        self.admin.post('/api/transits/pick-zones/copy-defaults/',
                        {'branch': branch.id}, format='json')
        # flip fridge role inside the BRANCH config only
        z = PickZone.objects.get(branch=branch, name='أقراص')
        r = self.admin.patch(f'/api/transits/pick-zones/{z.id}/',
                             {'is_fridge_zone': True}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        # default config's fridge zone must be untouched
        self.assertTrue(PickZone.objects.get(
            branch__isnull=True, name='غوالى-ثلاجه').is_fridge_zone)
        self.assertFalse(PickZone.objects.get(
            branch=branch, name='غوالى-ثلاجه').is_fridge_zone)

    def test_override_branch_precedence(self):
        branch = make_branch2('مخزن تخصيص', 'B81')
        item = make_item('PRECEDENCE ITEM', 'IT070')
        z_milk    = PickZone.objects.get(name='ألبان')
        z_shampoo = PickZone.objects.get(name='شامبو')
        ItemPickOverride.objects.create(item=item, zone=z_milk, branch=None)
        ItemPickOverride.objects.create(item=item, zone=z_shampoo, branch=branch)

        z, _ = build_ruleset(branch=branch.id).classify('PRECEDENCE ITEM', 10,
                                                        itemcode='IT070')
        self.assertEqual(z.name, 'شامبو')                 # branch-specific wins
        z, _ = build_ruleset().classify('PRECEDENCE ITEM', 10, itemcode='IT070')
        self.assertEqual(z.name, 'ألبان')                 # default elsewhere

    # ── master-column rules ───────────────────────────────────────────────────

    def test_field_rule_matches_shape(self):
        item = make_item('MYSTERY FORM ITEM', 'IT071')
        item.shape_code = 'SH1'
        item.shape_name = 'Tablet'
        item.save()
        zone = PickZone.objects.get(name='لبوس')
        PickZoneRule.objects.create(zone=zone, match_field='shape',
                                    keywords=['SH1'], priority=1)
        rs = build_ruleset()
        z, _ = rs.classify('MYSTERY FORM ITEM', 10, itemcode='IT071',
                           attrs={'shape_code': 'SH1'})
        self.assertEqual(z.name, 'لبوس')
        # explain names the field
        _z, _t, reason = rs.explain('MYSTERY FORM ITEM', 10,
                                    attrs={'shape_code': 'SH1'})
        self.assertIn('شكل الصنف', reason)

    def test_field_values_endpoint(self):
        it = make_item('SHAPED ITEM', 'IT072')
        it.shape_code = 'SH9'
        it.shape_name_ar = 'أقراص فوارة'
        it.save()
        r = self.admin.get('/api/transits/pick-zones/field-values/?field=shape')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        row = next(v for v in r.data['values'] if v['value'] == 'SH9')
        self.assertEqual(row['label'], 'أقراص فوارة')
        r = self.admin.get('/api/transits/pick-zones/field-values/?field=bogus')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_classify_items_endpoint(self):
        it = make_item('BULK CLASSIFY 20TAB', 'IT073')
        it.pack_price = Decimal('50')
        it.save()
        # readable by non-managers too (products-table column)
        r = self.pharmacist.post('/api/transits/pick-zones/classify-items/',
                                 {'codes': ['IT073']}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r.data['IT073']['zone_name'], 'أقراص')
        self.assertIn('TAB', r.data['IT073']['reason'])

    def test_uncategorized_scan(self):
        make_item('STRANGE GADGET X', 'IT060')        # matches no rule
        make_item('NORMAL 20TAB', 'IT061')            # TAB rule
        r = self.admin.get('/api/transits/pick-zones/uncategorized/')
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        codes = [x['item_code'] for x in r.data['results']]
        self.assertIn('IT060', codes)
        self.assertNotIn('IT061', codes)

        # after an override the item disappears from the list
        zone = PickZone.objects.get(name='ألبان')
        self.admin.post('/api/transits/item-overrides/',
                        {'item_code': 'IT060', 'zone': zone.id}, format='json')
        r = self.admin.get('/api/transits/pick-zones/uncategorized/')
        codes = [x['item_code'] for x in r.data['results']]
        self.assertNotIn('IT060', codes)
