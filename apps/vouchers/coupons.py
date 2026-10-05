"""
apps/vouchers/coupons.py

Gift-coupon stocking — archive + serial generator.
Spec: docs/architecture/SOFTECH_GIFT_VOUCHER_STOCKING.md

One paper coupon (serial ``NNNNN-AAA999``) is stocked in SOFTECH on TWO items, both bought
from supplier 1268 into HQ store 100, one purchase line per serial (qty 1):
  * POINTS leg  — item 102230 (COUPON FOR POINTS), issued against loyalty points;
  * SERVED leg  — item 118639 (COUPON SERVED TO CUSTOMER), the −50 EGP line on a sale.
The serial lives in ``stktrans.item_partno``. The line expiry is NOT a real expiry: it is a
per-serial uniqueness key, because SOFTECH keeps on-hand per (item, store, expiry). A reused
date merges two serials into one stkbalexpiry row, so the generator never reuses a date.

This module never writes to SOFTECH. The ``read_*`` helpers are SELECT-only.
"""
import datetime as dt
import math
import re
import secrets
import string

from django.conf import settings
from django.db import transaction
from django.db.models import Max

SERIAL_RE = re.compile(r'^(\d{4,6})-([A-Z]{3}\d{3})$')

DATALOAD_HEADER = ['', '', '', 'Itemcode', 'Button', 'Next Cell', 'qty', 'Next Cell',
                   'Expiry Date', 'Next Cell', 'Serial Number', 'Button F1']


def points_item():
    return str(getattr(settings, 'COUPON_POINTS_ITEM', '102230'))


def served_item():
    return str(getattr(settings, 'COUPON_SERVED_ITEM', '118639'))


def supplier_code():
    return str(getattr(settings, 'COUPON_SUPPLIER', '1268'))


def hq_branch():
    return str(getattr(settings, 'COUPON_BRANCH', '100'))


def batch_size():
    return int(getattr(settings, 'COUPON_BATCH_SIZE', 200))


def leg_for_item(itemcode):
    code = str(itemcode or '').strip()
    if code == points_item():
        return 'points'
    if code == served_item():
        return 'served'
    return None


# ── pure helpers ───────────────────────────────────────────────────────────────
def parse_serial(value):
    """'27101-ZWU704' → (27101, 'ZWU704', '27101-ZWU704'); None when not a coupon serial."""
    s = str(value or '').strip().upper()
    m = SERIAL_RE.match(s)
    if not m:
        return None
    return int(m.group(1)), m.group(2), s


def format_serial(number, code):
    return f'{int(number)}-{code}'


def random_code(rng=None):
    """Same shape as the Excel generator: 3 uppercase letters + RANDBETWEEN(100, 999)."""
    rng = rng or secrets.SystemRandom()
    letters = ''.join(rng.choice(string.ascii_uppercase) for _ in range(3))
    return f'{letters}{rng.randint(100, 999)}'


def to_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value).strip()[:10]
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def to_docnumber(value):
    if value in (None, ''):
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def next_free_dates(n, start, blocked):
    """``n`` consecutive-as-possible dates from ``start``, skipping every date in ``blocked``."""
    out, d = [], start
    while len(out) < n:
        if d not in blocked:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def collated_order(n, per_page=4):
    """Print order for a 4-up sheet cut into stacks: slot j on page p shows coupon p + P·j.

    Reproduces the Excel `Printing Sheet` formulas (F1=D2, F11=D27, F21=D52, F31=D77 for a
    100-coupon run), so after cutting, every stack runs in serial order.
    """
    pages = math.ceil(n / per_page) if n else 0
    order = []
    for p in range(pages):
        for j in range(per_page):
            i = p + pages * j
            if i < n:
                order.append(i)
    return order


# ── archive import ─────────────────────────────────────────────────────────────
def _status_for(c):
    if c.status == 'void':
        return 'void'
    has_p, has_s = bool(c.points_docnumber), bool(c.served_docnumber)
    if has_p and has_s:
        return 'stocked'
    if has_p or has_s:
        return 'partial'
    return 'generated' if c.source == 'generated' else 'unstocked'


def _add_conflict(c, note):
    existing = [x for x in (c.conflict_note or '').split(' | ') if x]
    if note not in existing:
        existing.append(note)
    c.conflict_note = ' | '.join(existing)[:2000]


@transaction.atomic
def import_purchase_lines(rows):
    """Upsert SOFTECH coupon purchase lines (doccode 10, supplier 1268) into the archive.

    ``rows`` are dicts with at least itemcode, item_partno, itemexpirydate, docnumber,
    docdate (the probe CSV columns, or a live SELECT). Idempotent: re-importing the
    same lines changes nothing. A leg bought on two different documents, or a leg
    whose expiry differs from the other leg, is kept (earliest doc wins) and flagged.
    """
    from .models import CouponSerial

    report = {'lines': 0, 'created': 0, 'updated': 0, 'skipped_item': 0,
              'skipped_serial': 0, 'conflicts': 0, 'bad_serial_samples': []}
    existing = {c.serial: c for c in CouponSerial.objects.all()}
    touched, created = {}, []

    for r in rows:
        report['lines'] += 1
        leg = leg_for_item(r.get('itemcode'))
        if not leg:
            report['skipped_item'] += 1
            continue
        parsed = parse_serial(r.get('item_partno'))
        if not parsed:
            report['skipped_serial'] += 1
            raw = str(r.get('item_partno') or '').strip()
            if raw and len(report['bad_serial_samples']) < 20 and raw not in report['bad_serial_samples']:
                report['bad_serial_samples'].append(raw)
            continue
        number, code, serial = parsed
        expiry = to_date(r.get('itemexpirydate'))
        docnumber = to_docnumber(r.get('docnumber'))
        docdate = to_date(r.get('docdate'))

        c = existing.get(serial)
        if c is None:
            c = CouponSerial(serial=serial, number=number, code=code, source='softech')
            existing[serial] = c
            created.append(c)

        doc_f, date_f, exp_f = f'{leg}_docnumber', f'{leg}_docdate', f'{leg}_expiry'
        cur_doc = getattr(c, doc_f)
        if cur_doc and docnumber and int(cur_doc) != docnumber:
            _add_conflict(c, f'{leg} leg purchased on two docs: {int(cur_doc)} & {docnumber}')
            cur_date = getattr(c, date_f)
            if docdate and cur_date and docdate < cur_date:
                setattr(c, doc_f, docnumber)
                setattr(c, date_f, docdate)
                setattr(c, exp_f, expiry)
        elif not cur_doc:
            setattr(c, doc_f, docnumber)
            setattr(c, date_f, docdate)
            setattr(c, exp_f, expiry)
        if c.points_expiry and c.served_expiry and c.points_expiry != c.served_expiry:
            _add_conflict(c, f'leg expiries differ: points {c.points_expiry} / served {c.served_expiry}')
        if c.source == 'excel':
            c.source = 'softech'
        c.status = _status_for(c)
        touched[serial] = c

    new_ids = {id(c) for c in created}
    for c in created:
        c.status = _status_for(c)
    CouponSerial.objects.bulk_create(created, batch_size=1000)
    to_update = [c for c in touched.values() if id(c) not in new_ids]
    if to_update:
        CouponSerial.objects.bulk_update(
            to_update,
            ['points_docnumber', 'points_docdate', 'points_expiry', 'served_docnumber',
             'served_docdate', 'served_expiry', 'status', 'source', 'conflict_note'],
            batch_size=1000)
    report['created'] = len(created)
    report['updated'] = len(to_update)
    report['conflicts'] = flag_number_collisions()
    return report


@transaction.atomic
def import_excel_serials(entries):
    """Add serials from the Excel `Serial Database` that SOFTECH never received.

    ``entries`` = iterable of serial strings (or (number, code) tuples). Existing serials
    are left untouched. Returns a small report.
    """
    from .models import CouponSerial

    report = {'rows': 0, 'created': 0, 'already': 0, 'invalid': 0}
    known = set(CouponSerial.objects.values_list('serial', flat=True))
    new = []
    for e in entries:
        report['rows'] += 1
        raw = format_serial(*e) if isinstance(e, tuple) else e
        parsed = parse_serial(raw)
        if not parsed:
            report['invalid'] += 1
            continue
        number, code, serial = parsed
        if serial in known:
            report['already'] += 1
            continue
        known.add(serial)
        new.append(CouponSerial(serial=serial, number=number, code=code,
                                source='excel', status='unstocked'))
    CouponSerial.objects.bulk_create(new, batch_size=1000)
    report['created'] = len(new)
    report['conflicts'] = flag_number_collisions()
    return report


def flag_number_collisions():
    """Flag serials that share a running number or a random code with another serial
    (the legacy sheet re-issued 22504–22600 and reused 17101–17200's codes)."""
    from django.db.models import Count
    from .models import CouponSerial

    flagged = 0
    for field, label in (('number', 'number'), ('code', 'random code')):
        dups = (CouponSerial.objects.values(field).annotate(n=Count('id')).filter(n__gt=1)
                .values_list(field, flat=True))
        dup_set = set(dups)
        if not dup_set:
            continue
        rows = list(CouponSerial.objects.filter(**{f'{field}__in': dup_set}))
        by_val = {}
        for c in rows:
            by_val.setdefault(getattr(c, field), []).append(c.serial)
        for c in rows:
            others = sorted(s for s in by_val[getattr(c, field)] if s != c.serial)
            _add_conflict(c, f'{label} shared with {", ".join(others[:5])}')
        CouponSerial.objects.bulk_update(rows, ['conflict_note'], batch_size=1000)
        flagged += len(rows)
    return CouponSerial.objects.exclude(conflict_note='').count()


# ── generator ──────────────────────────────────────────────────────────────────
def used_expiries():
    """Every expiry date the archive has ever used, on either leg."""
    from .models import CouponSerial

    out = set()
    for p, s in CouponSerial.objects.values_list('points_expiry', 'served_expiry'):
        if p:
            out.add(p)
        if s:
            out.add(s)
    return out


def default_start_date(today=None):
    """The day after the latest expiry ever used — dates only move forward, never cycle —
    and at least COUPON_MIN_EXPIRY_DAYS (default 730) ahead of today."""
    from .models import CouponSerial

    today = today or dt.date.today()
    floor = today + dt.timedelta(days=int(getattr(settings, 'COUPON_MIN_EXPIRY_DAYS', 730)))
    agg = CouponSerial.objects.aggregate(p=Max('points_expiry'), s=Max('served_expiry'))
    latest = max([d for d in (agg['p'], agg['s']) if d] or [floor])
    return max(floor, latest + dt.timedelta(days=1))


@transaction.atomic
def generate_batch(*, size=None, created_by=None, start_number=None, start_date=None,
                   blocked_dates=None, rng=None, notes=''):
    """Create a CouponBatch of ``size`` new serials (status ``generated``).

    * numbers continue after the highest number in the archive (or ``start_number``);
    * random codes are unique against every code ever issued;
    * each serial gets ONE expiry date used for both legs, unused by the archive and by
      ``blocked_dates`` (pass the live SOFTECH stkbalexpiry dates for both items).
    Nothing is sent to SOFTECH.
    """
    from .models import CouponBatch, CouponSerial

    size = int(size or batch_size())
    if size <= 0:
        raise ValueError('size must be positive')
    # Serialise concurrent generators on the batch table.
    list(CouponBatch.objects.select_for_update().order_by('-id')[:1])

    rng = rng or secrets.SystemRandom()
    top = CouponSerial.objects.aggregate(m=Max('number'))['m'] or 0
    first = int(start_number) if start_number else top + 1
    if first <= top and CouponSerial.objects.filter(number__gte=first,
                                                     number__lt=first + size).exists():
        raise ValueError(f'serial numbers {first}–{first + size - 1} overlap the archive')

    codes_used = set(CouponSerial.objects.values_list('code', flat=True))
    blocked = used_expiries() | set(blocked_dates or ())
    start = start_date or default_start_date()
    dates = next_free_dates(size, start, blocked)

    batch = CouponBatch.objects.create(
        source='generated', status='generated', size=size,
        serial_from=first, serial_to=first + size - 1,
        expiry_from=dates[0], expiry_to=dates[-1],
        created_by=created_by, notes=notes)
    serials = []
    for i in range(size):
        code = random_code(rng)
        while code in codes_used:
            code = random_code(rng)
        codes_used.add(code)
        number = first + i
        serials.append(CouponSerial(
            serial=format_serial(number, code), number=number, code=code, batch=batch,
            source='generated', status='generated',
            points_expiry=dates[i], served_expiry=dates[i]))
    CouponSerial.objects.bulk_create(serials)
    return batch


# ── exports ────────────────────────────────────────────────────────────────────
def dataload_rows(batch):
    """Tab rows for the WorkBench DataLoad grid — the exact column layout of
    `Dataload_for_Coupon_Serial_Number.dld` (points leg first, then served leg).
    Interim fallback for the RPA until the SOFTECH push is approved."""
    serials = list(batch.serials.order_by('number'))
    rows = [DATALOAD_HEADER]
    for item, exp_field in ((points_item(), 'points_expiry'), (served_item(), 'served_expiry')):
        for c in serials:
            exp = getattr(c, exp_field)
            rows.append(['', '', '', item, r'\{ENTER}', r'\{TAB}', '1', r'\{TAB}',
                         exp.strftime('%d/%m/%Y'), r'\{TAB}', c.serial, r'\{F2}'])
    return rows


def print_workbook(batch, title=None):
    """openpyxl Workbook reproducing the Excel `Printing Sheet`: 10-row coupon blocks,
    4 per page, collated so the cut stacks come out in serial order."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.pagebreak import Break

    title = title or getattr(settings, 'COUPON_PRINT_TITLE',
                             'قسيمة مشتروات من صيدليات الرزيقى بقيمة 50ج.م')
    serials = list(batch.serials.order_by('number'))
    order = collated_order(len(serials))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Printing Sheet'
    ws.sheet_view.rightToLeft = False
    bold = Font(name='Arial', size=11, bold=True)
    title_font = Font(name='Aharoni', size=18)
    for slot, idx in enumerate(order):
        c = serials[idx]
        r0 = slot * 10 + 1
        ws.cell(r0, 4, 'Coupon')
        ws.cell(r0, 5, 'Serial')
        ws.cell(r0, 6, c.serial).font = bold
        ws.cell(r0, 7, 'الكوبون')
        ws.cell(r0, 8, 'سريال')
        ws.cell(r0 + 2, 4, 'Client Name:')
        ws.cell(r0 + 2, 10, 'إسم العميل:')
        ws.cell(r0 + 4, 4, 'Phone No.:')
        ws.cell(r0 + 4, 10, 'رقم التليفون:')
        ws.cell(r0 + 6, 4, 'PIC Code:')
        ws.cell(r0 + 6, 10, 'كود العميل:')
        t = ws.cell(r0 + 8, 4, title)
        t.font = title_font
        t.alignment = Alignment(horizontal='center')
        ws.merge_cells(start_row=r0 + 8, start_column=4, end_row=r0 + 8, end_column=9)
        if (slot + 1) % 4 == 0 and slot + 1 < len(order):
            ws.row_breaks.append(Break(id=r0 + 9))
    ws.column_dimensions['F'].width = 16
    return wb


# ── SOFTECH reads (SELECT only) ────────────────────────────────────────────────
_DB = 'SOFTECHDB9.dbo'


def read_purchase_lines(conn, since='2015-01-01'):
    """All coupon purchase lines from supplier 1268 at HQ (doccode 10) as dicts."""
    cur = conn.cursor()
    cur.execute(f"""
        SELECT docnumber FROM {_DB}.stktransm
        WHERE doccode='10' AND branchcode=? AND cust_branch_code=? AND docdate >= ?
    """, [hq_branch(), supplier_code(), since])
    docs = [to_docnumber(r[0]) for r in cur.fetchall()]
    out = []
    for dn in docs:
        c2 = conn.cursor()
        c2.execute(f"""
            SELECT docnumber, docdate, itemcode, itemexpirydate, item_partno
            FROM {_DB}.stktrans WHERE branchcode=? AND doccode='10' AND docnumber=?
        """, [hq_branch(), dn])
        for r in c2.fetchall():
            out.append({'docnumber': r[0], 'docdate': r[1], 'itemcode': str(r[2]).strip(),
                        'itemexpirydate': r[3], 'item_partno': r[4]})
    return out


def read_blocked_expiries(conn):
    """Every expiry date SOFTECH holds a stkbalexpiry row for, on either coupon item
    (any store, any qty) — receiving into one of these would merge serials."""
    cur = conn.cursor()
    cur.execute(f"""
        SELECT DISTINCT itemexpirydate FROM {_DB}.stkbalexpiry WHERE itemcode IN (?, ?)
    """, [points_item(), served_item()])
    return {d for d in (to_date(r[0]) for r in cur.fetchall()) if d}
