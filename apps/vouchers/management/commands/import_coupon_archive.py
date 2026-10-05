"""
python manage.py import_coupon_archive [--csv PATH] [--excel PATH] [--softech [--since YYYY-MM-DD]]

Builds the gift-coupon archive (CouponSerial) from:
  --csv     the probe's line archive docs/architecture/softech_gift_vouchers_lines.csv
  --softech the same lines read LIVE from SOFTECH (SELECT only)
  --excel   Coupon_Printing.xlsx › `Serial Database` (serials SOFTECH never received)

Idempotent: safe to re-run. SOFTECH data is applied before Excel. Nothing is written
to SOFTECH. Spec: docs/architecture/SOFTECH_GIFT_VOUCHER_STOCKING.md
"""
import csv

from django.core.management.base import BaseCommand, CommandError

from apps.vouchers import coupons


class Command(BaseCommand):
    help = 'Import gift-coupon serials (SOFTECH purchases from supplier 1268 + the Excel archive)'

    def add_arguments(self, parser):
        parser.add_argument('--csv', default='', help='probe CSV of stktrans lines')
        parser.add_argument('--softech', action='store_true', help='read the lines live from SOFTECH')
        parser.add_argument('--since', default='2015-01-01')
        parser.add_argument('--excel', default='', help='Coupon_Printing.xlsx')

    def handle(self, *args, **opts):
        if not (opts['csv'] or opts['softech'] or opts['excel']):
            raise CommandError('Give at least one of --csv, --softech, --excel')

        if opts['csv']:
            with open(opts['csv'], encoding='utf-8-sig', newline='') as f:
                rows = list(csv.DictReader(f))
            self._report('CSV', coupons.import_purchase_lines(rows))

        if opts['softech']:
            from config.sybase import get_sybase_connection
            conn = get_sybase_connection()
            try:
                rows = coupons.read_purchase_lines(conn, since=opts['since'])
            finally:
                conn.close()
            self._report('SOFTECH', coupons.import_purchase_lines(rows))

        if opts['excel']:
            from openpyxl import load_workbook
            wb = load_workbook(opts['excel'], read_only=True, data_only=True)
            if 'Serial Database' not in wb.sheetnames:
                raise CommandError("sheet 'Serial Database' not found")
            entries = []
            for row in wb['Serial Database'].iter_rows(min_row=2, values_only=True):
                full = row[3] if len(row) > 3 else None
                if full:
                    entries.append(str(full))
                elif len(row) > 2 and row[1] and row[2]:
                    entries.append((int(row[2]), str(row[1]).strip()))
            self._report('Excel', coupons.import_excel_serials(entries))

    def _report(self, label, rep):
        self.stdout.write(self.style.SUCCESS(f'[{label}] ' + ', '.join(
            f'{k}={v}' for k, v in rep.items() if k != 'bad_serial_samples')))
        if rep.get('bad_serial_samples'):
            self.stdout.write(f'  non-coupon item_partno samples: {rep["bad_serial_samples"]}')
